"""claims/delivered_human_scores.py: majority label of record over two or more reviewers'
filled Delivered workbooks.

No network, no LLM, no backend import. Interpreter: the system python (openpyxl).

Every workbook used here is built fresh with openpyxl in a temp directory (mirroring how
test_claims_score_review_sheet.py builds its own fixtures) and never touches a real reviewer
file.
"""

from __future__ import annotations

import json
from pathlib import Path

import delivered_human_scores as dhs
import export_review_sheet as ex
import openpyxl
import pytest
import score_review_sheet as sc

DEMO_RUN_ID = "run-test-01"
EXPORT_TIME = "2026-09-17T00:00:00+00:00"


# --------------------------------------------------------------------------------------
# Fixture building: write a delivered_evidence.json, build a delivered-only workbook from
# it, then fill in one reviewer's own answers.
# --------------------------------------------------------------------------------------


def _write_evidence(path: Path, rows: list[dict]) -> None:
    record = {
        "run_id": DEMO_RUN_ID, "draft_id": "draft-1", "section_title": "Section One",
        "rows": rows,
    }
    path.write_text(json.dumps(record), encoding="utf-8")


def _evidence_row(row: int, sentence: str, doi: str = "10.1/a", **overrides) -> dict:
    out = {
        "row": row, "sentence": sentence, "citation_text": "(Author, 2020)",
        "paper_title": "Paper", "paper_doi": doi, "claim_text": f"claim for row {row}",
        "evidence_quotes": ["an evidence quote"], "source_passage": "a source passage",
        "source_located": "yes", "passage_located": "yes", "sentence_in_draft": "yes",
    }
    out.update(overrides)
    return out


def _build_reviewer_workbook(
    tmp_path: Path, evidence_path: Path, name: str, fills: list[tuple[str, str, str]],
    *, demo_run_id: str = DEMO_RUN_ID, export_time: str = EXPORT_TIME,
) -> Path:
    """Builds a delivered-only workbook from *evidence_path* and fills its Delivered sheet's
    human columns, one ``(sentence_correct, quote_supports, note)`` triple per row in sheet
    order. Returns the saved workbook's path."""
    wb, _record, _rows, _workbook_id = ex.build_delivered_only_workbook(
        evidence_path, max_cell_chars=32000, demo_run_id=demo_run_id,
        export_time=export_time, paper_authors=None,
    )
    path = tmp_path / name
    wb.save(path)
    filled = openpyxl.load_workbook(path)
    ws = filled["Delivered"]
    header = [c.value for c in ws[1]]
    sc_col = header.index("human_sentence_correct") + 1
    qs_col = header.index("human_quote_supports") + 1
    note_col = header.index("human_note") + 1
    for i, (sentence_correct, quote_supports, note) in enumerate(fills, start=2):
        ws.cell(row=i, column=sc_col, value=sentence_correct)
        ws.cell(row=i, column=qs_col, value=quote_supports)
        ws.cell(row=i, column=note_col, value=note)
    filled.save(path)
    return path


# --------------------------------------------------------------------------------------
# Majority arithmetic, including the two-reviewer tie rule
# --------------------------------------------------------------------------------------


def test_majority_threshold_three_reviewers_is_two_of_three():
    assert dhs._majority_threshold(3) == 2


def test_majority_threshold_two_reviewers_needs_both():
    assert dhs._majority_threshold(2) == 2


def test_row_majority_two_of_three_yes_is_yes():
    aligned_row = {
        "reviewer_1": {"human_sentence_correct": "yes"},
        "reviewer_2": {"human_sentence_correct": "yes"},
        "reviewer_3": {"human_sentence_correct": "no"},
    }
    assert dhs.row_majority(
        aligned_row, ["reviewer_1", "reviewer_2", "reviewer_3"], "human_sentence_correct",
    ) == "yes"


def test_row_majority_one_of_three_yes_is_no():
    aligned_row = {
        "reviewer_1": {"human_sentence_correct": "yes"},
        "reviewer_2": {"human_sentence_correct": "no"},
        "reviewer_3": {"human_sentence_correct": "no"},
    }
    assert dhs.row_majority(
        aligned_row, ["reviewer_1", "reviewer_2", "reviewer_3"], "human_sentence_correct",
    ) == "no"


def test_row_majority_two_reviewers_tie_is_no():
    """With two reviewers, one yes and one no is a no: a tie is never a yes."""
    aligned_row = {
        "reviewer_1": {"human_sentence_correct": "yes"},
        "reviewer_2": {"human_sentence_correct": "no"},
    }
    assert dhs.row_majority(
        aligned_row, ["reviewer_1", "reviewer_2"], "human_sentence_correct",
    ) == "no"


def test_row_majority_two_reviewers_both_yes_is_yes():
    aligned_row = {
        "reviewer_1": {"human_sentence_correct": "yes"},
        "reviewer_2": {"human_sentence_correct": "yes"},
    }
    assert dhs.row_majority(
        aligned_row, ["reviewer_1", "reviewer_2"], "human_sentence_correct",
    ) == "yes"


def test_end_to_end_three_reviewer_majority_and_disagreement(tmp_path: Path):
    evidence_path = tmp_path / "evidence.json"
    _write_evidence(evidence_path, [
        _evidence_row(1, "Sentence one"), _evidence_row(2, "Sentence two"),
    ])
    # Row 1: sentence_correct yes/yes/no -> majority yes despite one dissent.
    # Row 2: quote_supports yes/no/yes -> majority yes despite one dissent.
    p1 = _build_reviewer_workbook(tmp_path, evidence_path, "r1.xlsx", [
        ("yes", "yes", ""), ("yes", "yes", ""),
    ])
    p2 = _build_reviewer_workbook(tmp_path, evidence_path, "r2.xlsx", [
        ("yes", "yes", ""), ("yes", "no", "quote is off point"),
    ])
    p3 = _build_reviewer_workbook(tmp_path, evidence_path, "r3.xlsx", [
        ("no", "yes", "sentence overstates it"), ("yes", "yes", ""),
    ])
    out = tmp_path / "scores.json"
    rc = dhs.main([
        "--workbook", str(p1), "--workbook", str(p2), "--workbook", str(p3),
        "--demo-run-id", DEMO_RUN_ID, "--export-time", EXPORT_TIME, "--out", str(out),
    ])
    assert rc == 0
    result = json.loads(out.read_text(encoding="utf-8"))
    assert result["n_rows"] == 2
    assert result["majority"]["sentence_correct_count"] == 2
    assert result["majority"]["quote_supports_count"] == 2
    rows_no = result["rows_judged_no"]
    assert len(rows_no) == 2
    assert {r["row"] for r in rows_no} == {1, 2}


# --------------------------------------------------------------------------------------
# Refusal on mismatched rows and on a wrong run id
# --------------------------------------------------------------------------------------


def test_align_rows_refuses_when_row_sets_differ(tmp_path: Path):
    evidence_a = tmp_path / "evidence_a.json"
    _write_evidence(evidence_a, [_evidence_row(1, "Sentence one")])
    evidence_b = tmp_path / "evidence_b.json"
    _write_evidence(evidence_b, [_evidence_row(1, "A different sentence")])

    p1 = _build_reviewer_workbook(tmp_path, evidence_a, "r1.xlsx", [("yes", "yes", "")])
    p2 = _build_reviewer_workbook(tmp_path, evidence_b, "r2.xlsx", [("yes", "yes", "")])
    rows1 = sc.read_delivered_sheet(p1)
    rows2 = sc.read_delivered_sheet(p2)
    with pytest.raises(dhs.ReviewSheetError):
        dhs.align_rows({"reviewer_1": rows1, "reviewer_2": rows2})


def test_main_refuses_on_mismatched_rows_and_writes_no_output(tmp_path: Path):
    evidence_a = tmp_path / "evidence_a.json"
    _write_evidence(evidence_a, [_evidence_row(1, "Sentence one")])
    evidence_b = tmp_path / "evidence_b.json"
    _write_evidence(evidence_b, [_evidence_row(1, "A different sentence")])
    p1 = _build_reviewer_workbook(tmp_path, evidence_a, "r1.xlsx", [("yes", "yes", "")])
    p2 = _build_reviewer_workbook(tmp_path, evidence_b, "r2.xlsx", [("yes", "yes", "")])
    out = tmp_path / "scores.json"
    rc = dhs.main([
        "--workbook", str(p1), "--workbook", str(p2),
        "--demo-run-id", DEMO_RUN_ID, "--export-time", EXPORT_TIME, "--out", str(out),
    ])
    assert rc == 2
    assert not out.exists()


def test_main_refuses_on_wrong_demo_run_id(tmp_path: Path):
    evidence_path = tmp_path / "evidence.json"
    _write_evidence(evidence_path, [_evidence_row(1, "Sentence one")])
    p1 = _build_reviewer_workbook(tmp_path, evidence_path, "r1.xlsx", [("yes", "yes", "")])
    p2 = _build_reviewer_workbook(tmp_path, evidence_path, "r2.xlsx", [("yes", "yes", "")])
    out = tmp_path / "scores.json"
    rc = dhs.main([
        "--workbook", str(p1), "--workbook", str(p2),
        "--demo-run-id", "wrong-run-id", "--export-time", EXPORT_TIME, "--out", str(out),
    ])
    assert rc == 2
    assert not out.exists()


def test_main_refuses_with_only_one_workbook(tmp_path: Path):
    evidence_path = tmp_path / "evidence.json"
    _write_evidence(evidence_path, [_evidence_row(1, "Sentence one")])
    p1 = _build_reviewer_workbook(tmp_path, evidence_path, "r1.xlsx", [("yes", "yes", "")])
    out = tmp_path / "scores.json"
    rc = dhs.main([
        "--workbook", str(p1),
        "--demo-run-id", DEMO_RUN_ID, "--export-time", EXPORT_TIME, "--out", str(out),
    ])
    assert rc == 2
    assert not out.exists()


# --------------------------------------------------------------------------------------
# Per-sentence aggregation over a multi-row sentence
# --------------------------------------------------------------------------------------


def test_per_sentence_aggregation_over_multi_row_sentence(tmp_path: Path):
    evidence_path = tmp_path / "evidence.json"
    _write_evidence(evidence_path, [
        _evidence_row(1, "Sentence A"), _evidence_row(2, "Sentence A"),
        _evidence_row(3, "Sentence B"),
    ])
    # Sentence A: two rows, majority yes on row 1, majority no on row 2 -> not faithful.
    # Sentence B: one row, majority yes -> faithful.
    p1 = _build_reviewer_workbook(tmp_path, evidence_path, "r1.xlsx", [
        ("yes", "yes", ""), ("no", "yes", "row 2 overstates it"), ("yes", "yes", ""),
    ])
    p2 = _build_reviewer_workbook(tmp_path, evidence_path, "r2.xlsx", [
        ("yes", "yes", ""), ("no", "yes", "row 2 overstates it"), ("yes", "yes", ""),
    ])
    out = tmp_path / "scores.json"
    rc = dhs.main([
        "--workbook", str(p1), "--workbook", str(p2),
        "--demo-run-id", DEMO_RUN_ID, "--export-time", EXPORT_TIME, "--out", str(out),
    ])
    assert rc == 0
    result = json.loads(out.read_text(encoding="utf-8"))
    by_sentence = result["majority"]["by_sentence"]
    assert by_sentence["n_sentences"] == 2
    assert by_sentence["sentence_correct"]["n_yes"] == 1
    assert by_sentence["quote_supports"]["n_yes"] == 2


# --------------------------------------------------------------------------------------
# Kappa on a known case
# --------------------------------------------------------------------------------------


def test_pairwise_kappa_matches_a_hand_worked_case(tmp_path: Path):
    """Reviewer A: yes, yes, no, no. Reviewer B: yes, no, no, no. po = 3/4 = 0.75;
    pe = (2*1 + 2*3) / 16 = 0.5; kappa = (0.75 - 0.5) / (1 - 0.5) = 0.5."""
    evidence_path = tmp_path / "evidence.json"
    _write_evidence(evidence_path, [
        _evidence_row(1, "S1"), _evidence_row(2, "S2"),
        _evidence_row(3, "S3"), _evidence_row(4, "S4"),
    ])
    p1 = _build_reviewer_workbook(tmp_path, evidence_path, "r1.xlsx", [
        ("yes", "yes", ""), ("yes", "yes", ""),
        ("no", "yes", "no"), ("no", "yes", "no"),
    ])
    p2 = _build_reviewer_workbook(tmp_path, evidence_path, "r2.xlsx", [
        ("yes", "yes", ""), ("no", "yes", "no"),
        ("no", "yes", "no"), ("no", "yes", "no"),
    ])
    out = tmp_path / "scores.json"
    rc = dhs.main([
        "--workbook", str(p1), "--workbook", str(p2),
        "--demo-run-id", DEMO_RUN_ID, "--export-time", EXPORT_TIME, "--out", str(out),
    ])
    assert rc == 0
    result = json.loads(out.read_text(encoding="utf-8"))
    pair = result["agreement"]["human_sentence_correct"]["pairs"]["reviewer_1_reviewer_2"]
    assert pair["kappa"] == pytest.approx(0.5)
    assert result["agreement"]["human_sentence_correct"]["mean_kappa"] == pytest.approx(0.5)


# --------------------------------------------------------------------------------------
# Deterministic output
# --------------------------------------------------------------------------------------


def test_output_bytes_are_deterministic_across_repeated_runs(tmp_path: Path):
    evidence_path = tmp_path / "evidence.json"
    _write_evidence(evidence_path, [
        _evidence_row(1, "Sentence one"), _evidence_row(2, "Sentence two"),
    ])
    p1 = _build_reviewer_workbook(tmp_path, evidence_path, "r1.xlsx", [
        ("yes", "yes", ""), ("no", "yes", "a reason"),
    ])
    p2 = _build_reviewer_workbook(tmp_path, evidence_path, "r2.xlsx", [
        ("yes", "no", "a different reason"), ("yes", "yes", ""),
    ])
    out1 = tmp_path / "scores1.json"
    out2 = tmp_path / "scores2.json"
    rc1 = dhs.main([
        "--workbook", str(p1), "--workbook", str(p2),
        "--demo-run-id", DEMO_RUN_ID, "--export-time", EXPORT_TIME, "--out", str(out1),
    ])
    rc2 = dhs.main([
        "--workbook", str(p1), "--workbook", str(p2),
        "--demo-run-id", DEMO_RUN_ID, "--export-time", EXPORT_TIME, "--out", str(out2),
    ])
    assert rc1 == 0
    assert rc2 == 0
    assert out1.read_bytes() == out2.read_bytes()


# --------------------------------------------------------------------------------------
# End to end: every top-level key is present
# --------------------------------------------------------------------------------------


def test_end_to_end_produces_every_top_level_key(tmp_path: Path):
    evidence_path = tmp_path / "evidence.json"
    _write_evidence(evidence_path, [
        _evidence_row(1, "Sentence one"), _evidence_row(2, "Sentence two"),
    ])
    p1 = _build_reviewer_workbook(tmp_path, evidence_path, "r1.xlsx", [
        ("yes", "yes", ""), ("yes", "yes", ""),
    ])
    p2 = _build_reviewer_workbook(tmp_path, evidence_path, "r2.xlsx", [
        ("yes", "yes", ""), ("yes", "yes", ""),
    ])
    p3 = _build_reviewer_workbook(tmp_path, evidence_path, "r3.xlsx", [
        ("yes", "yes", ""), ("yes", "yes", ""),
    ])
    out = tmp_path / "scores.json"
    rc = dhs.main([
        "--workbook", str(p1), "--workbook", str(p2), "--workbook", str(p3),
        "--demo-run-id", DEMO_RUN_ID, "--export-time", EXPORT_TIME, "--out", str(out),
    ])
    assert rc == 0
    result = json.loads(out.read_text(encoding="utf-8"))
    for key in (
        "demo_run_id", "export_time", "workbook_id", "n_reviewers", "reviewers", "unit",
        "majority_rule", "n_rows", "per_reviewer", "majority", "agreement", "rows_judged_no",
    ):
        assert key in result, key
    assert result["unit"] == "verified claim row"
    assert result["n_reviewers"] == 3
    assert set(result["reviewers"]) == {"reviewer_1", "reviewer_2", "reviewer_3"}
    assert set(result["per_reviewer"]) == {"reviewer_1", "reviewer_2", "reviewer_3"}
    assert result["majority"]["sentence_correct_rate"] == 1.0
    assert result["majority"]["quote_supports_rate"] == 1.0
    assert result["rows_judged_no"] == []
