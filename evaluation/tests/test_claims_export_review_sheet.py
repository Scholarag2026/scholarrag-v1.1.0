"""Tests for claims/export_review_sheet.py, the human-review workbook builder.

No network, no LLM, no backend import: openpyxl round trips over small in-memory/tmp_path
fixtures. Interpreter: the evaluation venv's python (openpyxl) or the system python.
"""

from __future__ import annotations

import csv
import hashlib
import json
import re
from pathlib import Path

import export_review_sheet as ex
import openpyxl
import pytest
from conftest import load_script_module

from build_hss_set import doi_slug
from common import read_jsonl, write_json

# --------------------------------------------------------------------------------------
# Fixtures
# --------------------------------------------------------------------------------------

CONSTRUCTED_ITEM = {
    "item_id": "hss-verbatim-01", "rule": "verbatim", "expected": ["verified"],
    "claim": "Scores were significantly higher for the treatment group after training.",
    "source_doi": "10.1000/paper-a", "chunk_doi": "10.1000/paper-a", "chunk_index": 1,
    "title": "Paper A", "authors": ["A. One"],
}
NO_FULL_TEXT_ITEM = {
    "item_id": "hss-no_full_text-01", "rule": "no_full_text", "expected": ["no_full_text"],
    "claim": "A claim whose paper was never fetched for this fixture.",
    "source_doi": "10.1000/paper-c", "chunk_doi": None, "chunk_index": 0,
    "title": "Paper Missing", "authors": ["C. Three"],
}
REAL_ITEM = {
    "item_id": "real-test-01", "rule": "real", "expected": None,
    "claim": "Feedback improved learner motivation across the term for most participants.",
    "cited_title": "Paper B", "cited_doi": "10.1000/paper-b",
    "source_doi": "10.1000/paper-b", "chunk_doi": "10.1000/paper-b", "chunk_index": 0,
    "title": "Paper B (as cited in review)", "authors": ["B. Two"],
}


def make_cache(cache_dir: Path, doi: str, title: str, chunks: list[str]) -> None:
    cache_dir.mkdir(parents=True, exist_ok=True)
    payload = {
        "doi": doi, "title": title, "authors": ["Author One"],
        "source_url": "https://example.org/x.pdf", "fetched_at": "2026-09-05T00:00:00+00:00",
        "n_chars": sum(len(c) for c in chunks),
        "chunks": [{"section": "results", "text": c} for c in chunks],
    }
    write_json(cache_dir / f"{doi_slug(doi)}.json", payload)


def result_row(item_id, status, *, model_status=None, evidence_quotes=None,
                machine_reasons=None, diagnostics=None, unstated_details=None,
                assertions=None, explanation="the explanation", quote_relocations=None,
                passes=None):
    return {
        "item_id": item_id,
        "predicted_status": status,
        "model_status": model_status if model_status is not None else status,
        "evidence_quotes": evidence_quotes if evidence_quotes is not None else [],
        "evidence_quote": (evidence_quotes or [None])[0],
        "machine_reasons": machine_reasons or [],
        "diagnostics": diagnostics or [],
        "unstated_details": unstated_details or [],
        "assertions": assertions or [],
        "explanation": explanation,
        "quote_relocations": quote_relocations or [],
        "passes": passes or [],
    }


def write_jsonl(path: Path, rows) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as fh:
        for row in rows:
            fh.write(json.dumps(row) + "\n")


# --------------------------------------------------------------------------------------
# Pure helpers
# --------------------------------------------------------------------------------------


def test_classify_set_real_by_literal_rule():
    assert ex.classify_set({"rule": "real"}) == "real"


def test_classify_set_constructed_for_any_other_rule():
    for rule in ("verbatim", "paraphrase", "altered", "over_specified", "wrong_paper",
                 "no_full_text"):
        assert ex.classify_set({"rule": rule}) == "constructed"


def test_cited_title_and_doi_real_item_uses_its_own_fields():
    title, doi = ex.cited_title_and_doi(REAL_ITEM)
    assert title == "Paper B"
    assert doi == "10.1000/paper-b"


def test_cited_title_and_doi_constructed_item_uses_chunk_doi():
    title, doi = ex.cited_title_and_doi(CONSTRUCTED_ITEM)
    assert title == "Paper A"
    assert doi == "10.1000/paper-a"


def test_cited_title_and_doi_no_full_text_falls_back_to_source_doi():
    title, doi = ex.cited_title_and_doi(NO_FULL_TEXT_ITEM)
    assert title == "Paper Missing"
    assert doi == "10.1000/paper-c"


def test_joined_list_drops_falsy_and_joins_with_semicolon():
    assert ex.joined_list(["a", "", None, "b"]) == "a; b"
    assert ex.joined_list(None) == ""


def test_first_evidence_quote_prefers_evidence_quotes_list():
    row = {"evidence_quotes": ["first", "second"], "evidence_quote": "single"}
    assert ex.first_evidence_quote(row) == "first"


def test_first_evidence_quote_falls_back_to_singular_field():
    row = {"evidence_quotes": [], "evidence_quote": "single"}
    assert ex.first_evidence_quote(row) == "single"


def test_all_evidence_quotes_joins_every_span():
    row = {"evidence_quotes": ["a", "b", "c"]}
    assert ex.all_evidence_quotes(row) == "a; b; c"


# A bare "null" string in evidence_quote/evidence_quotes is the verifier's own placeholder for
# "no evidence", indistinguishable from a real quote once it lands in a cell; both functions
# must treat it the same as a genuinely empty field.


def test_first_evidence_quote_treats_bare_null_string_as_empty():
    row = {"evidence_quotes": [], "evidence_quote": "null"}
    assert ex.first_evidence_quote(row) == ""


def test_first_evidence_quote_treats_null_case_and_whitespace_insensitively():
    row = {"evidence_quotes": [], "evidence_quote": " NuLL \n"}
    assert ex.first_evidence_quote(row) == ""


def test_first_evidence_quote_skips_null_entries_in_the_list():
    row = {"evidence_quotes": ["null"], "evidence_quote": "single"}
    assert ex.first_evidence_quote(row) == "single"


def test_first_evidence_quote_real_null_word_in_context_is_kept():
    # A real quote that happens to contain the word "null" (not equal to it) is not the
    # placeholder and must be kept.
    row = {"evidence_quotes": [], "evidence_quote": "the result was null and void"}
    assert ex.first_evidence_quote(row) == "the result was null and void"


def test_all_evidence_quotes_drops_bare_null_entries():
    row = {"evidence_quotes": ["null"]}
    assert ex.all_evidence_quotes(row) == ""


def test_all_evidence_quotes_drops_null_but_keeps_real_spans():
    row = {"evidence_quotes": ["a", "null", "b"]}
    assert ex.all_evidence_quotes(row) == "a; b"


def test_all_evidence_quotes_singular_fallback_drops_null():
    row = {"evidence_quotes": [], "evidence_quote": "null"}
    assert ex.all_evidence_quotes(row) == ""


def test_central_assertion_text_single_central():
    row = {"assertions": [
        {"text": "peripheral one", "central": False},
        {"text": "the central one", "central": True},
    ]}
    assert ex.central_assertion_text(row) == "the central one"


def test_central_assertion_text_empty_when_zero_central():
    row = {"assertions": [{"text": "peripheral one", "central": False}]}
    assert ex.central_assertion_text(row) == ""


def test_central_assertion_text_empty_when_multiple_central():
    row = {"assertions": [
        {"text": "one", "central": True}, {"text": "two", "central": True},
    ]}
    assert ex.central_assertion_text(row) == ""


def test_run_b_status_if_differs_blank_when_same():
    row_a = {"predicted_status": "verified"}
    row_b = {"predicted_status": "verified"}
    assert ex.run_b_status_if_differs(row_a, row_b) == ""


def test_run_b_status_if_differs_filled_when_different():
    row_a = {"predicted_status": "verified"}
    row_b = {"predicted_status": "needs_nuance"}
    assert ex.run_b_status_if_differs(row_a, row_b) == "needs_nuance"


def test_run_b_status_if_differs_blank_when_no_run_b():
    assert ex.run_b_status_if_differs({"predicted_status": "verified"}, None) == ""


def test_cap_text_under_limit_unchanged():
    assert ex.cap_text("short", 100) == "short"


def test_cap_text_over_limit_appends_marker():
    capped = ex.cap_text("x" * 50, 20)
    assert len(capped) == 20
    assert capped.endswith(ex.TRUNCATION_MARKER[-5:])


def test_split_into_parts_empty_text():
    assert ex.split_into_parts("", 10) == []


def test_split_into_parts_splits_evenly():
    assert ex.split_into_parts("abcdefgh", 3) == ["abc", "def", "gh"]


# --------------------------------------------------------------------------------------
# Blinding
# --------------------------------------------------------------------------------------


def test_assign_opaque_ids_deterministic_for_same_seed():
    items = [{"item_id": f"i-{i}"} for i in range(10)]
    _, ids_a = ex.assign_opaque_ids(items, seed=42)
    _, ids_b = ex.assign_opaque_ids(items, seed=42)
    assert ids_a == ids_b


def test_assign_opaque_ids_never_reveals_real_item_id():
    items = [{"item_id": "hss-altered-07"}]
    _, ids = ex.assign_opaque_ids(items, seed=1)
    assert ids["hss-altered-07"] != "hss-altered-07"
    assert ids["hss-altered-07"].startswith("R-")


def test_compute_workbook_id_changes_with_seed():
    ids = ["a", "b", "c"]
    assert ex.compute_workbook_id(ids, 1) != ex.compute_workbook_id(ids, 2)


def test_compute_workbook_id_changes_with_item_order():
    assert ex.compute_workbook_id(["a", "b"], 1) != ex.compute_workbook_id(["b", "a"], 1)


# --------------------------------------------------------------------------------------
# Review row assembly
# --------------------------------------------------------------------------------------


def test_build_review_row_all_columns_present():
    row_a = result_row(
        "hss-verbatim-01", "verified", evidence_quotes=["Scores were significantly higher."],
        machine_reasons=["attribution_mismatch"], diagnostics=["centrality_unmarked"],
        unstated_details=["some detail"],
        assertions=[{"text": "the central claim", "central": True}],
    )
    row_b = result_row("hss-verbatim-01", "needs_nuance")
    results_by_run = {"A": {"hss-verbatim-01": row_a}, "B": {"hss-verbatim-01": row_b}}
    out = ex.build_review_row(
        CONSTRUCTED_ITEM, results_by_run, "R-001", runs=["A", "B"], max_cell_chars=32000,
    )
    assert out["item_id"] == "R-001"
    assert out["set"] == "constructed"
    assert out["cited_title"] == "Paper A"
    assert out["cited_doi"] == "10.1000/paper-a"
    assert out["evidence_quote"] == "Scores were significantly higher."
    assert out["quote_relocated_original"] == ""
    assert out["evidence_quotes_all"] == "Scores were significantly higher."
    assert out["verifier_status"] == "verified"
    assert out["model_status"] == "verified"
    assert out["machine_reasons"] == "attribution_mismatch"
    assert out["diagnostics"] == "centrality_unmarked"
    assert out["unstated_details"] == "some detail"
    assert out["central_assertion"] == "the central claim"
    assert out["run_B_status"] == "needs_nuance"
    for human_field in (
        "human_verdict", "human_agrees_with_verifier", "disagreement_axis", "human_quote",
        "human_justification", "countersigned_by", "countersigned_at",
    ):
        assert out[human_field] == ""


# --------------------------------------------------------------------------------------
# The stored evidence_quote is the source's own text once a quote is relocated, so the
# reviewer countersigning the row needs a separate look at what the model actually wrote.
# --------------------------------------------------------------------------------------


def test_relocated_quote_original_blank_when_no_relocation():
    row = result_row("hss-verbatim-01", "verified", evidence_quotes=["Scores were higher."])
    assert ex.relocated_quote_original(row) == ""


def test_relocated_quote_original_surfaces_the_models_own_text():
    row = result_row(
        "hss-verbatim-01", "verified", evidence_quotes=["Scores were significantly higher."],
        diagnostics=["quote_relocated"],
        quote_relocations=[
            {
                "type": "quote_relocated",
                "original": "Scores were we significantly higher.",
                "replacement": "Scores were significantly higher.",
                "edit_distance": 1,
                "chunk_index": 0,
            }
        ],
    )
    assert ex.relocated_quote_original(row) == "Scores were we significantly higher."


def test_relocated_quote_original_joins_multiple_relocations():
    row = result_row(
        "hss-verbatim-01", "verified",
        quote_relocations=[{"original": "first original"}, {"original": "second original"}],
    )
    assert ex.relocated_quote_original(row) == "first original; second original"


def test_build_review_row_surfaces_the_relocated_quotes_original_text():
    row_a = result_row(
        "hss-verbatim-01", "verified", evidence_quotes=["Scores were significantly higher."],
        diagnostics=["quote_relocated"],
        quote_relocations=[{"original": "Scores were we significantly higher."}],
    )
    results_by_run = {"A": {"hss-verbatim-01": row_a}}
    out = ex.build_review_row(
        CONSTRUCTED_ITEM, results_by_run, "R-001", runs=["A"], max_cell_chars=32000,
    )
    assert out["evidence_quote"] == "Scores were significantly higher."
    assert out["quote_relocated_original"] == "Scores were we significantly higher."


# --------------------------------------------------------------------------------------
# The bounded second pass surfaces the first pass's own demotion, shown beside the final
# verifier_status/model_status.
# --------------------------------------------------------------------------------------


def test_first_pass_demotion_blank_when_only_one_pass():
    row = result_row(
        "hss-verbatim-01", "verified",
        passes=[{"model_status": "verified", "machine_reasons": []}],
    )
    assert ex.first_pass_demotion(row) == ""


def test_first_pass_demotion_blank_when_no_passes_recorded_at_all():
    row = result_row("hss-verbatim-01", "verified")
    assert ex.first_pass_demotion(row) == ""


def test_first_pass_demotion_surfaces_the_first_passs_own_reasons():
    row = result_row(
        "hss-verbatim-01", "verified",
        passes=[
            {"model_status": "verified", "machine_reasons": ["quote_not_verbatim"]},
            {"model_status": "verified", "machine_reasons": []},
        ],
    )
    assert ex.first_pass_demotion(row) == "quote_not_verbatim"


def test_build_review_row_surfaces_the_first_pass_demotion():
    row_a = result_row(
        "hss-verbatim-01", "verified",
        passes=[
            {"model_status": "verified", "machine_reasons": ["quote_not_verbatim"]},
            {"model_status": "verified", "machine_reasons": []},
        ],
    )
    results_by_run = {"A": {"hss-verbatim-01": row_a}}
    out = ex.build_review_row(
        CONSTRUCTED_ITEM, results_by_run, "R-001", runs=["A"], max_cell_chars=32000,
    )
    assert out["first_pass_demotion"] == "quote_not_verbatim"
    assert out["verifier_status"] == "verified"


def test_build_review_row_no_full_text_item_leaves_relocated_column_blank():
    """A withheld row's evidence_quote is already overridden to the withheld marker; the
    relocation column must not leak whatever a stub row happens to carry either."""
    row_a = result_row(
        "hss-no_full_text-01", "no_full_text",
        quote_relocations=[{"original": "should never surface"}],
    )
    results_by_run = {"A": {"hss-no_full_text-01": row_a}}
    out = ex.build_review_row(
        NO_FULL_TEXT_ITEM, results_by_run, "R-003", runs=["A"], max_cell_chars=32000,
    )
    assert out["evidence_quote"] == ex.NO_FULL_TEXT_EVIDENCE_MARKER
    assert out["quote_relocated_original"] == ""


def test_review_header_places_relocated_column_next_to_evidence_quote():
    header = ex.review_header()
    assert header.index("quote_relocated_original") == header.index("evidence_quote") + 1


def test_review_header_places_first_pass_demotion_next_to_model_status():
    header = ex.review_header()
    assert header.index("first_pass_demotion") == header.index("model_status") + 1


def test_build_review_row_single_run_leaves_run_b_status_blank():
    row_a = result_row("real-test-01", "verified")
    results_by_run = {"A": {"real-test-01": row_a}}
    out = ex.build_review_row(
        REAL_ITEM, results_by_run, "R-002", runs=["A"], max_cell_chars=32000,
    )
    assert out["run_B_status"] == ""


def test_evidence_withheld_true_for_no_chunk_doi():
    assert ex.evidence_withheld(NO_FULL_TEXT_ITEM) is True


def test_evidence_withheld_false_when_chunk_doi_present():
    assert ex.evidence_withheld(CONSTRUCTED_ITEM) is False


def test_build_review_row_no_full_text_item_carries_withheld_marker():
    # A no_full_text item's evidence_quote must say plainly that the tool was denied the text,
    # even though the same underlying paper may appear on the Sources sheet for a *different*
    # item that does carry a chunk_doi.
    row_a = result_row("hss-no_full_text-01", "no_full_text")
    results_by_run = {"A": {"hss-no_full_text-01": row_a}}
    out = ex.build_review_row(
        NO_FULL_TEXT_ITEM, results_by_run, "R-003", runs=["A"], max_cell_chars=32000,
    )
    assert out["evidence_quote"] == ex.NO_FULL_TEXT_EVIDENCE_MARKER


def test_build_review_row_normal_item_never_carries_withheld_marker():
    row_a = result_row(
        "hss-verbatim-01", "verified", evidence_quotes=["Scores were significantly higher."],
    )
    results_by_run = {"A": {"hss-verbatim-01": row_a}}
    out = ex.build_review_row(
        CONSTRUCTED_ITEM, results_by_run, "R-001", runs=["A"], max_cell_chars=32000,
    )
    assert out["evidence_quote"] != ex.NO_FULL_TEXT_EVIDENCE_MARKER


def test_build_key_row_shape():
    row = ex.build_key_row(CONSTRUCTED_ITEM, "R-001", "wbid123")
    assert row == {
        "opaque_id": "R-001", "real_item_id": "hss-verbatim-01", "set": "constructed",
        "construction_category": "verbatim", "expected_label": "verified",
        "workbook_id": "wbid123",
    }


def test_build_key_row_real_item_has_no_expected_label():
    row = ex.build_key_row(REAL_ITEM, "R-002", "wbid123")
    assert row["expected_label"] == ""
    assert row["set"] == "real"


# --------------------------------------------------------------------------------------
# Model_labels
# --------------------------------------------------------------------------------------


def test_load_annotation_labels(tmp_path: Path):
    path = tmp_path / "ann.csv"
    with path.open("w", encoding="utf-8", newline="") as fh:
        writer = csv.DictWriter(
            fh, fieldnames=["item_id", "final_label", "source", "rationale"],
        )
        writer.writeheader()
        writer.writerow({
            "item_id": "hss-verbatim-01", "final_label": "verified", "source": "unanimous",
            "rationale": "matches verbatim",
        })
    out = ex.load_annotation_labels([path])
    assert out["hss-verbatim-01"] == {
        "label": "verified", "rationale": "matches verbatim", "agreement": "unanimous",
    }


def test_build_model_labels_rows_only_covers_annotated_items():
    items = [CONSTRUCTED_ITEM, REAL_ITEM]
    opaque_ids = {"hss-verbatim-01": "R-001", "real-test-01": "R-002"}
    annotations = {
        "hss-verbatim-01": {"label": "verified", "rationale": "r1", "agreement": "unanimous"},
    }
    rows = ex.build_model_labels_rows(items, opaque_ids, annotations, max_cell_chars=1000)
    assert len(rows) == 1
    assert rows[0]["item_id"] == "R-001"
    assert rows[0]["model_annotation_label"] == "verified"
    assert rows[0]["annotator_agreement"] == "unanimous"


# --------------------------------------------------------------------------------------
# Sources sheet: keyed by item_id, not DOI
# --------------------------------------------------------------------------------------


def test_build_source_rows_withheld_item_absent_under_any_key():
    caches = {"dir": {"10.1/x": {"chunks": [{"text": "full text of x"}]}}}
    items = [
        {"item_id": "i1", "chunk_doi": "10.1/x", "_cache_dir": "dir"},
        {"item_id": "i2", "chunk_doi": None, "_cache_dir": "dir"},
    ]
    opaque_ids = {"i1": "R-001", "i2": "R-002"}
    rows = ex.build_source_rows(items, caches, opaque_ids, max_chars=1000)
    item_ids = {r[0] for r in rows}
    assert item_ids == {"R-001"}


def test_build_source_rows_duplicates_text_per_citing_item():
    caches = {"dir": {"10.1/x": {"chunks": [{"text": "shared paper text"}]}}}
    items = [
        {"item_id": "i1", "chunk_doi": "10.1/x", "_cache_dir": "dir"},
        {"item_id": "i2", "chunk_doi": "10.1/x", "_cache_dir": "dir"},
    ]
    opaque_ids = {"i1": "R-001", "i2": "R-002"}
    rows = ex.build_source_rows(items, caches, opaque_ids, max_chars=1000)
    by_item = {item_id: text for item_id, _part, text in rows}
    assert len(rows) == 2
    assert by_item["R-001"] == "shared paper text"
    assert by_item["R-002"] == "shared paper text"


def test_write_sources_sheet_header_is_item_id():
    wb = openpyxl.Workbook()
    ex.write_sources_sheet(wb, [("R-001", 1, "some text")])
    header = [c.value for c in wb["Sources"][1]]
    assert header == ["item_id", "part", "text"]


# --------------------------------------------------------------------------------------
# Provenance
# --------------------------------------------------------------------------------------


def test_resolve_guard_commit_explicit_wins_when_clean(monkeypatch):
    monkeypatch.setattr(ex, "_guard_path_is_dirty", lambda: False)
    assert ex.resolve_guard_commit("abc123") == "abc123"


def test_resolve_guard_commit_falls_back_when_no_git(monkeypatch):
    def fake_run(*args, **kwargs):
        raise FileNotFoundError("no git")
    monkeypatch.setattr(ex.subprocess, "run", fake_run)
    assert ex.resolve_guard_commit(None) == "unknown"


def test_resolve_guard_commit_appends_dirty_marker_when_explicit(monkeypatch):
    monkeypatch.setattr(ex, "_guard_path_is_dirty", lambda: True)
    result = ex.resolve_guard_commit("abc123")
    assert result.startswith("abc123")
    assert "dirty" in result
    assert "fulltext.py" in result


def test_resolve_guard_commit_appends_dirty_marker_when_auto_detected(monkeypatch):
    def fake_run(args, **kwargs):
        class Result:
            returncode = 0
            stdout = "deadbeef\n" if args[1] == "log" else ""
        return Result()
    monkeypatch.setattr(ex.subprocess, "run", fake_run)
    monkeypatch.setattr(ex, "_guard_path_is_dirty", lambda: True)
    result = ex.resolve_guard_commit(None)
    assert result.startswith("deadbeef")
    assert "dirty" in result


def test_resolve_guard_commit_no_dirty_marker_when_clean(monkeypatch):
    monkeypatch.setattr(ex, "_guard_path_is_dirty", lambda: False)
    assert ex.resolve_guard_commit("abc123") == "abc123"


def test_guard_path_is_dirty_false_when_porcelain_empty(monkeypatch):
    def fake_run(*args, **kwargs):
        class Result:
            returncode = 0
            stdout = ""
        return Result()
    monkeypatch.setattr(ex.subprocess, "run", fake_run)
    assert ex._guard_path_is_dirty() is False


def test_guard_path_is_dirty_true_when_porcelain_nonempty(monkeypatch):
    def fake_run(*args, **kwargs):
        class Result:
            returncode = 0
            stdout = " M backend/app/services/fulltext.py\n"
        return Result()
    monkeypatch.setattr(ex.subprocess, "run", fake_run)
    assert ex._guard_path_is_dirty() is True


def test_guard_path_is_dirty_false_when_no_git(monkeypatch):
    def fake_run(*args, **kwargs):
        raise FileNotFoundError("no git")
    monkeypatch.setattr(ex.subprocess, "run", fake_run)
    assert ex._guard_path_is_dirty() is False


def test_provenance_rows_shape():
    rows = ex.provenance_rows(
        verifier_prompt_shas=["sha256:abc"], guard_commit="deadbeef",
        results_dir=Path("results/v3"), export_time="2026-09-08T00:00:00+00:00",
        demo_run_id="screening_record", screening_prompt_version="sha256:def",
    )
    fields = {r["field"]: r["value"] for r in rows}
    assert fields["verifier_prompt_sha"] == "sha256:abc"
    assert fields["guard_commit"] == "deadbeef"
    assert fields["demo_run_id"] == "screening_record"
    assert fields["screening_prompt_version"] == "sha256:def"


# The provenance table's nine facts must be on the sheet itself; provenance_rows carries six of
# them with defaults that hold even when a caller (or an old test) does not pass them.


def test_provenance_rows_carries_the_six_facts_m4_added():
    rows = ex.provenance_rows(
        verifier_prompt_shas=["sha256:abc"], guard_commit="deadbeef",
        results_dir=Path("results/v3"), export_time="2026-09-08T00:00:00+00:00",
        demo_run_id="screening_record", screening_prompt_version="sha256:def",
        run_shown=["hss_runA.jsonl", "real_runA.jsonl"], run_commit="e0acccb",
        model="deepseek-v4-flash, temperature 0.0, 1 session per run",
        labels="hss_annotation_adjudicated_v3.csv (sha256:c9e4297...)",
    )
    fields = {r["field"]: r["value"] for r in rows}
    assert fields["run_shown"] == "hss_runA.jsonl; real_runA.jsonl"
    assert fields["run_commit"] == "e0acccb"
    assert fields["model"] == "deepseek-v4-flash, temperature 0.0, 1 session per run"
    assert fields["labels"] == "hss_annotation_adjudicated_v3.csv (sha256:c9e4297...)"
    assert fields["labels_provenance"] == ex.LABELS_PROVENANCE_FACT
    assert fields["blindness"] == ex.BLINDNESS_FACT
    assert "not human experts" in fields["labels_provenance"]
    assert "not blind" in fields["blindness"]


def test_provenance_rows_defaults_when_the_new_facts_are_not_given():
    rows = ex.provenance_rows(
        verifier_prompt_shas=["sha256:abc"], guard_commit="deadbeef",
        results_dir=Path("results/v3"), export_time="2026-09-08T00:00:00+00:00",
        demo_run_id="screening_record", screening_prompt_version="sha256:def",
    )
    fields = {r["field"]: r["value"] for r in rows}
    assert fields["run_shown"] == "unknown"
    assert fields["run_commit"] == "unknown"
    assert fields["model"] == "unknown"
    assert fields["labels"] == "n/a"


# Two per-run meta fields (``verification_policy_version``, ``repair_prompt_version``) have no
# v3/v4 equivalent; the Provenance sheet must show them whenever a run records them, and fall
# back the same way the other optional facts do when a run (v3, v4) does not.


def test_provenance_rows_carries_policy_and_repair_prompt_version():
    rows = ex.provenance_rows(
        verifier_prompt_shas=["sha256:abc"], guard_commit="deadbeef",
        results_dir=Path("results/v5"), export_time="2026-09-11T00:00:00+00:00",
        demo_run_id="screening_record", screening_prompt_version="sha256:def",
        policy_version="667bcc9209126ddc", repair_prompt_version="sha256:7ec7cffe852d",
    )
    fields = {r["field"]: r["value"] for r in rows}
    assert fields["verification_policy_version"] == "667bcc9209126ddc"
    assert fields["repair_prompt_version"] == "sha256:7ec7cffe852d"


def test_provenance_rows_defaults_policy_and_repair_prompt_version_when_not_given():
    rows = ex.provenance_rows(
        verifier_prompt_shas=["sha256:abc"], guard_commit="deadbeef",
        results_dir=Path("results/v3"), export_time="2026-09-08T00:00:00+00:00",
        demo_run_id="screening_record", screening_prompt_version="sha256:def",
    )
    fields = {r["field"]: r["value"] for r in rows}
    assert fields["verification_policy_version"] == "unknown"
    assert fields["repair_prompt_version"] == "unknown"


def test_load_run_meta_field_versions_reads_the_named_field(tmp_path: Path):
    results_dir = tmp_path / "results"
    write_json(results_dir / "hss_runA.meta.json", {"verification_policy_version": "abc123"})
    write_json(results_dir / "real_runA.meta.json", {"verification_policy_version": "abc123"})
    assert ex.load_run_meta_field_versions(
        results_dir, "A", "verification_policy_version",
    ) == ["abc123"]


def test_load_run_meta_field_versions_distinct_values_only(tmp_path: Path):
    results_dir = tmp_path / "results"
    write_json(results_dir / "hss_runA.meta.json", {"repair_prompt_version": "sha256:one"})
    write_json(results_dir / "real_runA.meta.json", {"repair_prompt_version": "sha256:two"})
    assert ex.load_run_meta_field_versions(
        results_dir, "A", "repair_prompt_version",
    ) == ["sha256:one", "sha256:two"]


def test_load_run_meta_field_versions_empty_when_field_absent(tmp_path: Path):
    results_dir = tmp_path / "results"
    write_json(results_dir / "hss_runA.meta.json", {"prompt_version": "sha256:testprompt"})
    assert ex.load_run_meta_field_versions(results_dir, "A", "repair_prompt_version") == []


def test_load_run_prompt_versions_is_the_prompt_version_case_of_the_general_reader(
    tmp_path: Path,
):
    results_dir = tmp_path / "results"
    write_json(results_dir / "hss_runA.meta.json", {"prompt_version": "sha256:testprompt"})
    assert ex.load_run_prompt_versions(results_dir, "A") == ["sha256:testprompt"]


def test_run_files_shown_lists_the_primary_runs_own_files(tmp_path: Path):
    results_dir = tmp_path / "results"
    write_jsonl(results_dir / "hss_runA.jsonl", [{"item_id": "x"}])
    write_jsonl(results_dir / "real_runA.jsonl", [{"item_id": "y"}])
    write_jsonl(results_dir / "hss_runB.jsonl", [{"item_id": "x"}])
    assert ex.run_files_shown(results_dir, "A") == ["hss_runA.jsonl", "real_runA.jsonl"]


def test_run_files_shown_empty_when_no_run_files():
    assert ex.run_files_shown(Path("no/such/dir"), "A") == []


def test_resolve_run_commit_falls_back_when_no_git(monkeypatch):
    def fake_run(*args, **kwargs):
        raise FileNotFoundError("no git")
    monkeypatch.setattr(ex.subprocess, "run", fake_run)
    assert ex.resolve_run_commit(Path("results/v3")) == "unknown"


def test_resolve_run_commit_appends_dirty_marker(monkeypatch):
    def fake_run(args, **kwargs):
        class Result:
            returncode = 0
            stdout = "deadbeef\n" if args[1] == "log" else " M results/v3/hss_runA.jsonl\n"
        return Result()
    monkeypatch.setattr(ex.subprocess, "run", fake_run)
    result = ex.resolve_run_commit(Path("results/v3"))
    assert result.startswith("deadbeef")
    assert "dirty" in result


def test_resolve_run_commit_no_dirty_marker_when_clean(monkeypatch):
    def fake_run(args, **kwargs):
        class Result:
            returncode = 0
            stdout = "deadbeef\n" if args[1] == "log" else ""
        return Result()
    monkeypatch.setattr(ex.subprocess, "run", fake_run)
    assert ex.resolve_run_commit(Path("results/v3")) == "deadbeef"


def test_resolve_run_commit_resolves_a_relative_results_dir_to_absolute(monkeypatch, tmp_path):
    # Regression: git subprocess calls run with cwd=REPO_ROOT, so a --results-dir given
    # relative to some other directory (e.g. "claims/results/v3" typed from evaluation/, while
    # REPO_ROOT is the repo root one level up) must be resolved to an absolute path first, or
    # git silently matches nothing under the wrong directory and this always reports "unknown".
    seen_paths = []

    def fake_run(args, **kwargs):
        seen_paths.append(args[-1])

        class Result:
            returncode = 0
            stdout = "deadbeef\n"
        return Result()

    monkeypatch.setattr(ex.subprocess, "run", fake_run)
    monkeypatch.setattr(ex, "_run_results_are_dirty", lambda results_dir: False)
    results_dir = tmp_path / "claims" / "results" / "v3"
    results_dir.mkdir(parents=True)
    monkeypatch.chdir(tmp_path)
    ex.resolve_run_commit(Path("claims/results/v3"))
    assert Path(seen_paths[0]).is_absolute()
    assert Path(seen_paths[0]) == results_dir.resolve()


def test_run_results_are_dirty_resolves_a_relative_results_dir_to_absolute(monkeypatch, tmp_path):
    seen_paths = []

    def fake_run(args, **kwargs):
        seen_paths.append(args[-1])

        class Result:
            returncode = 0
            stdout = ""
        return Result()

    monkeypatch.setattr(ex.subprocess, "run", fake_run)
    results_dir = tmp_path / "claims" / "results" / "v3"
    results_dir.mkdir(parents=True)
    monkeypatch.chdir(tmp_path)
    ex._run_results_are_dirty(Path("claims/results/v3"))
    assert Path(seen_paths[0]).is_absolute()
    assert Path(seen_paths[0]) == results_dir.resolve()


def _write_run_meta(
    results_dir: Path, name: str, *, model_reported, temperature, n_sessions,
):
    write_json(results_dir / name, {
        "model_reported": model_reported, "temperature": temperature,
        "sessions": [{"started": "t", "finished": "t"}] * n_sessions,
    })


def test_load_run_model_facts_reads_the_meta_files(tmp_path: Path):
    results_dir = tmp_path / "results"
    _write_run_meta(
        results_dir, "hss_runA.meta.json",
        model_reported=["deepseek-v4-flash"], temperature=0.0, n_sessions=1,
    )
    _write_run_meta(
        results_dir, "real_runA.meta.json",
        model_reported=["deepseek-v4-flash"], temperature=0.0, n_sessions=1,
    )
    assert ex.load_run_model_facts(results_dir, "A") == (
        "deepseek-v4-flash, temperature 0.0, 1 session per run"
    )


def test_load_run_model_facts_pluralises_multiple_sessions(tmp_path: Path):
    results_dir = tmp_path / "results"
    _write_run_meta(
        results_dir, "hss_runA.meta.json",
        model_reported=["deepseek-v4-flash"], temperature=0.0, n_sessions=3,
    )
    assert ex.load_run_model_facts(results_dir, "A") == (
        "deepseek-v4-flash, temperature 0.0, 3 sessions per run"
    )


def test_load_run_model_facts_unknown_when_no_meta_files(tmp_path: Path):
    assert ex.load_run_model_facts(tmp_path / "results", "A") == "unknown"


def test_load_run_model_facts_names_disagreement_instead_of_averaging(tmp_path: Path):
    results_dir = tmp_path / "results"
    _write_run_meta(
        results_dir, "hss_runA.meta.json",
        model_reported=["deepseek-v4-flash"], temperature=0.0, n_sessions=1,
    )
    _write_run_meta(
        results_dir, "real_runA.meta.json",
        model_reported=["a-different-model"], temperature=0.5, n_sessions=1,
    )
    result = ex.load_run_model_facts(results_dir, "A")
    assert "inconsistent" in result
    assert "deepseek-v4-flash" in result and "a-different-model" in result


def test_annotation_label_facts_hashes_each_file(tmp_path: Path):
    path_a = tmp_path / "hss_annotation_adjudicated_v3.csv"
    path_a.write_bytes(b"item_id,final_label\n")
    result = ex.annotation_label_facts([path_a])
    digest = hashlib.sha256(b"item_id,final_label\n").hexdigest()
    assert result == f"hss_annotation_adjudicated_v3.csv (sha256:{digest})"


def test_annotation_label_facts_joins_multiple_files(tmp_path: Path):
    path_a = tmp_path / "a.csv"
    path_a.write_bytes(b"a")
    path_b = tmp_path / "b.csv"
    path_b.write_bytes(b"b")
    result = ex.annotation_label_facts([path_a, path_b])
    assert "a.csv (sha256:" in result
    assert "b.csv (sha256:" in result
    assert "; " in result


def test_annotation_label_facts_empty_is_not_applicable():
    assert ex.annotation_label_facts([]) == "n/a"


# --------------------------------------------------------------------------------------
# Guide sheet
# --------------------------------------------------------------------------------------


def test_guide_paragraphs_opens_with_the_required_sentence():
    assert ex.guide_paragraphs()[0] == ex.GUIDE_OPENING_SENTENCE


def test_guide_paragraphs_carries_the_four_status_definitions_verbatim():
    paragraphs = ex.guide_paragraphs()
    for definition in ex.STATUS_DEFINITIONS_V3:
        assert definition in paragraphs


def test_guide_paragraphs_no_full_text_definition_has_the_plausibility_rule():
    no_full_text_def = next(d for d in ex.STATUS_DEFINITIONS_V3 if d.startswith("no_full_text"))
    assert "regardless of how plausible the claim sounds" in no_full_text_def


def test_guide_paragraphs_states_sources_sheet_caveat_for_withheld_rows():
    paragraphs = ex.guide_paragraphs()
    assert any(ex.NO_FULL_TEXT_EVIDENCE_MARKER in p for p in paragraphs)
    assert any("no_full_text" in p and "Sources sheet" in p for p in paragraphs)


# The Guide's quote_is_verbatim wording must match the tool's own rule: the title is part of the
# shown text, and the comparison folds case, whitespace and dash/quote style, not exact
# substring containment against abstract_shown alone.


def test_guide_paragraphs_quote_is_verbatim_allows_a_title_quote():
    paragraphs = ex.guide_paragraphs(has_screening=True)
    quote_para = next(p for p in paragraphs if "quote_is_verbatim" in p)
    assert "title" in quote_para
    assert "allowed" in quote_para


def test_guide_paragraphs_quote_is_verbatim_states_the_folding_rule():
    paragraphs = ex.guide_paragraphs(has_screening=True)
    quote_para = next(p for p in paragraphs if "quote_is_verbatim" in p)
    assert "case" in quote_para
    assert "whitespace" in quote_para
    assert "dash" in quote_para or "quote" in quote_para


# A "(no abstract)" row is a decision to confirm, not a record to judge on absent evidence.
# This applies beyond needs_review rows: an included row can also carry the same placeholder.


def test_guide_paragraphs_explains_no_abstract_rows():
    paragraphs = ex.guide_paragraphs(has_screening=True)
    assert any(
        "(no abstract)" in p and "title alone" in p for p in paragraphs
    )


def test_guide_paragraphs_no_abstract_note_covers_included_rows_too():
    paragraphs = ex.guide_paragraphs(has_screening=True)
    no_abstract_para = next(p for p in paragraphs if "(no abstract)" in p)
    assert "included" in no_abstract_para


def test_guide_paragraphs_no_abstract_note_absent_without_screening():
    paragraphs = ex.guide_paragraphs(has_screening=False)
    assert not any("(no abstract)" in p for p in paragraphs)


# The no-abstract note must not read as forbidding a legitimate disagreement with a
# title-only included row.


def test_guide_paragraphs_no_abstract_note_allows_disagreeing_with_a_title_only_include():
    paragraphs = ex.guide_paragraphs(has_screening=True)
    no_abstract_para = next(p for p in paragraphs if "(no abstract)" in p)
    assert "legitimate disagreement" in no_abstract_para
    assert "not enough to support an include" in no_abstract_para


def test_guide_paragraphs_no_abstract_note_still_restricts_exclude_and_needs_review():
    paragraphs = ex.guide_paragraphs(has_screening=True)
    no_abstract_para = next(p for p in paragraphs if "(no abstract)" in p)
    assert "not on its own a reason to mark a row exclude" in no_abstract_para
    assert "needs_review routing" in no_abstract_para


# The Guide must explain unanchored_exclude the same way it already explains no_abstract:
# parallel treatment for the other needs_review_reason value.


def test_guide_paragraphs_explains_unanchored_exclude_rows():
    paragraphs = ex.guide_paragraphs(has_screening=True)
    para = next(p for p in paragraphs if "unanchored_exclude" in p)
    assert "not anchored" in para
    assert "human_note" in para


def test_guide_paragraphs_unanchored_exclude_note_absent_without_screening():
    paragraphs = ex.guide_paragraphs(has_screening=False)
    assert not any("unanchored_exclude" in p for p in paragraphs)


def test_guide_paragraphs_defines_the_quote_relocated_slug():
    paragraphs = ex.guide_paragraphs()
    para = next(p for p in paragraphs if "quote_relocated" in p)
    assert "quote_relocated_original" in para
    assert "evidence_quote" in para


def test_guide_paragraphs_defines_the_second_pass_note():
    """One Guide sentence explains first_pass_demotion and the bounded second pass."""
    paragraphs = ex.guide_paragraphs()
    para = next(p for p in paragraphs if "first_pass_demotion" in p)
    assert "verified twice" in para or "second" in para.lower()


# SECOND_PASS_NOTE describes the follow-up as a RepairRequest naming the non-verbatim quote
# segments and asking for a corrected span or a changed status, not the same question asked
# twice; one of the quotes shown may be the follow-up's own corrected set with the failing
# quote withdrawn.


def test_second_pass_note_describes_the_repair_turn_not_an_identical_question():
    para = ex.SECOND_PASS_NOTE
    assert "the same question a second time" not in para
    assert "follow-up" in para.lower()
    assert "not verbatim" in para or "not match" in para
    assert "corrected" in para.lower()
    assert "withdraw" in para.lower() or "drop" in para.lower()


# UNANCHORED_EXCLUDE_NOTE tells the reviewer to answer quote_is_verbatim and
# criterion_is_right on their own merits, and explains that quote_is_verbatim is pulled into
# its own nested rate (score_review_sheet.py's quote_is_verbatim_rate["unanchored_exclude"])
# while criterion_is_right stays in the headline criterion_is_right_rate.


def test_unanchored_exclude_note_explains_the_nested_quote_rate():
    para = ex.UNANCHORED_EXCLUDE_NOTE
    assert "quote_is_verbatim_rate" in para
    assert "criterion_is_right_rate" in para
    assert "nested" in para.lower() or "own rate" in para.lower()


# BLINDNESS_FACT must account for the quote-relocation policy and the quote-repair turn, both
# applied with full knowledge of the annotation labels and both able to change a row's status
# on this same sheet.


def test_blindness_fact_names_the_two_mechanisms_added_after_annotation():
    fact = ex.BLINDNESS_FACT
    assert "not blind" in fact
    assert "relocation" in fact.lower()
    assert "repair" in fact.lower()


# One Guide sentence states the Review sheet is a census (every constructed and every real
# item), not a sample of either.


def test_guide_paragraphs_states_review_sheet_is_a_census():
    paragraphs = ex.guide_paragraphs(n_constructed=60, n_real=26)
    para = next(p for p in paragraphs if "60" in p and "26" in p)
    assert "census" in para
    assert "not a sample" in para or "not a subset" in para


def test_guide_paragraphs_census_sentence_absent_when_counts_not_given():
    paragraphs = ex.guide_paragraphs()
    assert not any("census, not a sample" in p for p in paragraphs)


# The Guide's full-text census paragraph (FULL_TEXT_CENSUS_STRUCTURAL_TEXT) must not describe
# excluded_full_text_violation as a live invariant when the record's own criteria stage no
# criterion full-text, the case every run has carried so far.


def test_guide_paragraphs_census_note_describes_full_text_stratum_as_structural_by_default():
    paragraphs = ex.guide_paragraphs(has_screening=True)
    census_para = next(p for p in paragraphs if "excluded_full_text_violation" in p)
    assert "cannot occur at all" in census_para
    assert "structural fact" in census_para
    # only two of the four are described as guard invariants
    assert "excluded_missing_anchor" in census_para
    assert "excluded_no_abstract" in census_para


def test_guide_paragraphs_census_note_describes_full_text_stratum_as_live_when_possible():
    paragraphs = ex.guide_paragraphs(has_screening=True, full_text_stratum_possible=True)
    census_para = next(
        p for p in paragraphs if "guard invariant expected to be" in p
    )
    assert "cannot occur at all" not in census_para


# unscreened is a deliberate operational outcome (a cancelled job, an exhausted time budget, a
# failed screening batch call), not a guard invariant whose non-empty state is by itself a bug
# report.


def test_guide_paragraphs_census_note_does_not_call_unscreened_a_guard_invariant():
    paragraphs = ex.guide_paragraphs(has_screening=True)
    census_para = next(p for p in paragraphs if "excluded_full_text_violation" in p)
    assert "unscreened" in census_para
    assert "and unscreened -- are guard invariants" not in census_para
    assert "its own case, not a guard invariant" in census_para


def test_guide_paragraphs_census_note_restores_the_fill_in_instruction():
    # The structural variant must carry the same fill-in instruction as the live variant
    # (used when a full-text-staged criterion exists).
    paragraphs = ex.guide_paragraphs(has_screening=True)
    census_para = next(p for p in paragraphs if "excluded_full_text_violation" in p)
    assert "Fill in the same human columns for these rows as for any other" in census_para


# The Guide explains the "(no abstract)" placeholder but never the "[...]" cut marker
# render_abstract_for_screening puts on a very long abstract.


def test_guide_paragraphs_explains_the_abstract_cut_marker():
    paragraphs = ex.guide_paragraphs(has_screening=True)
    para = next(p for p in paragraphs if "[...]" in p)
    assert "cut" in para
    assert "screener" in para


def test_guide_paragraphs_abstract_cut_note_absent_without_screening():
    paragraphs = ex.guide_paragraphs(has_screening=False)
    assert not any("[...]" in p for p in paragraphs)


# The Guide's UNANCHORED_EXCLUDE_NOTE must tell the reviewer that the dictated
# unanchored_exclude answer does not count as a disagreement in score_review_sheet.py's own
# per_status/overall numbers.


def test_guide_paragraphs_unanchored_exclude_note_explains_the_scoring_treatment():
    paragraphs = ex.guide_paragraphs(has_screening=True)
    para = next(p for p in paragraphs if "unanchored_exclude" in p)
    assert "per_status" in para
    assert "overall" in para


# Wrapping must never merge text from opposite sides of an explicit newline (e.g. two
# different numbered criteria) onto the same output line.


def test_wrap_paragraph_never_merges_across_a_newline():
    text = "I4. " + ("word " * 40).strip() + "\nE1. next criterion entirely"
    lines = ex._wrap_paragraph(text, 80)
    assert not any("I4." in line and "E1." in line for line in lines)
    assert any(line.startswith("E1.") for line in lines)


def test_wrap_paragraph_still_wraps_a_single_long_line():
    text = "word " * 400
    lines = ex._wrap_paragraph(text.strip(), 100)
    assert len(lines) > 1
    assert all(len(line) <= 100 for line in lines)


def test_guide_paragraphs_criteria_block_keeps_each_criterion_on_its_own_lines():
    long_i4 = "Reports at least one outcome concerning writing accuracy, revision or " * 3
    criteria_text = f"I1. short one\nI4. {long_i4}\nE1. an exclusion criterion"
    paragraphs = ex.guide_paragraphs(criteria_text=criteria_text, has_screening=True)
    assert not any("I4." in p and "E1." in p for p in paragraphs)


def test_write_guide_sheet_sanitizes_control_characters():
    # sanitize_for_excel applies to Guide as well as Review/Model_labels/Sources; a form feed
    # in a screening record's own free text would otherwise crash the export with an
    # unhandled IllegalCharacterError.
    wb = openpyxl.Workbook()
    ex.write_guide_sheet(
        wb.active, workbook_id="wbid", research_question="bad\x0cquestion",
        criteria_text="I1. also bad\x0c", has_screening=True, screening_seed=1,
    )  # no raise
    for row in wb.active.iter_rows():
        for cell in row:
            if isinstance(cell.value, str):
                assert "\x0c" not in cell.value


# The workbook id row must also carry the demo run id and export time, so a returned copy
# identifies its own export without the reviewer needing the key file.


def test_write_guide_sheet_prints_demo_run_id_and_export_time_next_to_workbook_id():
    wb = openpyxl.Workbook()
    ex.write_guide_sheet(
        wb.active, workbook_id="wbid123", demo_run_id="20260910-095226",
        export_time="2026-09-10T12:00:00+00:00",
    )
    values = [c.value for row in wb.active.iter_rows() for c in row if isinstance(c.value, str)]
    id_line = next(v for v in values if v.startswith(ex.WORKBOOK_ID_PREFIX))
    assert "20260910-095226" in id_line
    assert "2026-09-10T12:00:00+00:00" in id_line


def test_write_guide_sheet_defaults_when_demo_run_id_and_export_time_not_given():
    wb = openpyxl.Workbook()
    ex.write_guide_sheet(wb.active, workbook_id="wbid123")  # no raise
    values = [c.value for row in wb.active.iter_rows() for c in row if isinstance(c.value, str)]
    id_line = next(v for v in values if v.startswith(ex.WORKBOOK_ID_PREFIX))
    assert "n/a" in id_line


# census_counts feeds the Guide sheet's own census sentence.


def test_census_counts_splits_constructed_and_real():
    items = [
        {"item_id": "hss-verbatim-01", "rule": "verbatim"},
        {"item_id": "hss-altered-01", "rule": "altered"},
        {"item_id": "real-test-01", "rule": "real"},
    ]
    assert ex.census_counts(items) == (2, 1)


def test_census_counts_empty_items():
    assert ex.census_counts([]) == (0, 0)


def test_write_provenance_sheet_sanitizes_control_characters():
    wb = openpyxl.Workbook()
    ex.write_provenance_sheet(wb, [{"field": "results_directory", "value": "bad\x0cpath"}])
    ws = wb["Provenance"]
    assert "\x0c" not in ws.cell(row=2, column=2).value


# --------------------------------------------------------------------------------------
# Screening sampling
# --------------------------------------------------------------------------------------


def test_check_screening_prompt_version_raises_on_mismatch():
    record = {"provenance": {"screener": {"prompt_version": "sha256:old"}}}
    with pytest.raises(ex.ExportError, match="prompt_version"):
        ex.check_screening_prompt_version(record, "sha256:new")


def test_check_screening_prompt_version_passes_on_match():
    record = {"provenance": {"screener": {"prompt_version": "sha256:new"}}}
    ex.check_screening_prompt_version(record, "sha256:new")  # no raise


def test_allocate_stratified_with_floor_respects_total():
    sizes = {"a": 100, "b": 100, "c": 100}
    alloc = ex.allocate_stratified_with_floor(sizes, total_n=30, floor=5, min_bucket=5)
    assert sum(alloc.values()) == 30
    assert all(v >= 5 for v in alloc.values())


def test_allocate_stratified_with_floor_small_bucket_takes_all():
    sizes = {"big": 100, "tiny": 3}
    alloc = ex.allocate_stratified_with_floor(sizes, total_n=20, floor=5, min_bucket=5)
    assert alloc["tiny"] == 3
    assert alloc["big"] <= 100


def test_allocate_stratified_with_floor_never_exceeds_total_available():
    sizes = {"a": 2, "b": 3}
    alloc = ex.allocate_stratified_with_floor(sizes, total_n=40, floor=5, min_bucket=5)
    assert alloc == {"a": 2, "b": 3}


def test_allocate_stratified_with_floor_proportional_allocation():
    sizes = {"a": 90, "b": 10}
    alloc = ex.allocate_stratified_with_floor(sizes, total_n=20, floor=5, min_bucket=5)
    assert sum(alloc.values()) == 20
    assert alloc["a"] > alloc["b"]


def test_allocate_stratified_with_floor_empty_buckets():
    assert ex.allocate_stratified_with_floor({}, total_n=10, floor=5, min_bucket=5) == {}


def test_allocate_stratified_with_floor_never_exceeds_total_n_when_floor_infeasible():
    # 12 buckets * floor 5 = 60 exceeds total_n 40; allocation must not exceed total_n.
    sizes = {f"b{i}": 6 for i in range(12)}
    alloc = ex.allocate_stratified_with_floor(sizes, total_n=40, floor=5, min_bucket=5)
    assert sum(alloc.values()) == 40
    assert all(0 <= v <= 6 for v in alloc.values())


# The five measured cases below each assert the total (min(total_n, sum of sizes)) and that
# the largest bucket is never left at zero.


def test_allocate_stratified_with_floor_case_12_of_6_n_40():
    sizes = {f"b{i}": 6 for i in range(12)}
    alloc = ex.allocate_stratified_with_floor(sizes, total_n=40, floor=5, min_bucket=5)
    assert sum(alloc.values()) == 40
    assert max(alloc.values()) > 0


def test_allocate_stratified_with_floor_case_11_of_4_n_40():
    sizes = {f"b{i}": 4 for i in range(11)}
    alloc = ex.allocate_stratified_with_floor(sizes, total_n=40, floor=5, min_bucket=5)
    assert sum(alloc.values()) == 40
    assert max(alloc.values()) > 0


def test_allocate_stratified_with_floor_case_7_of_4_n_20():
    sizes = {f"b{i}": 4 for i in range(7)}
    alloc = ex.allocate_stratified_with_floor(sizes, total_n=20, floor=5, min_bucket=5)
    assert sum(alloc.values()) == 20
    assert max(alloc.values()) > 0


def test_allocate_stratified_with_floor_case_100_plus_11_of_4_n_40():
    # This guards against the largest bucket being zeroed while the smaller buckets keep
    # their whole size.
    sizes = {"big": 100, **{f"b{i}": 4 for i in range(11)}}
    alloc = ex.allocate_stratified_with_floor(sizes, total_n=40, floor=5, min_bucket=5)
    assert sum(alloc.values()) == 40
    assert alloc["big"] > 0
    assert max(alloc.values()) > 0


def test_allocate_stratified_with_floor_case_200_plus_6_of_4_n_20():
    sizes = {"big": 200, **{f"b{i}": 4 for i in range(6)}}
    alloc = ex.allocate_stratified_with_floor(sizes, total_n=20, floor=5, min_bucket=5)
    assert sum(alloc.values()) == 20
    assert alloc["big"] > 0
    assert max(alloc.values()) > 0


def test_allocate_stratified_with_floor_below_one_only_when_total_n_smaller_than_bucket_count():
    # total_n (2) is smaller than the number of non-empty buckets (3): giving every bucket at
    # least 1 is not possible, so the reduction floor relaxes to 0 rather than raising the
    # total above total_n.
    sizes = {"a": 10, "b": 10, "c": 10}
    alloc = ex.allocate_stratified_with_floor(sizes, total_n=2, floor=5, min_bucket=5)
    assert sum(alloc.values()) == 2
    assert all(v >= 0 for v in alloc.values())


def _records(n, outcome, **extra):
    out = []
    for i in range(n):
        row = {"outcome": outcome, "title": f"paper {outcome} {i}", "doi": f"10.1/{outcome}-{i}"}
        row.update({k: (v[i % len(v)] if isinstance(v, list) else v) for k, v in extra.items()})
        out.append(row)
    return out


def test_full_text_criterion_ids_reads_both_lists():
    criteria = {
        "inclusion_criteria_stages": ["abstract", "full_text"],
        "exclusion_criteria_stages": ["full_text", "abstract", "full_text"],
    }
    assert ex.full_text_criterion_ids(criteria) == {"I2", "E1", "E3"}


def test_full_text_criterion_ids_empty_for_v1_shaped_criteria():
    assert ex.full_text_criterion_ids({}) == set()


def test_classify_excluded_llm_census_no_abstract_wins_over_other_invariants():
    record = {"abstract": "", "quote": "", "criterion": ""}
    assert ex.classify_excluded_llm_census(record, set()) == "excluded_no_abstract"


def test_classify_excluded_llm_census_missing_anchor_over_full_text():
    record = {"abstract": "abs", "quote": "", "criterion": "E1"}
    assert ex.classify_excluded_llm_census(record, {"E1"}) == "excluded_missing_anchor"


def test_classify_excluded_llm_census_empty_criterion_is_missing_anchor():
    record = {"abstract": "abs", "quote": "some quote", "criterion": ""}
    assert ex.classify_excluded_llm_census(record, set()) == "excluded_missing_anchor"


def test_classify_excluded_llm_census_full_text_violation():
    record = {"abstract": "abs", "quote": "q", "criterion": "E1"}
    assert ex.classify_excluded_llm_census(record, {"E1"}) == "excluded_full_text_violation"


def test_classify_excluded_llm_census_none_for_normal_exclusion():
    record = {"abstract": "abs", "quote": "q", "criterion": "E1"}
    assert ex.classify_excluded_llm_census(record, set()) is None


def test_build_screening_sample_composition():
    records = (
        _records(50, "included", abstract="abs")
        + _records(30, "needs_review", needs_review_reason=["no_abstract", "undecidable"])
        + _records(
            40, "excluded", stage="llm", abstract="abs", quote="q",
            criterion=["E1", "E2", "TOPIC"],
        )
        + _records(5, "unscreened")
    )
    sampled, counts = ex.build_screening_sample(
        records, seed=1, included_n=10, needs_review_n=10, excluded_numbered_n=10,
        excluded_off_topic_n=5, floor=3, min_bucket=3,
    )
    strata = [r["_stratum"] for r in sampled]
    assert strata.count("included") == 10
    assert strata.count("needs_review") == 10
    assert strata.count("excluded_numbered") == 10
    assert strata.count("excluded_off_topic") == 5
    # unscreened is a whole census, not a sample -- every one of the 5 planted is present.
    assert strata.count("unscreened") == 5
    assert counts["included"] == {"drawn": 10, "available": 50}
    assert counts["needs_review"]["drawn"] == 10
    assert counts["unscreened"] == {"drawn": 5, "available": 5}


def test_build_screening_sample_stage_wos_excluded_never_sampled():
    records = (
        _records(3, "excluded", stage="wos")
        + _records(5, "excluded", stage="llm", abstract="abs", quote="q", criterion="E1")
    )
    sampled, counts = ex.build_screening_sample(
        records, seed=1, included_n=0, needs_review_n=0, excluded_numbered_n=10,
        excluded_off_topic_n=10, floor=1, min_bucket=1,
    )
    assert len(sampled) == 5
    assert all(r.get("stage") == "llm" for r in sampled)
    assert counts["excluded_numbered"] == {"drawn": 5, "available": 5}


def test_build_screening_sample_excluded_numbered_is_uniform_not_stratified():
    # A criterion-stratified draw with a floor forces a heavily-skewed population toward an
    # even split (5 of each here); a plain uniform draw does not.
    records = (
        _records(95, "excluded", stage="llm", abstract="abs", quote="q", criterion="BIG")
        + _records(5, "excluded", stage="llm", abstract="abs", quote="q", criterion="SMALL")
    )
    sampled, counts = ex.build_screening_sample(
        records, seed=1, included_n=0, needs_review_n=0, excluded_numbered_n=10,
        excluded_off_topic_n=0, floor=5, min_bucket=5,
    )
    numbered = [r for r in sampled if r["_stratum"] == "excluded_numbered"]
    criteria = [r["criterion"] for r in numbered]
    assert len(numbered) == 10
    assert criteria.count("SMALL") <= 2
    assert counts["excluded_numbered"] == {"drawn": 10, "available": 100}


def test_build_screening_sample_off_topic_kept_separate_from_numbered():
    records = (
        _records(3, "excluded", stage="llm", abstract="abs", quote="q", criterion="TOPIC")
        + _records(3, "excluded", stage="llm", abstract="abs", quote="q", criterion="E1")
    )
    sampled, counts = ex.build_screening_sample(
        records, seed=1, included_n=0, needs_review_n=0, excluded_numbered_n=10,
        excluded_off_topic_n=10, floor=1, min_bucket=1,
    )
    numbered = {r["criterion"] for r in sampled if r["_stratum"] == "excluded_numbered"}
    off_topic = {r["criterion"] for r in sampled if r["_stratum"] == "excluded_off_topic"}
    assert numbered == {"E1"}
    assert off_topic == {"TOPIC"}
    assert counts["excluded_numbered"] == {"drawn": 3, "available": 3}
    assert counts["excluded_off_topic"] == {"drawn": 3, "available": 3}


def test_build_screening_sample_census_no_abstract():
    records = _records(3, "excluded", stage="llm", abstract="", quote="q", criterion="E1")
    sampled, counts = ex.build_screening_sample(
        records, seed=1, included_n=0, needs_review_n=0, excluded_numbered_n=10,
        excluded_off_topic_n=10, floor=1, min_bucket=1,
    )
    assert [r["_stratum"] for r in sampled] == ["excluded_no_abstract"] * 3
    assert counts["excluded_no_abstract"] == {"drawn": 3, "available": 3}
    assert counts["excluded_numbered"] == {"drawn": 0, "available": 0}


def test_build_screening_sample_census_missing_anchor():
    records = (
        _records(2, "excluded", stage="llm", abstract="abs", quote="", criterion="E1")
        + _records(2, "excluded", stage="llm", abstract="abs", quote="q", criterion="")
    )
    sampled, counts = ex.build_screening_sample(
        records, seed=1, included_n=0, needs_review_n=0, excluded_numbered_n=10,
        excluded_off_topic_n=10, floor=1, min_bucket=1,
    )
    assert [r["_stratum"] for r in sampled] == ["excluded_missing_anchor"] * 4
    assert counts["excluded_missing_anchor"] == {"drawn": 4, "available": 4}


def test_build_screening_sample_census_full_text_violation():
    records = _records(2, "excluded", stage="llm", abstract="abs", quote="q", criterion="E1")
    criteria = {"exclusion_criteria_stages": ["full_text"]}
    sampled, counts = ex.build_screening_sample(
        records, seed=1, included_n=0, needs_review_n=0, excluded_numbered_n=10,
        excluded_off_topic_n=10, floor=1, min_bucket=1, criteria=criteria,
    )
    assert [r["_stratum"] for r in sampled] == ["excluded_full_text_violation"] * 2
    assert counts["excluded_full_text_violation"] == {"drawn": 2, "available": 2}
    assert counts["excluded_numbered"] == {"drawn": 0, "available": 0}


def test_build_screening_sample_no_abstract_takes_precedence_over_other_census():
    records = _records(1, "excluded", stage="llm", abstract="", quote="", criterion="E1")
    criteria = {"exclusion_criteria_stages": ["full_text"]}
    sampled, counts = ex.build_screening_sample(
        records, seed=1, included_n=0, needs_review_n=0, excluded_numbered_n=10,
        excluded_off_topic_n=10, floor=1, min_bucket=1, criteria=criteria,
    )
    assert [r["_stratum"] for r in sampled] == ["excluded_no_abstract"]


def test_build_screening_sample_unscreened_is_a_census_not_dropped():
    records = _records(4, "unscreened")
    sampled, counts = ex.build_screening_sample(
        records, seed=1, included_n=0, needs_review_n=0, excluded_numbered_n=0,
        excluded_off_topic_n=0, floor=1, min_bucket=1,
    )
    assert len(sampled) == 4
    assert all(r["_stratum"] == "unscreened" for r in sampled)
    assert counts["unscreened"] == {"drawn": 4, "available": 4}


def test_shuffle_records_deterministic():
    records = [{"i": i} for i in range(20)]
    a = ex.shuffle_records(records, seed=5)
    b = ex.shuffle_records(records, seed=5)
    assert a == b
    assert [r["i"] for r in a] != list(range(20))  # actually shuffled


def test_match_screening_key_prefers_doi():
    assert ex.match_screening_key({"doi": "10.1/X", "title": "t"}) == "doi:10.1/x"


def test_match_screening_key_falls_back_to_title():
    assert ex.match_screening_key({"title": "Some Title"}) == "title:some title"


def test_screening_run_b_status_blank_when_same():
    record = {"doi": "10.1/x", "status": "INCLUDE"}
    by_key = {"doi:10.1/x": {"status": "INCLUDE"}}
    assert ex.screening_run_b_status(record, by_key) == ""


def test_screening_run_b_status_filled_when_different():
    record = {"doi": "10.1/x", "status": "INCLUDE"}
    by_key = {"doi:10.1/x": {"status": "EXCLUDE"}}
    assert ex.screening_run_b_status(record, by_key) == "EXCLUDE"


def test_screening_run_b_status_blank_when_not_matched():
    record = {"doi": "10.1/x", "status": "INCLUDE"}
    assert ex.screening_run_b_status(record, {}) == ""


def test_screening_status_value_prefers_status_field():
    assert ex.screening_status_value({"status": "INCLUDE", "outcome": "included"}) == "INCLUDE"


def test_screening_status_value_folds_v1_outcome_to_present_tense():
    assert ex.screening_status_value({"outcome": "included"}) == "include"
    assert ex.screening_status_value({"outcome": "excluded"}) == "exclude"
    assert ex.screening_status_value({"outcome": "needs_review"}) == "needs_review"


def test_screening_status_value_unknown_outcome_passes_through():
    assert ex.screening_status_value({"outcome": "unscreened"}) == "unscreened"


def test_build_screening_row_shape_and_identifying_columns():
    record = {
        "outcome": "excluded", "stage": "llm", "status": "EXCLUDE", "criterion": "E1",
        "quote": "the verbatim span", "needs_review_reason": "", "to_confirm": "",
        "reason": "the model's sentence", "title": "Paper X", "doi": "10.1/x", "year": 2021,
        "journal": "J", "journal_issn": "1234-5678", "openalex_id": "W1",
        "abstract": "the abstract text", "_stratum": "excluded_numbered",
    }
    row = ex.build_screening_row(record, "S-001", max_cell_chars=1000, records_b_by_key=None)
    assert row["record_id"] == "S-001"
    assert row["outcome"] == "excluded"
    assert row["stage"] == "llm"
    assert row["reason"] == "the model's sentence"
    assert row["doi"] == "10.1/x"
    assert row["year"] == 2021
    assert row["journal"] == "J"
    assert row["journal_issn"] == "1234-5678"
    assert row["openalex_id"] == "W1"
    assert row["abstract_shown"] == "the abstract text"
    assert row["stratum"] == "excluded_numbered"
    for human_field in (
        "human_status", "human_agrees", "quote_is_verbatim", "criterion_is_right",
        "human_note",
    ):
        assert row[human_field] == ""


def test_build_screening_row_missing_run_b_key_absent():
    row = ex.build_screening_row(
        {"outcome": "included", "title": "t"}, "S-001", max_cell_chars=1000,
        records_b_by_key=None,
    )
    assert "run_B_status" not in row


# --------------------------------------------------------------------------------------
# is_no_abstract_exempt / the abstract gate exemption covers all four evidence strata: see the
# function's own docstring for why (the screener can return an included decision from the
# title alone with no abstract at all).
# --------------------------------------------------------------------------------------


def test_is_no_abstract_exempt_true_for_a_needs_review_record_with_no_abstract():
    record = {"_stratum": "needs_review", "abstract": ""}
    assert ex.is_no_abstract_exempt(record) is True


def test_is_no_abstract_exempt_true_for_an_included_record_with_no_abstract():
    record = {"_stratum": "included", "abstract": ""}
    assert ex.is_no_abstract_exempt(record) is True


def test_is_no_abstract_exempt_true_for_excluded_numbered_and_off_topic_with_no_abstract():
    for stratum in ("excluded_numbered", "excluded_off_topic"):
        record = {"_stratum": stratum, "abstract": ""}
        assert ex.is_no_abstract_exempt(record) is True


def test_is_no_abstract_exempt_false_for_a_stratum_outside_the_four_evidence_strata():
    for stratum in (
        "excluded_no_abstract", "excluded_missing_anchor", "excluded_full_text_violation",
        "unscreened",
    ):
        record = {"_stratum": stratum, "abstract": ""}
        assert ex.is_no_abstract_exempt(record) is False


def test_is_no_abstract_exempt_false_when_abstract_present():
    for stratum in ex.EVIDENCE_STRATA:
        record = {"_stratum": stratum, "abstract": "text"}
        assert ex.is_no_abstract_exempt(record) is False


def test_build_screening_row_no_abstract_needs_review_gets_placeholder():
    record = {
        "outcome": "needs_review", "needs_review_reason": "no_abstract", "abstract": "",
        "title": "Paper Y", "_stratum": "needs_review",
    }
    row = ex.build_screening_row(record, "S-010", max_cell_chars=1000, records_b_by_key=None)
    assert row["abstract_shown"] == ex.NO_ABSTRACT_PLACEHOLDER
    assert row["_no_abstract_exempt"] is True


def test_build_screening_row_normal_needs_review_not_exempt():
    record = {
        "outcome": "needs_review", "needs_review_reason": "undecidable", "abstract": "the text",
        "_stratum": "needs_review",
    }
    row = ex.build_screening_row(record, "S-011", max_cell_chars=1000, records_b_by_key=None)
    assert row["abstract_shown"] == "the text"
    assert row["_no_abstract_exempt"] is False


def test_build_screening_row_included_with_no_abstract_gets_placeholder():
    # An included row can legitimately carry no abstract at all, decided from the title alone.
    record = {
        "outcome": "included", "abstract": "", "title": "Paper Z", "_stratum": "included",
    }
    row = ex.build_screening_row(record, "S-012", max_cell_chars=1000, records_b_by_key=None)
    assert row["abstract_shown"] == ex.NO_ABSTRACT_PLACEHOLDER
    assert row["_no_abstract_exempt"] is True


def test_build_screening_row_included_with_real_abstract_not_exempt():
    record = {
        "outcome": "included", "abstract": "the real abstract", "title": "Paper W",
        "_stratum": "included",
    }
    row = ex.build_screening_row(record, "S-013", max_cell_chars=1000, records_b_by_key=None)
    assert row["abstract_shown"] == "the real abstract"
    assert row["_no_abstract_exempt"] is False


# --------------------------------------------------------------------------------------
# render_abstract_for_screening / abstract_shown cut: the sheet must show the human the
# screener's own head/marker/tail cut, not the whole stored abstract, once the abstract runs
# past the screener's own limit.
# --------------------------------------------------------------------------------------


def test_render_abstract_for_screening_short_abstract_unchanged():
    assert ex.render_abstract_for_screening("a short one") == "a short one"


def test_render_abstract_for_screening_exactly_at_limit_unchanged():
    text = "x" * ex.SCREENER_ABSTRACT_CHAR_LIMIT
    assert ex.render_abstract_for_screening(text) == text


def test_render_abstract_for_screening_over_limit_cuts_head_marker_tail():
    text = "H" * 8000 + "T" * 8000  # 16,000 chars, over the 10,000 limit
    rendered = ex.render_abstract_for_screening(text)
    assert ex.SCREENER_ABSTRACT_CUT_MARKER in rendered
    head, _, tail = rendered.partition(ex.SCREENER_ABSTRACT_CUT_MARKER)
    assert head == "H" * 7000  # 7:3 split of the 10,000 limit
    assert tail == "T" * 3000
    assert text.startswith(head)
    assert text.endswith(tail)


def test_render_abstract_for_screening_matches_the_documented_10000_character_limit():
    # backend/app/agents/relevance_screener_agent.py::ABSTRACT_CHAR_LIMIT is 10,000 (this
    # module makes no backend import -- see the module docstring -- so the value is checked
    # by its documented effect instead: an abstract one character over 10,000 is cut, one
    # character at 10,000 is not).
    assert ex.SCREENER_ABSTRACT_CHAR_LIMIT == 10_000
    at_limit = "x" * 10_000
    over_limit = "x" * 10_001
    assert ex.SCREENER_ABSTRACT_CUT_MARKER not in ex.render_abstract_for_screening(at_limit)
    assert ex.SCREENER_ABSTRACT_CUT_MARKER in ex.render_abstract_for_screening(over_limit)


def test_render_abstract_for_screening_none_abstract_is_empty_string():
    assert ex.render_abstract_for_screening(None) == ""


# build_shown_texts strips the abstract before ever calling render_abstract, so the mirror
# here must strip too, otherwise a stored abstract with leading/trailing whitespace can fall
# on the wrong side of the limit, or land on the right side but with every head/tail offset
# shifted by the whitespace count.


def test_render_abstract_for_screening_strips_whitespace_before_measuring_length():
    text = "x" * ex.SCREENER_ABSTRACT_CHAR_LIMIT
    padded = "   " + text + "\n\t"
    rendered = ex.render_abstract_for_screening(padded)
    assert rendered == text
    assert ex.SCREENER_ABSTRACT_CUT_MARKER not in rendered


def test_render_abstract_for_screening_strips_whitespace_around_content_already_over_the_limit():
    # inner is over the limit on its own; padding it with whitespace must not change where
    # the head/tail cut lands, since strip() removes the padding first.
    inner = "x" * (ex.SCREENER_ABSTRACT_CHAR_LIMIT + 500)
    padded = f"  {inner}  "
    rendered_padded = ex.render_abstract_for_screening(padded)
    rendered_plain = ex.render_abstract_for_screening(inner)
    assert rendered_padded == rendered_plain
    assert ex.SCREENER_ABSTRACT_CUT_MARKER in rendered_padded


def test_build_screening_row_cuts_a_long_abstract_through_the_screener_rule():
    record = {
        "outcome": "included", "abstract": "A" * 20_000, "title": "Long abstract paper",
        "_stratum": "included",
    }
    row = ex.build_screening_row(record, "S-020", max_cell_chars=32_000, records_b_by_key=None)
    assert ex.SCREENER_ABSTRACT_CUT_MARKER in row["abstract_shown"]
    assert len(row["abstract_shown"]) < 20_000


def test_build_screening_row_short_abstract_not_cut():
    record = {
        "outcome": "included", "abstract": "a short abstract", "title": "Paper",
        "_stratum": "included",
    }
    row = ex.build_screening_row(record, "S-021", max_cell_chars=32_000, records_b_by_key=None)
    assert row["abstract_shown"] == "a short abstract"


# --------------------------------------------------------------------------------------
# Workbook integration
# --------------------------------------------------------------------------------------


@pytest.fixture()
def small_export(tmp_path: Path):
    items_path = tmp_path / "items.jsonl"
    write_jsonl(items_path, [CONSTRUCTED_ITEM, REAL_ITEM, NO_FULL_TEXT_ITEM])
    cache_dir = tmp_path / "cache"
    make_cache(cache_dir, "10.1000/paper-a", "Paper A", [
        "Introduction chunk not relevant here at all.",
        "Scores were significantly higher for the treatment group after training.",
    ])
    make_cache(cache_dir, "10.1000/paper-b", "Paper B", [
        "Feedback improved learner motivation across the term for most participants.",
    ])
    results_dir = tmp_path / "results"
    write_jsonl(results_dir / "hss_runA.jsonl", [
        result_row(
            "hss-verbatim-01", "verified",
            evidence_quotes=["Scores were significantly higher for the treatment group "
                             "after training."],
        ),
        result_row("hss-no_full_text-01", "no_full_text"),
    ])
    write_jsonl(results_dir / "real_runA.jsonl", [
        result_row("real-test-01", "needs_nuance"),
    ])
    write_json(results_dir / "hss_runA.meta.json", {"prompt_version": "sha256:testprompt"})
    return items_path, cache_dir, results_dir


def test_build_workbook_writes_expected_sheets(small_export, tmp_path: Path):
    items_path, cache_dir, results_dir = small_export
    item_specs = ex.resolve_item_specs([items_path], [cache_dir])
    wb, kept_items, opaque_ids, missing, workbook_id = ex.build_workbook(
        item_specs, results_dir, ["A"], max_cell_chars=32000, sources_max_chars=30000,
    )
    assert missing == []
    assert len(kept_items) == 3
    assert set(wb.sheetnames) >= {"Guide", "Review", "Sources", "Provenance"}
    review_ws = wb["Review"]
    header = [c.value for c in review_ws[1]]
    assert header == ex.review_header()
    assert review_ws.max_row == 4  # header + 3 items


# build_workbook must surface the repair turn's own two new meta fields on the Provenance
# sheet when a run's own meta.json carries them, and fall back to "unknown" (the same default
# every other optional provenance fact uses) when it does not (v3, v4, and this fixture's own
# default meta.json all lack them).


def test_build_workbook_provenance_reads_policy_and_repair_prompt_version_from_meta(
    small_export,
):
    items_path, cache_dir, results_dir = small_export
    write_json(results_dir / "hss_runA.meta.json", {
        "prompt_version": "sha256:testprompt",
        "verification_policy_version": "667bcc9209126ddc",
        "repair_prompt_version": "sha256:7ec7cffe852d",
    })
    item_specs = ex.resolve_item_specs([items_path], [cache_dir])
    wb, *_ = ex.build_workbook(
        item_specs, results_dir, ["A"], max_cell_chars=32000, sources_max_chars=30000,
    )
    prov = {row[0].value: row[1].value for row in wb["Provenance"].iter_rows(min_row=2)}
    assert prov["verification_policy_version"] == "667bcc9209126ddc"
    assert prov["repair_prompt_version"] == "sha256:7ec7cffe852d"


def test_build_workbook_provenance_defaults_policy_and_repair_prompt_version_when_absent(
    small_export,
):
    items_path, cache_dir, results_dir = small_export
    item_specs = ex.resolve_item_specs([items_path], [cache_dir])
    wb, *_ = ex.build_workbook(
        item_specs, results_dir, ["A"], max_cell_chars=32000, sources_max_chars=30000,
    )
    prov = {row[0].value: row[1].value for row in wb["Provenance"].iter_rows(min_row=2)}
    assert prov["verification_policy_version"] == "unknown"
    assert prov["repair_prompt_version"] == "unknown"


# The Guide sheet's own workbook-id row and the Provenance sheet must print the same
# export_time and demo_run_id for one export.
def test_build_workbook_guide_and_provenance_share_export_time_and_demo_run_id(small_export):
    items_path, cache_dir, results_dir = small_export
    item_specs = ex.resolve_item_specs([items_path], [cache_dir])
    wb, kept_items, opaque_ids, missing, workbook_id = ex.build_workbook(
        item_specs, results_dir, ["A"], max_cell_chars=32000, sources_max_chars=30000,
        demo_run_id="20260910-095226", export_time="2026-09-10T12:00:00+00:00",
    )
    prov = {row[0].value: row[1].value for row in wb["Provenance"].iter_rows(min_row=2)}
    assert prov["export_time"] == "2026-09-10T12:00:00+00:00"
    assert prov["demo_run_id"] == "20260910-095226"
    guide_values = [
        c.value for row in wb["Guide"].iter_rows() for c in row if isinstance(c.value, str)
    ]
    id_line = next(v for v in guide_values if v.startswith(ex.WORKBOOK_ID_PREFIX))
    assert "2026-09-10T12:00:00+00:00" in id_line
    assert "20260910-095226" in id_line


# build_workbook's own Guide sheet must carry the census sentence, computed from the kept
# items, without any caller having to pass it explicitly.
def test_build_workbook_guide_sheet_states_the_census_counts(small_export):
    items_path, cache_dir, results_dir = small_export
    item_specs = ex.resolve_item_specs([items_path], [cache_dir])
    wb, kept_items, opaque_ids, missing, workbook_id = ex.build_workbook(
        item_specs, results_dir, ["A"], max_cell_chars=32000, sources_max_chars=30000,
    )
    n_constructed, n_real = ex.census_counts(kept_items)
    guide_values = [
        c.value for row in wb["Guide"].iter_rows() for c in row if isinstance(c.value, str)
    ]
    assert any(
        "census" in v and str(n_constructed) in v and str(n_real) in v for v in guide_values
    )


def test_build_workbook_sources_sheet_keyed_by_item_id(small_export):
    # Sources is keyed by the citing item's own opaque item_id, not the cited paper's DOI, so
    # a withheld (no_full_text) row has no rows of its own there under any key.
    items_path, cache_dir, results_dir = small_export
    item_specs = ex.resolve_item_specs([items_path], [cache_dir])
    wb, kept_items, opaque_ids, missing, workbook_id = ex.build_workbook(
        item_specs, results_dir, ["A"], max_cell_chars=32000, sources_max_chars=30000,
    )
    sources_ws = wb["Sources"]
    header = [c.value for c in sources_ws[1]]
    assert header == ["item_id", "part", "text"]
    item_ids_on_sheet = {row[0].value for row in sources_ws.iter_rows(min_row=2)}
    assert opaque_ids["hss-no_full_text-01"] not in item_ids_on_sheet
    assert opaque_ids["hss-verbatim-01"] in item_ids_on_sheet


def test_build_workbook_missing_primary_result_raises(small_export):
    items_path, cache_dir, results_dir = small_export
    # add an item with no run-A result
    extra_item = {**CONSTRUCTED_ITEM, "item_id": "hss-verbatim-99"}
    items = [json.loads(line) for line in items_path.read_text(encoding="utf-8").splitlines()]
    items.append(extra_item)
    write_jsonl(items_path, items)
    item_specs = ex.resolve_item_specs([items_path], [cache_dir])
    with pytest.raises(ex.ExportError, match="no run A result"):
        ex.build_workbook(item_specs, results_dir, ["A"], max_cell_chars=32000,
                           sources_max_chars=30000)


def test_build_workbook_allow_missing_results_drops_item(small_export):
    items_path, cache_dir, results_dir = small_export
    extra_item = {**CONSTRUCTED_ITEM, "item_id": "hss-verbatim-99"}
    items = [json.loads(line) for line in items_path.read_text(encoding="utf-8").splitlines()]
    items.append(extra_item)
    write_jsonl(items_path, items)
    item_specs = ex.resolve_item_specs([items_path], [cache_dir])
    wb, kept_items, opaque_ids, missing, workbook_id = ex.build_workbook(
        item_specs, results_dir, ["A"], max_cell_chars=32000, sources_max_chars=30000,
        allow_missing_results=True,
    )
    assert missing == ["hss-verbatim-99"]
    assert all(it["item_id"] != "hss-verbatim-99" for it in kept_items)


def test_build_workbook_with_annotation_adds_model_labels_sheet(small_export, tmp_path: Path):
    items_path, cache_dir, results_dir = small_export
    ann_path = tmp_path / "ann.csv"
    with ann_path.open("w", encoding="utf-8", newline="") as fh:
        writer = csv.DictWriter(fh, fieldnames=["item_id", "final_label", "source", "rationale"])
        writer.writeheader()
        writer.writerow({
            "item_id": "hss-verbatim-01", "final_label": "verified", "source": "unanimous",
            "rationale": "matches",
        })
    item_specs = ex.resolve_item_specs([items_path], [cache_dir])
    wb, *_ = ex.build_workbook(
        item_specs, results_dir, ["A"], max_cell_chars=32000, sources_max_chars=30000,
        annotation_paths=[ann_path],
    )
    assert "Model_labels" in wb.sheetnames
    ml_header = [c.value for c in wb["Model_labels"][1]]
    assert ml_header == ex.model_labels_header()


def test_key_csv_round_trip(small_export, tmp_path: Path):
    items_path, cache_dir, results_dir = small_export
    item_specs = ex.resolve_item_specs([items_path], [cache_dir])
    wb, kept_items, opaque_ids, missing, workbook_id = ex.build_workbook(
        item_specs, results_dir, ["A"], max_cell_chars=32000, sources_max_chars=30000,
    )
    key_path = tmp_path / "key.csv"
    ex.write_key_csv(key_path, kept_items, opaque_ids, workbook_id=workbook_id)
    with key_path.open(encoding="utf-8") as fh:
        rows = list(csv.DictReader(fh))
    assert len(rows) == 3
    assert all(r["workbook_id"] == workbook_id for r in rows)
    real_item_ids = {r["real_item_id"] for r in rows}
    assert real_item_ids == {"hss-verbatim-01", "real-test-01", "hss-no_full_text-01"}


def test_build_screening_sheet_end_to_end(tmp_path: Path):
    record = {
        "criteria": {
            "query": "does X help Y?",
            "inclusion_criteria": ["studies X"], "exclusion_criteria": ["no abstract"],
        },
        "provenance": {"screener": {"prompt_version": "sha256:demov1"}},
        "records": (
            [{"outcome": "included", "title": f"inc {i}", "doi": f"10.1/inc-{i}",
              "abstract": f"abstract text {i}"} for i in range(5)]
            + [{"outcome": "excluded", "title": f"exc {i}", "doi": f"10.1/exc-{i}",
                "stage": "llm", "criterion": "E1", "quote": "the quote", "reason": "why",
                "abstract": f"abstract text exc {i}"} for i in range(5)]
        ),
    }
    record_path = tmp_path / "screening_record.json"
    write_json(record_path, record)
    wb = openpyxl.Workbook()
    ex.write_guide_sheet(wb.active, workbook_id="wbid")
    loaded_record, rows, stratum_counts = ex.build_screening_sheet(
        wb, record_path, expected_prompt_version="sha256:demov1", seed=1,
        included_n=40, needs_review_n=40, excluded_numbered_n=40, excluded_off_topic_n=20,
        floor=5, min_bucket=5, max_cell_chars=32000,
    )
    assert "Screening" in wb.sheetnames
    assert len(rows) == 10
    header = [c.value for c in wb["Screening"][1]]
    assert header == ex.screening_header(has_run_b=False)
    # B1: the sheet's evidence column is actually populated, not structurally blank.
    evidence_rows = [r for r in rows if r["stratum"] in ex.EVIDENCE_STRATA]
    assert evidence_rows
    assert all(r["abstract_shown"].strip() for r in evidence_rows)
    assert stratum_counts["included"] == {"drawn": 5, "available": 5}
    assert stratum_counts["excluded_numbered"] == {"drawn": 5, "available": 5}


def test_build_screening_sheet_refuses_wrong_prompt_version(tmp_path: Path):
    record = {
        "criteria": {}, "provenance": {"screener": {"prompt_version": "sha256:old"}},
        "records": [],
    }
    record_path = tmp_path / "screening_record.json"
    write_json(record_path, record)
    wb = openpyxl.Workbook()
    with pytest.raises(ex.ExportError, match="prompt_version"):
        ex.build_screening_sheet(
            wb, record_path, expected_prompt_version="sha256:new", seed=1, included_n=1,
            needs_review_n=1, excluded_numbered_n=1, excluded_off_topic_n=1, floor=1,
            min_bucket=1, max_cell_chars=1000,
        )


def test_build_screening_sheet_sampled_no_abstract_needs_review_row_exports(tmp_path: Path):
    # A needs_review row the screener routed there for having no abstract at all must export
    # with the placeholder, not refuse the whole sheet.
    record = {
        "criteria": {}, "provenance": {"screener": {"prompt_version": "sha256:demov1"}},
        "records": [{
            "outcome": "needs_review", "needs_review_reason": "no_abstract", "abstract": "",
            "title": "no abstract paper", "doi": "10.1/na-0",
        }],
    }
    record_path = tmp_path / "screening_record.json"
    write_json(record_path, record)
    wb = openpyxl.Workbook()
    loaded_record, rows, stratum_counts = ex.build_screening_sheet(
        wb, record_path, expected_prompt_version="sha256:demov1", seed=1, included_n=1,
        needs_review_n=1, excluded_numbered_n=1, excluded_off_topic_n=1, floor=1,
        min_bucket=1, max_cell_chars=1000,
    )  # no raise
    assert len(rows) == 1
    assert rows[0]["abstract_shown"] == ex.NO_ABSTRACT_PLACEHOLDER


def test_build_screening_sheet_included_row_with_no_abstract_exports(tmp_path: Path):
    # An included row can legitimately carry no abstract at all, decided from the title alone.
    # Not a build defect; must export with the placeholder, not refuse the whole sheet. The
    # exemption covers all four evidence strata, not just needs_review.
    record = {
        "criteria": {}, "provenance": {"screener": {"prompt_version": "sha256:demov1"}},
        "records": [
            {"outcome": "included", "title": "inc 0", "doi": "10.1/inc-0", "abstract": ""},
        ],
    }
    record_path = tmp_path / "screening_record.json"
    write_json(record_path, record)
    wb = openpyxl.Workbook()
    loaded_record, rows, stratum_counts = ex.build_screening_sheet(
        wb, record_path, expected_prompt_version="sha256:demov1", seed=1, included_n=1,
        needs_review_n=1, excluded_numbered_n=1, excluded_off_topic_n=1, floor=1,
        min_bucket=1, max_cell_chars=1000,
    )  # no raise
    assert len(rows) == 1
    assert rows[0]["abstract_shown"] == ex.NO_ABSTRACT_PLACEHOLDER


# --------------------------------------------------------------------------------------
# check_screening_abstracts
# --------------------------------------------------------------------------------------


def test_check_screening_abstracts_passes_when_evidence_rows_have_text():
    rows = [{"record_id": "S-001", "stratum": "included", "abstract_shown": "text"}]
    ex.check_screening_abstracts(rows)  # no raise


def test_check_screening_abstracts_raises_on_blank_evidence_row():
    rows = [{"record_id": "S-001", "stratum": "included", "abstract_shown": ""}]
    with pytest.raises(ex.ExportError, match="abstract"):
        ex.check_screening_abstracts(rows)


def test_check_screening_abstracts_ignores_no_abstract_census_row():
    rows = [{"record_id": "S-001", "stratum": "excluded_no_abstract", "abstract_shown": ""}]
    ex.check_screening_abstracts(rows)  # no raise -- that stratum is for exactly this


def test_check_screening_abstracts_ignores_unscreened_row():
    rows = [{"record_id": "S-001", "stratum": "unscreened", "abstract_shown": ""}]
    ex.check_screening_abstracts(rows)  # no raise -- never shown one at all


def test_check_screening_abstracts_exempts_flagged_needs_review_row():
    # build_screening_row's own _no_abstract_exempt flag, not just a non-blank
    # abstract_shown, is what this checks.
    rows = [{
        "record_id": "S-001", "stratum": "needs_review",
        "abstract_shown": ex.NO_ABSTRACT_PLACEHOLDER, "_no_abstract_exempt": True,
    }]
    ex.check_screening_abstracts(rows)  # no raise


def test_check_screening_abstracts_exempts_flagged_included_row():
    # The exemption flag covers included/excluded_numbered/excluded_off_topic too, not just
    # needs_review.
    rows = [{
        "record_id": "S-001", "stratum": "included",
        "abstract_shown": ex.NO_ABSTRACT_PLACEHOLDER, "_no_abstract_exempt": True,
    }]
    ex.check_screening_abstracts(rows)  # no raise


def test_check_screening_abstracts_still_raises_on_unflagged_needs_review_row():
    rows = [{
        "record_id": "S-001", "stratum": "needs_review", "abstract_shown": "",
        "_no_abstract_exempt": False,
    }]
    with pytest.raises(ex.ExportError, match="abstract"):
        ex.check_screening_abstracts(rows)


# --------------------------------------------------------------------------------------
# Screening provenance / vocabulary
# --------------------------------------------------------------------------------------


def test_append_screening_stratum_rows_writes_one_row_per_stratum():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Provenance"
    ws.append(["field", "value"])
    counts = {s: {"drawn": 1, "available": 2} for s in ex.SCREENING_STRATA}
    ex.append_screening_stratum_rows(ws, counts)
    fields = [row[0].value for row in ws.iter_rows(min_row=2)]
    assert fields == [f"screening_stratum_{s}" for s in ex.SCREENING_STRATA]
    values = [row[1].value for row in ws.iter_rows(min_row=2)]
    assert all(v == "1 of 2" for v in values)


# --------------------------------------------------------------------------------------
# retrieval_cutoff: the record's own criteria.publication_date_max, the retrieval cap
# forwarded to every OpenAlex call on the run, dropped nowhere on the sheet.
# --------------------------------------------------------------------------------------


def test_retrieval_cutoff_value_present():
    criteria = {"publication_date_max": "2026-09-08", "inclusion_criteria": []}
    assert ex.retrieval_cutoff_value(criteria) == "2026-09-08"


def test_retrieval_cutoff_value_none_when_absent():
    assert ex.retrieval_cutoff_value({"inclusion_criteria": []}) is None


def test_retrieval_cutoff_value_none_when_blank():
    assert ex.retrieval_cutoff_value({"publication_date_max": ""}) is None


def test_append_retrieval_cutoff_row_writes_field_value_pair():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Provenance"
    ws.append(["field", "value"])
    ex.append_retrieval_cutoff_row(ws, "2026-09-08")
    assert ws.cell(row=2, column=1).value == "retrieval_cutoff"
    assert ws.cell(row=2, column=2).value == "2026-09-08"


def test_write_screening_sheet_hides_stratum_column():
    wb = openpyxl.Workbook()
    rows = [ex.build_screening_row(
        {"outcome": "included", "title": "t", "_stratum": "included"}, "S-001",
        max_cell_chars=1000, records_b_by_key=None,
    )]
    ex.write_screening_sheet(wb, rows, has_run_b=False)
    ws = wb["Screening"]
    header = [c.value for c in ws[1]]
    col_letter = ex.get_column_letter(header.index("stratum") + 1)
    assert ws.column_dimensions[col_letter].hidden is True


def test_disagreement_axis_choices_includes_construct_naming():
    assert "construct naming" in ex.DISAGREEMENT_AXIS_CHOICES


def test_topic_criterion_id_is_topic():
    assert ex.TOPIC_CRITERION_ID == "TOPIC"


# --------------------------------------------------------------------------------------
# Delivered sheet: the final product a user receives, judged sentence by sentence
# --------------------------------------------------------------------------------------

DELIVERED_FIXTURE_PATH = (
    Path(__file__).resolve().parent / "fixtures" / "delivered_evidence_sample.json"
)


def test_delivered_header_shape():
    assert ex.delivered_header() == [
        "row", "section_title", "draft_id", "sentence", "claim_checked", "citation_text",
        "paper_title", "paper_authors", "paper_doi", "evidence_quotes", "source_passage",
        "source_located", "passage_located", "unlocated_quotes", "location_section_mismatch",
        "sentence_in_draft", "human_sentence_correct", "human_quote_supports", "human_note",
    ]


def test_load_delivered_evidence_reads_the_fixture():
    record = ex.load_delivered_evidence(DELIVERED_FIXTURE_PATH)
    assert record["run_id"] == "20260911-062145"
    assert record["draft_id"] == "ef8cd27d-0e97-4265-ac17-6e8de30f4b84"
    assert record["section_title"] == "Literature Review"
    assert len(record["rows"]) == 2


def test_load_delivered_evidence_raises_named_reason_when_missing(tmp_path: Path):
    missing = tmp_path / "no-such-file.json"
    with pytest.raises(ex.ExportError, match=re.escape(str(missing))):
        ex.load_delivered_evidence(missing)


def test_load_delivered_evidence_raises_when_not_valid_json(tmp_path: Path):
    bad = tmp_path / "bad.json"
    bad.write_text("not json", encoding="utf-8")
    with pytest.raises(ex.ExportError, match="not found or not valid json"):
        ex.load_delivered_evidence(bad)


def test_load_delivered_evidence_raises_when_rows_missing_or_malformed(tmp_path: Path):
    bad = tmp_path / "bad.json"
    write_json(bad, {"run_id": "r", "draft_id": "d", "section_title": "s"})
    with pytest.raises(ex.ExportError, match="rows"):
        ex.load_delivered_evidence(bad)


def test_build_delivered_row_shape():
    record = ex.load_delivered_evidence(DELIVERED_FIXTURE_PATH)
    row = ex.build_delivered_row(record["rows"][0], max_cell_chars=32000)
    assert row["row"] == 1
    assert row["sentence"] == record["rows"][0]["sentence"]
    assert row["citation_text"] == "(Fixture & Sample, 2026)"
    assert row["paper_title"] == record["rows"][0]["paper_title"]
    assert row["paper_doi"] == "10.1234/fixture.2026"
    assert row["evidence_quotes"] == "1) " + record["rows"][0]["evidence_quotes"][0]
    assert row["source_passage"] == record["rows"][0]["source_passage"]
    # Neither fixture row carries source_located/passage_located/sentence_in_draft (all
    # predate these fields on the record schema) -- build_delivered_row must not fabricate
    # a yes or no for a field the source record never recorded.
    assert row["source_located"] == ""
    assert row["passage_located"] == ""
    assert row["unlocated_quotes"] == ""
    assert row["location_section_mismatch"] == ""
    assert row["sentence_in_draft"] == ""
    assert row["human_sentence_correct"] == ""
    assert row["human_quote_supports"] == ""
    assert row["human_note"] == ""
    # The fixture row does carry claim_text; no --paper-authors map is given here, so
    # paper_authors stays blank rather than fabricated.
    assert row["claim_checked"] == record["rows"][0]["claim_text"]
    assert row["paper_authors"] == ""


# A run can write more than one generated section into one delivered_evidence.json; a row
# that carries its own section_title/draft_id is shown under that value, and a row that does
# not falls back to the file's own top-level value.
def test_build_delivered_row_defaults_section_title_and_draft_id_to_blank_with_no_fallback():
    out = ex.build_delivered_row({"row": 1}, max_cell_chars=32000)
    assert out["section_title"] == ""
    assert out["draft_id"] == ""


def test_build_delivered_row_uses_its_own_section_title_and_draft_id_when_present():
    row = {"row": 1, "section_title": "Discussion", "draft_id": "draft-2"}
    out = ex.build_delivered_row(
        row, max_cell_chars=32000,
        fallback_section_title="Introduction", fallback_draft_id="draft-1",
    )
    assert out["section_title"] == "Discussion"
    assert out["draft_id"] == "draft-2"


def test_build_delivered_row_falls_back_to_the_file_top_level_when_a_row_lacks_its_own():
    out = ex.build_delivered_row(
        {"row": 1}, max_cell_chars=32000,
        fallback_section_title="Introduction", fallback_draft_id="draft-1",
    )
    assert out["section_title"] == "Introduction"
    assert out["draft_id"] == "draft-1"


def test_build_delivered_sheet_uses_the_file_top_level_section_title_as_fallback():
    wb = openpyxl.Workbook()
    record, rows = ex.build_delivered_sheet(wb, DELIVERED_FIXTURE_PATH, max_cell_chars=32000)
    # DELIVERED_FIXTURE_PATH's own rows carry no section_title/draft_id of their own.
    assert all(r["section_title"] == "Literature Review" for r in rows)
    assert all(r["draft_id"] == "ef8cd27d-0e97-4265-ac17-6e8de30f4b84" for r in rows)


def test_build_delivered_sheet_prefers_a_rows_own_section_title_over_the_file_top_level(
    tmp_path: Path,
):
    path = tmp_path / "multi_section.json"
    write_json(path, {
        "run_id": "r", "draft_id": "d-top", "section_title": "Fallback Section",
        "rows": [
            {"row": 1, "sentence": "s1", "section_title": "Introduction", "draft_id": "d-1"},
            {"row": 2, "sentence": "s2"},
        ],
    })
    wb = openpyxl.Workbook()
    record, rows = ex.build_delivered_sheet(wb, path, max_cell_chars=32000)
    assert rows[0]["section_title"] == "Introduction"
    assert rows[0]["draft_id"] == "d-1"
    assert rows[1]["section_title"] == "Fallback Section"
    assert rows[1]["draft_id"] == "d-top"


def test_build_delivered_row_joins_multiple_evidence_quotes():
    row = {"row": 1, "evidence_quotes": ["one quote", "another quote"]}
    out = ex.build_delivered_row(row, max_cell_chars=32000)
    assert out["evidence_quotes"] == "1) one quote\n2) another quote"


# The record's own fidelity flags must reach the sheet.
def test_build_delivered_row_maps_source_located_true_and_false():
    assert ex.build_delivered_row(
        {"row": 1, "source_located": True}, max_cell_chars=32000
    )["source_located"] == "yes"
    assert ex.build_delivered_row(
        {"row": 1, "source_located": False}, max_cell_chars=32000
    )["source_located"] == "no"


# passage_located must reach the sheet just like source_located does, and
# unlocated_quotes/location_section_mismatch must be readable columns, not silently dropped.
def test_build_delivered_row_maps_passage_located_true_and_false():
    assert ex.build_delivered_row(
        {"row": 1, "passage_located": True}, max_cell_chars=32000
    )["passage_located"] == "yes"
    assert ex.build_delivered_row(
        {"row": 1, "passage_located": False}, max_cell_chars=32000
    )["passage_located"] == "no"


def test_build_delivered_row_joins_unlocated_quotes():
    row = {"row": 1, "unlocated_quotes": ["one quote", "another quote"]}
    out = ex.build_delivered_row(row, max_cell_chars=32000)
    assert out["unlocated_quotes"] == "1) one quote\n2) another quote"


def test_build_delivered_row_maps_location_section_mismatch_true_and_false():
    assert ex.build_delivered_row(
        {"row": 1, "location_section_mismatch": True}, max_cell_chars=32000
    )["location_section_mismatch"] == "yes"
    assert ex.build_delivered_row(
        {"row": 1, "location_section_mismatch": False}, max_cell_chars=32000
    )["location_section_mismatch"] == "no"


def test_build_delivered_row_maps_sentence_in_draft_true_and_false():
    assert ex.build_delivered_row(
        {"row": 1, "sentence_in_draft": True}, max_cell_chars=32000
    )["sentence_in_draft"] == "yes"
    assert ex.build_delivered_row(
        {"row": 1, "sentence_in_draft": False}, max_cell_chars=32000
    )["sentence_in_draft"] == "no"


def test_build_delivered_row_caps_long_text():
    row = {
        "row": 1, "sentence": "s" * 100, "source_passage": "p" * 100,
    }
    out = ex.build_delivered_row(row, max_cell_chars=20)
    assert len(out["sentence"]) == 20
    assert out["sentence"].endswith(ex.TRUNCATION_MARKER)
    assert len(out["source_passage"]) == 20


def test_build_delivered_sheet_adds_sheet_with_rows_and_validation():
    wb = openpyxl.Workbook()
    record, rows = ex.build_delivered_sheet(
        wb, DELIVERED_FIXTURE_PATH, max_cell_chars=32000,
    )
    assert record["run_id"] == "20260911-062145"
    assert len(rows) == 2
    assert "Delivered" in wb.sheetnames
    ws = wb["Delivered"]
    header = [c.value for c in ws[1]]
    assert header == ex.delivered_header()
    assert ws.max_row == 3  # header + 2 rows
    dv_sqrefs = " ".join(str(dv.sqref) for dv in ws.data_validations.dataValidation)
    correct_col = openpyxl.utils.get_column_letter(header.index("human_sentence_correct") + 1)
    supports_col = openpyxl.utils.get_column_letter(header.index("human_quote_supports") + 1)
    assert f"{correct_col}2" in dv_sqrefs
    assert f"{supports_col}2" in dv_sqrefs


def test_build_delivered_sheet_raises_named_reason_when_file_missing(tmp_path: Path):
    wb = openpyxl.Workbook()
    missing = tmp_path / "no-such-file.json"
    with pytest.raises(ex.ExportError, match=re.escape(str(missing))):
        ex.build_delivered_sheet(wb, missing, max_cell_chars=32000)
    assert "Delivered" not in wb.sheetnames


def test_append_delivered_provenance_rows():
    wb = openpyxl.Workbook()
    ex.write_provenance_sheet(wb, [{"field": "model", "value": "deepseek-flash"}])
    ex.append_delivered_provenance_rows(
        wb["Provenance"], n_rows=2, run_id="20260911-062145",
        draft_id="ef8cd27d-0e97-4265-ac17-6e8de30f4b84",
    )
    prov = {row[0].value: row[1].value for row in wb["Provenance"].iter_rows(min_row=2)}
    assert prov["delivered_rows"] == "2"
    assert prov["delivered_run_id"] == "20260911-062145"
    assert prov["delivered_draft_id"] == "ef8cd27d-0e97-4265-ac17-6e8de30f4b84"


def test_distinct_delivered_draft_ids_joins_every_distinct_id_in_first_seen_order():
    """A multi-section run's merged record carries no top-level draft_id at all; the
    Provenance sheet's own value must be built from the rows themselves instead."""
    rows = [
        {"row": 1, "draft_id": "draft-a"},
        {"row": 1, "draft_id": "draft-b"},
        {"row": 2, "draft_id": "draft-a"},
    ]
    assert ex.distinct_delivered_draft_ids(rows) == "draft-a; draft-b"


def test_distinct_delivered_draft_ids_of_a_single_section_run_is_just_that_one_id():
    rows = [{"row": 1, "draft_id": "draft-a"}, {"row": 2, "draft_id": "draft-a"}]
    assert ex.distinct_delivered_draft_ids(rows) == "draft-a"


def test_distinct_delivered_draft_ids_of_no_rows_is_blank():
    assert ex.distinct_delivered_draft_ids([]) == ""


def test_guide_paragraphs_includes_delivered_note_with_count_when_present():
    paragraphs = ex.guide_paragraphs(has_delivered=True, n_delivered=2)
    joined = " ".join(paragraphs)
    assert "Delivered" in joined
    assert "final product" in joined
    assert "2" in joined
    assert "human_sentence_correct" in joined
    assert "human_quote_supports" in joined
    # The Guide must say what source_located and sentence_in_draft are and that a row where
    # either is no is excluded from the headline rates.
    assert "source_located" in joined
    assert "sentence_in_draft" in joined
    assert "excluded" in joined
    # The Guide must describe passage_located as distinct from source_located
    # (source_located records only that the chunk was fetched, not that a passage was
    # found), and must not still claim source_located means the evidence passage could
    # not be found.
    assert "passage_located" in joined
    assert "unlocated_quotes" in joined
    assert "location_section_mismatch" in joined
    assert "source_located is no when the verifier's own evidence passage" not in joined


def test_guide_paragraphs_omits_delivered_note_by_default():
    paragraphs = ex.guide_paragraphs()
    joined = " ".join(paragraphs)
    assert "final product" not in joined


# DELIVERED_GUIDE_NOTE_TEMPLATE must not claim the sheet carries "every sentence of the
# generated report that actually reached the user". The row set is one row per
# final_report.verifications entry (demo/run_demo.py build_delivered_evidence_rows, which
# raises on any entry whose own status is not "verified"), so a delivered sentence with no
# citation, or a cited sentence whose verification did not reach "verified", gets no row.


def test_guide_paragraphs_delivered_note_does_not_overclaim_full_coverage():
    paragraphs = ex.guide_paragraphs(has_delivered=True, n_delivered=6)
    joined = " ".join(paragraphs)
    assert "every sentence of the generated report that actually reached the user" not in joined
    assert "cited" in joined
    assert "no citation" in joined or "not cited" in joined


# The rows come from several generated sections of one run; the Guide must say so, must say
# finalize is what removes an uncited finding and an unverified cited sentence, and must
# state how many sections and how many rows there are.
def test_guide_paragraphs_delivered_note_states_section_and_row_counts():
    paragraphs = ex.guide_paragraphs(has_delivered=True, n_delivered=35, n_delivered_sections=5)
    joined = " ".join(paragraphs)
    assert "5 generated sections" in joined
    assert "35 rows" in joined
    assert "one demo run" in joined
    assert "finalize removes" in joined


# DELIVERED_GUIDE_NOTE_TEMPLATE must not claim finalize "removes both kinds of sentence ...
# so only the cited and verified sentences of each section reach this sheet": a section's own
# framing-tagged uncited sentences are never removed by finalize
# (backend/app/services/fulltext.py, _finalize_paragraph_text) and are delivered in the saved
# draft without ever getting a row on this sheet. The note must name the two removal rules it
# actually applies and must not claim the Delivered sheet's rows are all a section's saved
# draft contains.
def test_guide_paragraphs_delivered_note_does_not_claim_finalize_removes_framing_sentences():
    paragraphs = ex.guide_paragraphs(has_delivered=True, n_delivered=35, n_delivered_sections=5)
    joined = " ".join(paragraphs)
    assert "removes both kinds of sentence" not in joined
    assert "only the cited and verified sentences of each section reach this sheet" not in joined
    assert "empirical finding" in joined
    assert "framing sentences are delivered" in joined


# `_drop_dangling_framing_sentences` (backend/app/services/fulltext.py) removes a third kind
# of sentence: an uncited framing sentence left dangling by one of the first two removals,
# incrementing the same finalize_stats["sentences_removed_dangling"] counter. The note must
# name all three removal rules and must not claim every framing sentence in a section
# survives.
def test_guide_paragraphs_delivered_note_names_dangling_framing_removal():
    paragraphs = ex.guide_paragraphs(has_delivered=True, n_delivered=35, n_delivered_sections=5)
    joined = " ".join(paragraphs)
    assert "a section's own framing sentences are delivered" not in joined
    assert "left dangling by one of those removals" in joined
    assert "section's remaining framing sentences are delivered" in joined


def test_guide_paragraphs_delivered_note_singular_section_and_row_count():
    paragraphs = ex.guide_paragraphs(has_delivered=True, n_delivered=1, n_delivered_sections=1)
    joined = " ".join(paragraphs)
    assert "1 generated section" in joined
    assert "1 row" in joined
    assert "1 generated sections" not in joined
    assert "1 rows" not in joined


def test_guide_paragraphs_delivered_note_defaults_section_count_when_not_given():
    # Unchanged behaviour for a caller that has not been updated to pass n_delivered_sections
    # (e.g. an older Guide-rewrite call site): the row count still renders correctly.
    paragraphs = ex.guide_paragraphs(has_delivered=True, n_delivered=6)
    joined = " ".join(paragraphs)
    assert "6 rows" in joined


# --------------------------------------------------------------------------------------
# Author page templates
# --------------------------------------------------------------------------------------


def test_render_review_singleton_note_takes_no_item_id_parameter():
    import inspect
    params = inspect.signature(ex.render_review_singleton_note).parameters
    assert not any("id" in name.lower() for name in params)


def test_render_review_singleton_note_never_mentions_an_item_id():
    text = ex.render_review_singleton_note("a bounded follow-up turn", 1)
    assert "hss-" not in text
    assert "real-" not in text
    assert "not identifiable" not in text
    assert "1 row" in text


def test_render_review_singleton_note_pluralises_the_count():
    assert "2 rows" in ex.render_review_singleton_note("a bounded follow-up turn", 2)


def test_render_delivered_author_paragraph_always_carries_the_qualifying_clause():
    text = ex.render_delivered_author_paragraph(n_rows=6, n_sections=1)
    assert "only the cited and verified part of it" in text
    assert "6 sentences" in text
    assert "1 generated section" in text


def test_render_delivered_author_paragraph_states_section_count_when_plural():
    text = ex.render_delivered_author_paragraph(n_rows=35, n_sections=5)
    assert "35 sentences" in text
    assert "5 generated sections" in text


def test_render_screening_author_paragraph_says_unchanged_when_nothing_changed():
    text = ex.render_screening_author_paragraph([])
    assert "Unchanged" in text


def test_render_screening_author_paragraph_names_every_change():
    text = ex.render_screening_author_paragraph([
        "the unanchored_exclude Guide paragraph gained a sentence",
        "a new stratum row",
    ])
    assert "the unanchored_exclude Guide paragraph gained a sentence" in text
    assert "a new stratum row" in text
    assert "Unchanged" not in text


def test_assert_author_page_safe_raises_on_a_real_item_id():
    with pytest.raises(ex.AuthorPageError):
        ex.assert_author_page_safe("one row in this run, hss-verbatim-07, carries it")


def test_assert_author_page_safe_raises_on_a_real_item_id_from_the_real_claims_set():
    with pytest.raises(ex.AuthorPageError):
        ex.assert_author_page_safe("this row is real-test-06 under its own real id")


def test_assert_author_page_safe_raises_on_the_false_assurance_phrase():
    with pytest.raises(ex.AuthorPageError):
        ex.assert_author_page_safe("this row is not identifiable from the opaque sheet")


def test_assert_author_page_safe_passes_clean_text():
    ex.assert_author_page_safe(
        ex.render_delivered_author_paragraph(n_rows=6, n_sections=1)
        + " " + ex.render_review_singleton_note("a bounded follow-up turn", 1)
        + " " + ex.render_screening_author_paragraph(["a new stratum row"])
    )


def test_assert_author_page_safe_does_not_flag_an_opaque_review_row_id():
    # Opaque ids (R-013, R-074, ...) are the point of the sheet and are not a leak on their
    # own: only a *real* item id is forbidden.
    ex.assert_author_page_safe("14 rows changed status: R-013, R-026, R-074.")


def test_assert_author_page_safe_raises_on_every_item_id_in_both_shipped_item_files():
    """`REAL_ITEM_ID_PATTERN`'s slug character class must include ``_``, or every
    ``hss-no_full_text-NN``, ``hss-over_specified-NN`` and ``hss-wrong_paper-NN`` id fails
    to match, even though those three categories name the most disclosive thing about a
    row's construction and expected label. Runs the guard over EVERY item_id of both
    shipped item files, not only ``hss-verbatim-07`` and ``real-test-06``."""
    claims_dir = Path(__file__).resolve().parents[1] / "claims"
    item_ids = [
        item["item_id"]
        for filename in ("hss_test_v3_claims.jsonl", "real_claims_test.jsonl")
        for item in read_jsonl(claims_dir / filename)
    ]
    assert len(item_ids) == 86  # 60 + 26: a guard on the fixture files themselves
    for item_id in item_ids:
        with pytest.raises(ex.AuthorPageError):
            ex.assert_author_page_safe(f"one row, {item_id}, carries it")


# --------------------------------------------------------------------------------------
# Parity with the annotation round
# --------------------------------------------------------------------------------------


def test_no_full_text_marker_matches_aggregate_annotations_v3_withheld_marker():
    agg3 = load_script_module("claims/annotation", "aggregate_annotations_v3")
    assert ex.NO_FULL_TEXT_EVIDENCE_MARKER == agg3.WITHHELD_MARKER


def _normalise_whitespace(text: str) -> str:
    return " ".join(text.split())


def _instructions_v3_labels_list() -> list[str]:
    """The "## Labels" section's own bullet list from INSTRUCTIONS_v3.md, one entry per
    bullet, wrapped continuation lines joined back into a single string and whitespace
    normalised -- the same normalisation applied to STATUS_DEFINITIONS_V3 before comparing."""
    path = ex.REPO_ROOT / "evaluation" / "claims" / "annotation" / "INSTRUCTIONS_v3.md"
    text = path.read_text(encoding="utf-8")
    section = text.split("## Labels", 1)[1].split("\n## ", 1)[0]
    chunks = re.split(r"\n(?=- )", section.strip())
    bullets = [c for c in chunks if c.startswith("- ")]
    return [_normalise_whitespace(b[2:]) for b in bullets]


def test_status_definitions_v3_matches_instructions_v3_labels_list():
    expected = _instructions_v3_labels_list()
    actual = [_normalise_whitespace(d) for d in ex.STATUS_DEFINITIONS_V3]
    assert actual == expected


# --------------------------------------------------------------------------------------
# v4: claim_checked, paper_authors, delivered-only mode
# --------------------------------------------------------------------------------------

PAPER_AUTHORS_FIXTURE_PATH = (
    Path(__file__).resolve().parent / "fixtures" / "paper_authors_sample.json"
)


def test_delivered_header_inserts_claim_checked_after_sentence():
    header = ex.delivered_header()
    assert header[header.index("sentence") + 1] == "claim_checked"


def test_delivered_header_inserts_paper_authors_after_paper_title():
    header = ex.delivered_header()
    assert header[header.index("paper_title") + 1] == "paper_authors"


def test_build_delivered_row_carries_claim_checked_from_the_record():
    out = ex.build_delivered_row({"row": 1, "claim_text": "the tool checked this"}, max_cell_chars=32000)
    assert out["claim_checked"] == "the tool checked this"


def test_build_delivered_row_claim_checked_blank_when_the_record_has_none():
    out = ex.build_delivered_row({"row": 1}, max_cell_chars=32000)
    assert out["claim_checked"] == ""


def test_build_delivered_row_caps_a_long_claim_checked():
    out = ex.build_delivered_row(
        {"row": 1, "claim_text": "c" * 100}, max_cell_chars=20,
    )
    assert len(out["claim_checked"]) == 20
    assert out["claim_checked"].endswith(ex.TRUNCATION_MARKER)


def test_build_delivered_row_leaves_paper_authors_blank_when_no_map_is_given():
    out = ex.build_delivered_row({"row": 1, "paper_doi": "10.1/x"}, max_cell_chars=32000)
    assert out["paper_authors"] == ""


def test_build_delivered_row_fills_paper_authors_from_the_map_by_doi():
    out = ex.build_delivered_row(
        {"row": 1, "paper_doi": "10.1234/fixture.2026"}, max_cell_chars=32000,
        paper_authors={"10.1234/fixture.2026": "Ada Fixture; Sam Sample (2026)"},
    )
    assert out["paper_authors"] == "Ada Fixture; Sam Sample (2026)"


def test_build_delivered_row_raises_named_reason_when_a_doi_has_no_author_record():
    with pytest.raises(ex.ExportError, match="no-such-doi"):
        ex.build_delivered_row(
            {"row": 7, "paper_doi": "no-such-doi"}, max_cell_chars=32000,
            paper_authors={"10.1234/fixture.2026": "Ada Fixture; Sam Sample (2026)"},
        )


def test_build_delivered_row_numbers_every_evidence_quote_on_its_own_line():
    out = ex.build_delivered_row(
        {"row": 1, "evidence_quotes": ["first", "second", "third"]}, max_cell_chars=32000,
    )
    assert out["evidence_quotes"] == "1) first\n2) second\n3) third"


def test_build_delivered_row_numbered_quotes_survive_a_semicolon_inside_a_quote():
    out = ex.build_delivered_row(
        {"row": 1, "evidence_quotes": ["one; two", "three"]}, max_cell_chars=32000,
    )
    lines = out["evidence_quotes"].split("\n")
    assert lines == ["1) one; two", "2) three"]


def test_build_delivered_sheet_with_no_author_map_still_writes_every_row():
    wb = openpyxl.Workbook()
    record, rows = ex.build_delivered_sheet(wb, DELIVERED_FIXTURE_PATH, max_cell_chars=32000)
    assert len(rows) == 2
    assert all(r["paper_authors"] == "" for r in rows)


def test_build_delivered_sheet_raises_named_reason_when_a_doi_has_no_author_record():
    wb = openpyxl.Workbook()
    with pytest.raises(ex.ExportError, match="10.1234/fixture.2026"):
        ex.build_delivered_sheet(
            wb, DELIVERED_FIXTURE_PATH, max_cell_chars=32000, paper_authors={},
        )


def test_build_delivered_sheet_fills_paper_authors_from_a_real_map():
    wb = openpyxl.Workbook()
    authors = ex.load_paper_authors(PAPER_AUTHORS_FIXTURE_PATH)
    record, rows = ex.build_delivered_sheet(
        wb, DELIVERED_FIXTURE_PATH, max_cell_chars=32000, paper_authors=authors,
    )
    assert all(r["paper_authors"] == "Ada Fixture; Sam Sample (2026)" for r in rows)


def test_assert_delivered_claims_present_raises_naming_the_rows_with_no_claim_text(
    tmp_path: Path,
):
    path = tmp_path / "missing_claims.json"
    write_json(path, {
        "run_id": "r", "draft_id": "d", "rows": [
            {"row": 1, "claim_text": "has one"},
            {"row": 2, "claim_text": ""},
            {"row": 3},
        ],
    })
    record = ex.load_delivered_evidence(path)
    with pytest.raises(ex.ExportError, match="2, 3"):
        ex.assert_delivered_claims_present(record, path)


def test_assert_delivered_claims_present_passes_on_the_demo_record():
    record = ex.load_delivered_evidence(DELIVERED_FIXTURE_PATH)
    ex.assert_delivered_claims_present(record, DELIVERED_FIXTURE_PATH)  # no raise


def test_load_paper_authors_indexes_selected_json_by_doi():
    authors = ex.load_paper_authors(PAPER_AUTHORS_FIXTURE_PATH)
    assert authors == {"10.1234/fixture.2026": "Ada Fixture; Sam Sample (2026)"}


def test_load_paper_authors_keeps_publisher_name_order_and_appends_the_year(tmp_path: Path):
    path = tmp_path / "authors.json"
    write_json(path, [
        {"doi": "10.1/a", "authors_crossref": ["Zed Zebra", "Alice Apple"], "year": 2020},
    ])
    authors = ex.load_paper_authors(path)
    assert authors["10.1/a"] == "Zed Zebra; Alice Apple (2020)"


def test_load_paper_authors_lookup_is_case_insensitive_on_the_doi(tmp_path: Path):
    path = tmp_path / "authors.json"
    write_json(path, [{"doi": "10.1/A-Mixed-Case", "authors_crossref": ["A. One"], "year": 2020}])
    authors = ex.load_paper_authors(path)
    assert authors["10.1/a-mixed-case"] == "A. One (2020)"


def test_load_paper_authors_formats_a_single_author_paper(tmp_path: Path):
    path = tmp_path / "authors.json"
    write_json(path, [{"doi": "10.1/solo", "authors_crossref": ["Ken Hyland"], "year": 2025}])
    authors = ex.load_paper_authors(path)
    assert authors["10.1/solo"] == "Ken Hyland (2025)"


def _delivered_only_rows():
    return [
        {"row": 1, "section_title": "S1", "draft_id": "d1", "sentence": "sent one",
         "claim_checked": "claim one", "paper_doi": "10.1/a"},
        {"row": 2, "section_title": "S1", "draft_id": "d1", "sentence": "sent two",
         "claim_checked": "claim two", "paper_doi": "10.1/b"},
    ]


def test_compute_delivered_workbook_id_is_deterministic_for_the_same_rows():
    rows = _delivered_only_rows()
    id_a = ex.compute_delivered_workbook_id("run-1", rows)
    id_b = ex.compute_delivered_workbook_id("run-1", rows)
    assert id_a == id_b
    assert len(id_a) == 16


def test_compute_delivered_workbook_id_changes_when_a_claim_changes():
    rows = _delivered_only_rows()
    id_a = ex.compute_delivered_workbook_id("run-1", rows)
    rows2 = [dict(r) for r in rows]
    rows2[0]["claim_checked"] = "a different claim"
    id_b = ex.compute_delivered_workbook_id("run-1", rows2)
    assert id_a != id_b


def test_compute_delivered_workbook_id_changes_when_the_demo_run_id_changes():
    rows = _delivered_only_rows()
    id_a = ex.compute_delivered_workbook_id("run-1", rows)
    id_b = ex.compute_delivered_workbook_id("run-2", rows)
    assert id_a != id_b


def test_compute_delivered_workbook_id_differs_from_the_item_list_workbook_id():
    rows = _delivered_only_rows()
    delivered_id = ex.compute_delivered_workbook_id("run-1", rows)
    item_list_id = ex.compute_workbook_id(["item-1", "item-2"], seed=42)
    assert delivered_id != item_list_id


def test_distinct_delivered_sentence_count_counts_groups_not_rows():
    rows = [
        {"draft_id": "d1", "sentence": "s1"},
        {"draft_id": "d1", "sentence": "s1"},
        {"draft_id": "d1", "sentence": "s2"},
    ]
    assert ex.distinct_delivered_sentence_count(rows) == 2


def test_distinct_delivered_sentence_count_separates_the_same_text_in_two_drafts():
    rows = [{"draft_id": "d1", "sentence": "same"}, {"draft_id": "d2", "sentence": "same"}]
    assert ex.distinct_delivered_sentence_count(rows) == 2


def test_longest_source_passage_chars_reports_the_longest_cell():
    rows = [{"source_passage": "short"}, {"source_passage": "a much longer passage here"}]
    assert ex.longest_source_passage_chars(rows) == len("a much longer passage here")


def test_longest_source_passage_chars_is_zero_for_no_rows():
    assert ex.longest_source_passage_chars([]) == 0


def test_count_containing_claim_pairs_finds_an_overlapping_pair_in_one_group():
    rows = [
        {"draft_id": "d1", "sentence": "s1", "claim_checked": "four ignored X"},
        {"draft_id": "d1", "sentence": "s1", "claim_checked": "Uptake selective, as four ignored X"},
    ]
    assert ex.count_containing_claim_pairs(rows) == 1


def test_count_containing_claim_pairs_is_zero_when_claims_are_disjoint():
    rows = [
        {"draft_id": "d1", "sentence": "s1", "claim_checked": "alpha"},
        {"draft_id": "d1", "sentence": "s1", "claim_checked": "beta"},
    ]
    assert ex.count_containing_claim_pairs(rows) == 0


def test_guide_paragraphs_delivered_only_returns_exactly_the_thirteen_specified_cells():
    cells = ex.guide_paragraphs(
        delivered_only=True, n_delivered=68, n_delivered_sections=7,
        n_delivered_sentences=50, longest_passage_chars=18040,
    )
    assert len(cells) == 13
    assert cells[0].startswith("The Delivered sheet carries the final product")
    assert cells[-1].startswith("human_note is required")


def test_guide_paragraphs_delivered_only_writes_no_review_or_screening_paragraph():
    cells = ex.guide_paragraphs(delivered_only=True, n_delivered=1, n_delivered_sections=1)
    joined = " ".join(cells)
    assert "Review" not in joined
    assert "Screening" not in joined
    assert "Model_labels" not in joined
    assert "countersign" not in joined


def test_guide_paragraphs_delivered_only_ends_with_the_note_rule():
    cells = ex.guide_paragraphs(delivered_only=True, n_delivered=1, n_delivered_sections=1)
    assert "human_note is required whenever" in cells[-1]


def test_guide_paragraphs_delivered_note_states_row_sentence_and_section_counts():
    cells = ex.guide_paragraphs(
        delivered_only=True, n_delivered=68, n_delivered_sections=7, n_delivered_sentences=50,
    )
    joined = " ".join(cells)
    assert "68 rows" in joined
    assert "7 generated sections" in joined
    assert "50 distinct sentences" in joined


def test_guide_paragraphs_delivered_note_states_the_longest_passage_length():
    cells = ex.guide_paragraphs(
        delivered_only=True, n_delivered=1, n_delivered_sections=1, longest_passage_chars=18040,
    )
    joined = " ".join(cells)
    assert "18,040" in joined


def test_guide_paragraphs_delivered_note_names_claim_checked_and_paper_authors():
    cells = ex.guide_paragraphs(delivered_only=True, n_delivered=1, n_delivered_sections=1)
    joined = " ".join(cells)
    assert "claim_checked" in joined
    assert "paper_authors" in joined


def test_guide_paragraphs_delivered_note_intro_cell_does_not_define_the_two_questions():
    cells = ex.guide_paragraphs(delivered_only=True, n_delivered=1, n_delivered_sections=1)
    assert "human_sentence_correct" not in cells[0]
    assert "human_quote_supports" not in cells[0]


def test_guide_paragraphs_delivered_note_allows_two_rows_of_one_sentence_to_differ():
    cells = ex.guide_paragraphs(delivered_only=True, n_delivered=1, n_delivered_sections=1)
    assert "do not assume two such rows must get the same answer" in cells[0]


def test_guide_paragraphs_delivered_note_tells_the_reader_to_resolve_a_back_reference():
    cells = ex.guide_paragraphs(delivered_only=True, n_delivered=1, n_delivered_sections=1)
    joined = " ".join(cells)
    assert "back reference" in joined


def test_guide_paragraphs_delivered_note_requires_a_note_on_every_no_row():
    cells = ex.guide_paragraphs(delivered_only=True, n_delivered=1, n_delivered_sections=1)
    assert "human_note is required whenever human_sentence_correct or human_quote_supports is no" in cells[-1]


def test_guide_paragraphs_delivered_note_has_no_em_dash_or_spaced_double_hyphen():
    cells = ex.guide_paragraphs(
        delivered_only=True, n_delivered=68, n_delivered_sections=7, n_delivered_sentences=50,
        longest_passage_chars=18040,
    )
    for cell in cells:
        assert "—" not in cell
        assert " -- " not in cell


def test_guide_paragraphs_every_cell_is_non_empty_under_900_chars_and_ends_in_a_full_stop():
    cells = ex.guide_paragraphs(
        delivered_only=True, n_delivered=68, n_delivered_sections=7, n_delivered_sentences=50,
        longest_passage_chars=18040,
    )
    for cell in cells:
        assert cell
        assert len(cell) < 900
        assert cell.endswith(".")
        assert "\n" not in cell


def test_guide_paragraphs_full_export_and_delivered_only_share_the_delivered_cells():
    delivered_only_cells = ex.guide_paragraphs(
        delivered_only=True, n_delivered=5, n_delivered_sections=2, n_delivered_sentences=4,
        longest_passage_chars=100,
    )
    full_cells = ex.guide_paragraphs(
        has_delivered=True, n_delivered=5, n_delivered_sections=2, n_delivered_sentences=4,
        longest_passage_chars=100,
    )
    for cell in delivered_only_cells:
        assert cell in full_cells


def test_guide_paragraphs_full_export_adds_the_review_sheet_contrast_cell():
    full_cells = ex.guide_paragraphs(
        has_delivered=True, n_delivered=5, n_delivered_sections=2, n_delivered_sentences=4,
    )
    assert ex.DELIVERED_GUIDE_REVIEW_CONTRAST_CELL in full_cells


def test_write_guide_sheet_sets_no_freeze_pane():
    wb = openpyxl.Workbook()
    ex.write_guide_sheet(
        wb.active, workbook_id="abc123", delivered_only=True, has_delivered=True,
        n_delivered=1, n_delivered_sections=1,
    )
    assert wb.active.freeze_panes is None


def test_write_guide_sheet_delivered_only_id_line_names_no_key_file():
    wb = openpyxl.Workbook()
    ex.write_guide_sheet(
        wb.active, workbook_id="abc123", delivered_only=True, has_delivered=True,
        n_delivered=1, n_delivered_sections=1,
    )
    id_text = wb.active.cell(row=wb.active.max_row, column=1).value
    assert "key-out" not in id_text
    assert "refuses to score" not in id_text


def test_write_guide_sheet_delivered_only_id_line_parses_back_to_the_bare_id():
    wb = openpyxl.Workbook()
    ex.write_guide_sheet(
        wb.active, workbook_id="6ee61df4ad585e41", delivered_only=True, has_delivered=True,
        n_delivered=1, n_delivered_sections=1,
    )
    id_text = wb.active.cell(row=wb.active.max_row, column=1).value
    assert id_text.startswith(ex.WORKBOOK_ID_PREFIX)
    token = id_text[len(ex.WORKBOOK_ID_PREFIX):].split()[0].strip()
    assert token == "6ee61df4ad585e41"


def test_write_guide_sheet_full_export_keeps_the_key_out_clause_on_the_id_line():
    wb = openpyxl.Workbook()
    ex.write_guide_sheet(
        wb.active, workbook_id="abc123", has_delivered=True,
        n_delivered=1, n_delivered_sections=1,
    )
    id_text = wb.active.cell(row=wb.active.max_row, column=1).value
    assert "key-out" in id_text
    assert "refuses to score" in id_text


def test_build_delivered_only_workbook_writes_only_guide_provenance_and_delivered():
    authors = ex.load_paper_authors(PAPER_AUTHORS_FIXTURE_PATH)
    wb, record, rows, workbook_id = ex.build_delivered_only_workbook(
        DELIVERED_FIXTURE_PATH, max_cell_chars=32000, demo_run_id="run-x",
        export_time="2026-09-16T00:00:00Z", paper_authors=authors,
    )
    assert wb.sheetnames == ["Guide", "Provenance", "Delivered"]
    assert len(rows) == 2


def test_build_delivered_only_workbook_freezes_delivered_and_provenance_at_a2():
    authors = ex.load_paper_authors(PAPER_AUTHORS_FIXTURE_PATH)
    wb, *_ = ex.build_delivered_only_workbook(
        DELIVERED_FIXTURE_PATH, max_cell_chars=32000, demo_run_id="run-x",
        export_time="t", paper_authors=authors,
    )
    assert wb["Delivered"].freeze_panes == "A2"
    assert wb["Provenance"].freeze_panes == "A2"
    assert wb["Guide"].freeze_panes is None


def test_build_delivered_only_workbook_leaves_delivered_row_heights_unset():
    authors = ex.load_paper_authors(PAPER_AUTHORS_FIXTURE_PATH)
    wb, *_ = ex.build_delivered_only_workbook(
        DELIVERED_FIXTURE_PATH, max_cell_chars=32000, demo_run_id="run-x",
        export_time="t", paper_authors=authors,
    )
    ws = wb["Delivered"]
    for r in range(2, ws.max_row + 1):
        assert ws.row_dimensions[r].height is None


def test_build_delivered_only_workbook_validates_both_human_columns():
    authors = ex.load_paper_authors(PAPER_AUTHORS_FIXTURE_PATH)
    wb, *_ = ex.build_delivered_only_workbook(
        DELIVERED_FIXTURE_PATH, max_cell_chars=32000, demo_run_id="run-x",
        export_time="t", paper_authors=authors,
    )
    ws = wb["Delivered"]
    dv_sqrefs = " ".join(str(dv.sqref) for dv in ws.data_validations.dataValidation)
    header = [c.value for c in ws[1]]
    correct_col = openpyxl.utils.get_column_letter(header.index("human_sentence_correct") + 1)
    supports_col = openpyxl.utils.get_column_letter(header.index("human_quote_supports") + 1)
    assert f"{correct_col}2" in dv_sqrefs
    assert f"{supports_col}2" in dv_sqrefs


def test_write_key_csv_delivered_only_holds_the_header_plus_one_workbook_id_row(tmp_path: Path):
    path = tmp_path / "key.csv"
    ex.write_delivered_only_key_csv(path, workbook_id="abc123")
    with path.open(encoding="utf-8") as fh:
        rows = list(csv.DictReader(fh))
    assert len(rows) == 1
    assert rows[0]["workbook_id"] == "abc123"
    assert rows[0]["real_item_id"] == ""


def test_main_delivered_only_does_not_require_items_cache_dir_results_dir_or_key_out(
    tmp_path: Path,
):
    out = tmp_path / "out.xlsx"
    rc = ex.main([
        "--delivered-only", "--demo-run-id", "run-x",
        "--delivered-evidence", str(DELIVERED_FIXTURE_PATH),
        "--paper-authors", str(PAPER_AUTHORS_FIXTURE_PATH),
        "--out", str(out),
    ])
    assert rc == 0
    assert out.exists()


def test_main_delivered_only_requires_delivered_evidence_and_paper_authors(
    tmp_path: Path, capsys,
):
    out = tmp_path / "out.xlsx"
    rc = ex.main([
        "--delivered-only", "--demo-run-id", "run-x", "--out", str(out),
    ])
    assert rc == 2
    assert "delivered-evidence" in capsys.readouterr().out


def test_main_delivered_only_refuses_an_annotation_or_screening_record(tmp_path: Path, capsys):
    out = tmp_path / "out.xlsx"
    rc = ex.main([
        "--delivered-only", "--demo-run-id", "run-x",
        "--delivered-evidence", str(DELIVERED_FIXTURE_PATH),
        "--paper-authors", str(PAPER_AUTHORS_FIXTURE_PATH),
        "--out", str(out), "--annotation", str(PAPER_AUTHORS_FIXTURE_PATH),
    ])
    assert rc == 2
    assert "refused" in capsys.readouterr().out


def test_main_full_export_requires_paper_authors_with_delivered_evidence(tmp_path: Path, capsys):
    out = tmp_path / "out.xlsx"
    key = tmp_path / "key.csv"
    rc = ex.main([
        "--items", str(Path(__file__).resolve().parents[1] / "claims" / "hss_test_v3_claims.jsonl"),
        "--cache-dir", str(Path(__file__).resolve().parents[1] / "claims" / "data" / "hss_test_v3_fulltext"),
        "--results-dir", str(Path(__file__).resolve().parents[1] / "claims" / "results" / "v6" / "inputs"),
        "--out", str(out), "--key-out", str(key),
        "--delivered-evidence", str(DELIVERED_FIXTURE_PATH),
        "--allow-missing-results", "--allow-missing-chunks",
    ])
    assert rc == 2
    assert "paper-authors" in capsys.readouterr().out


def test_main_without_delivered_only_still_requires_items_cache_dir_results_dir_and_key_out(
    tmp_path: Path, capsys,
):
    out = tmp_path / "out.xlsx"
    rc = ex.main(["--out", str(out)])
    assert rc == 2
    assert "required" in capsys.readouterr().out
