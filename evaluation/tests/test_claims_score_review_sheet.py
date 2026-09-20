"""claims/score_review_sheet.py -- scores a filled review workbook.

No network, no LLM, no backend import. Interpreter: the evaluation venv's python (openpyxl) or
the system python.
"""

from __future__ import annotations

import ast
import csv
import inspect
import json
from pathlib import Path

import export_review_sheet as ex
import openpyxl
import pytest
import score_review_sheet as sc

# --------------------------------------------------------------------------------------
# Normalisation
# --------------------------------------------------------------------------------------


def test_normalise_status_lowercases_and_folds_spaces():
    assert sc.normalise_status("Needs Nuance") == "needs_nuance"
    assert sc.normalise_status(None) == ""


def test_normalise_yes_no_lowercases():
    assert sc.normalise_yes_no("YES") == "yes"


# --------------------------------------------------------------------------------------
# Review-sheet blank/consistency validation
# --------------------------------------------------------------------------------------


def base_row(**overrides):
    row = {
        "_row_number": 2, "item_id": "R-001", "verifier_status": "verified",
        "human_verdict": "verified", "human_agrees_with_verifier": "yes",
        "disagreement_axis": "", "human_quote": "", "human_justification": "",
    }
    row.update(overrides)
    return row


def test_find_review_blanks_clean_agree_row_has_no_problems():
    assert sc.find_review_blanks([base_row()]) == []


def test_find_review_blanks_empty_verdict():
    row = base_row(human_verdict="")
    problems = sc.find_review_blanks([row])
    assert any("empty human_verdict" in p for p in problems)


def test_find_review_blanks_invalid_verdict_value():
    row = base_row(human_verdict="maybe")
    problems = sc.find_review_blanks([row])
    assert any("must be one of" in p and "human_verdict" in p for p in problems)


def test_find_review_blanks_empty_agrees():
    row = base_row(human_agrees_with_verifier="")
    problems = sc.find_review_blanks([row])
    assert any("empty human_agrees_with_verifier" in p for p in problems)


def test_find_review_blanks_contradiction_yes_but_different_verdict():
    row = base_row(human_verdict="unsupported", human_agrees_with_verifier="yes")
    problems = sc.find_review_blanks([row])
    assert any("contradicts" in p for p in problems)


def test_find_review_blanks_contradiction_no_but_same_verdict():
    row = base_row(human_verdict="verified", human_agrees_with_verifier="no")
    problems = sc.find_review_blanks([row])
    assert any("contradicts" in p for p in problems)


def test_find_review_blanks_disagreement_requires_axis_quote_justification():
    row = base_row(
        human_verdict="unsupported", human_agrees_with_verifier="no",
        disagreement_axis="", human_quote="", human_justification="",
    )
    problems = sc.find_review_blanks([row])
    joined = "; ".join(problems)
    assert "disagreement_axis" in joined
    assert "human_quote" in joined
    assert "human_justification" in joined


def test_find_review_blanks_valid_disagreement_row_clean():
    row = base_row(
        human_verdict="unsupported", human_agrees_with_verifier="no",
        disagreement_axis="added detail", human_quote="the paper never says this",
        human_justification="the source omits the detail entirely",
    )
    assert sc.find_review_blanks([row]) == []


def test_find_review_blanks_invalid_disagreement_axis_value():
    row = base_row(
        human_verdict="unsupported", human_agrees_with_verifier="no",
        disagreement_axis="not a real axis", human_quote="q", human_justification="j",
    )
    problems = sc.find_review_blanks([row])
    assert any("disagreement_axis must be one of" in p for p in problems)


def test_validate_review_rows_raises_on_any_problem():
    with pytest.raises(sc.ReviewSheetError):
        sc.validate_review_rows([base_row(human_verdict="")])


def test_validate_review_rows_passes_clean_rows():
    sc.validate_review_rows([base_row()])  # no raise


# --------------------------------------------------------------------------------------
# Screening-sheet blank validation
# --------------------------------------------------------------------------------------


def screening_row(**overrides):
    row = {
        "_row_number": 2, "record_id": "S-001", "status": "include",
        "human_status": "include", "human_agrees": "yes", "human_note": "",
    }
    row.update(overrides)
    return row


def test_find_screening_blanks_clean_row():
    assert sc.find_screening_blanks([screening_row()]) == []


def test_find_screening_blanks_empty_human_status():
    problems = sc.find_screening_blanks([screening_row(human_status="")])
    assert any("empty human_status" in p for p in problems)


def test_find_screening_blanks_bad_human_status():
    problems = sc.find_screening_blanks([screening_row(human_status="maybe")])
    assert any("must be one of" in p for p in problems)


def test_find_screening_blanks_disagree_requires_note():
    problems = sc.find_screening_blanks(
        [screening_row(human_status="exclude", human_agrees="no", human_note="")]
    )
    assert any("empty human_note" in p for p in problems)


def test_find_screening_blanks_disagree_with_note_is_clean():
    problems = sc.find_screening_blanks(
        [screening_row(human_status="exclude", human_agrees="no", human_note="wrong call")]
    )
    assert problems == []


def test_validate_screening_rows_raises():
    with pytest.raises(sc.ReviewSheetError):
        sc.validate_screening_rows([screening_row(human_agrees="")])


def test_find_screening_blanks_quote_present_requires_quote_is_verbatim():
    problems = sc.find_screening_blanks([screening_row(quote="the quote")])
    assert any("empty quote_is_verbatim" in p for p in problems)


def test_find_screening_blanks_quote_is_verbatim_filled_is_clean():
    problems = sc.find_screening_blanks(
        [screening_row(quote="the quote", quote_is_verbatim="yes")]
    )
    assert problems == []


def test_find_screening_blanks_bad_quote_is_verbatim_value():
    problems = sc.find_screening_blanks(
        [screening_row(quote="the quote", quote_is_verbatim="maybe")]
    )
    assert any("quote_is_verbatim must be yes or no" in p for p in problems)


def test_find_screening_blanks_no_quote_does_not_require_quote_is_verbatim():
    assert sc.find_screening_blanks([screening_row(quote="")]) == []


def test_find_screening_blanks_criterion_present_requires_criterion_is_right():
    problems = sc.find_screening_blanks([screening_row(criterion="E1")])
    assert any("empty criterion_is_right" in p for p in problems)


def test_find_screening_blanks_criterion_is_right_filled_is_clean():
    problems = sc.find_screening_blanks(
        [screening_row(criterion="E1", criterion_is_right="no")]
    )
    assert problems == []


def test_find_screening_blanks_bad_criterion_is_right_value():
    problems = sc.find_screening_blanks(
        [screening_row(criterion="E1", criterion_is_right="maybe")]
    )
    assert any("criterion_is_right must be yes or no" in p for p in problems)


def test_find_screening_blanks_no_criterion_does_not_require_criterion_is_right():
    assert sc.find_screening_blanks([screening_row(criterion="")]) == []


# --------------------------------------------------------------------------------------
# Screening sheet left entirely unfilled: treated as not provided, not refused.
# --------------------------------------------------------------------------------------


def blank_screening_row(**overrides):
    row = screening_row(human_status="", human_agrees="", human_note="")
    row.update(overrides)
    return row


def test_screening_sheet_is_blank_true_when_every_human_cell_blank():
    rows = [blank_screening_row(record_id="S-001"), blank_screening_row(record_id="S-002")]
    assert sc.screening_sheet_is_blank(rows) is True


def test_screening_sheet_is_blank_true_for_no_rows():
    assert sc.screening_sheet_is_blank([]) is True


def test_screening_sheet_is_blank_false_when_one_row_filled():
    rows = [blank_screening_row(record_id="S-001"), screening_row(record_id="S-002")]
    assert sc.screening_sheet_is_blank(rows) is False


def test_screening_sheet_is_blank_false_when_quote_is_verbatim_filled():
    rows = [blank_screening_row(record_id="S-001", quote_is_verbatim="yes")]
    assert sc.screening_sheet_is_blank(rows) is False


def test_score_screening_skips_when_sheet_entirely_blank():
    rows = [blank_screening_row(record_id="S-001"), blank_screening_row(record_id="S-002")]
    result = sc.score_screening(rows)
    assert result["screening_skipped"] is True
    assert "blank" in result["screening_skip_reason"]


def test_score_screening_still_raises_when_partly_filled():
    rows = [blank_screening_row(record_id="S-001"), screening_row(record_id="S-002")]
    with pytest.raises(sc.ReviewSheetError):
        sc.score_screening(rows)


def test_render_markdown_summary_states_screening_not_filled_when_skipped():
    result = {"screening": sc.score_screening(
        [blank_screening_row(record_id="S-001"), blank_screening_row(record_id="S-002")]
    )}
    md = sc.render_markdown_summary(result)
    assert "Screening sheet" in md
    assert "not filled" in md


def test_render_markdown_summary_falls_through_to_single_return_when_screening_skipped():
    # The skipped-screening branch must not return from inside the "screening" if block
    # instead of falling through to the function's own trailing return. That would be
    # harmless only because screening happens to be the last section rendered; any section
    # appended after it would silently vanish for every blank-Screening-sheet workbook.
    # This asserts the return structure directly, since there is currently no later section
    # in the render order to exercise the regression behaviourally.
    source = inspect.getsource(sc.render_markdown_summary)
    func_def = ast.parse(source).body[0]
    assert isinstance(func_def, ast.FunctionDef)
    returns = [node for node in ast.walk(func_def) if isinstance(node, ast.Return)]
    assert len(returns) == 1, (
        "render_markdown_summary should have exactly one return statement (the function's "
        "own trailing return); an early return inside the screening-skipped branch would "
        "drop any section appended after screening"
    )


# --------------------------------------------------------------------------------------
# Key loading / binding
# --------------------------------------------------------------------------------------


def write_key(path: Path, rows):
    with path.open("w", encoding="utf-8", newline="") as fh:
        writer = csv.DictWriter(
            fh, fieldnames=["opaque_id", "real_item_id", "set", "construction_category",
                            "expected_label", "workbook_id"],
        )
        writer.writeheader()
        for row in rows:
            writer.writerow(row)


def test_load_key(tmp_path: Path):
    path = tmp_path / "key.csv"
    write_key(path, [{
        "opaque_id": "R-001", "real_item_id": "hss-verbatim-01", "set": "constructed",
        "construction_category": "verbatim", "expected_label": "verified",
        "workbook_id": "wbid",
    }])
    key = sc.load_key(path)
    assert key["R-001"]["real_item_id"] == "hss-verbatim-01"
    assert key["R-001"]["expected_label"] == ["verified"]


def test_load_key_workbook_id(tmp_path: Path):
    path = tmp_path / "key.csv"
    write_key(path, [{
        "opaque_id": "R-001", "real_item_id": "x", "set": "constructed",
        "construction_category": "verbatim", "expected_label": "verified",
        "workbook_id": "wbid42",
    }])
    assert sc.load_key_workbook_id(path) == "wbid42"


def test_validate_key_membership_raises_on_unknown_id():
    with pytest.raises(sc.ReviewSheetError, match="not present in the key"):
        sc.validate_key_membership([{"_row_number": 2, "item_id": "R-999"}], {"R-001": {}})


def test_validate_key_membership_passes_on_exact_match():
    rows = [{"_row_number": 2, "item_id": "R-001"}, {"_row_number": 3, "item_id": "R-002"}]
    key = {"R-001": {}, "R-002": {}}
    sc.validate_key_membership(rows, key)  # no raise


def test_validate_key_membership_raises_when_key_item_has_no_row():
    # a reviewer deleted or blanked a Review row: the key still has 2 items, the sheet has 1.
    rows = [{"_row_number": 2, "item_id": "R-001"}]
    key = {"R-001": {}, "R-002": {}}
    with pytest.raises(sc.ReviewSheetError, match="R-002"):
        sc.validate_key_membership(rows, key)


def test_validate_key_membership_names_both_directions():
    rows = [{"_row_number": 2, "item_id": "R-999"}]
    key = {"R-001": {}}
    with pytest.raises(sc.ReviewSheetError) as excinfo:
        sc.validate_key_membership(rows, key)
    message = str(excinfo.value)
    assert "not present in the key" in message  # extra: R-999 at row 2
    assert "R-001" in message  # missing: the key's own R-001 has no row


def test_validate_key_membership_raises_on_duplicated_item_id():
    rows = [
        {"_row_number": 2, "item_id": "R-001"}, {"_row_number": 3, "item_id": "R-001"},
    ]
    key = {"R-001": {}}
    with pytest.raises(sc.ReviewSheetError, match="more than one Review row"):
        sc.validate_key_membership(rows, key)


def test_validate_workbook_binding_raises_on_mismatch():
    with pytest.raises(sc.ReviewSheetError, match="does not match"):
        sc.validate_workbook_binding("wbid1", "wbid2")


def test_validate_workbook_binding_passes_when_both_none():
    sc.validate_workbook_binding(None, None)  # no raise


def test_validate_workbook_binding_passes_when_equal():
    sc.validate_workbook_binding("wbid", "wbid")  # no raise


# --------------------------------------------------------------------------------------
# Provenance sheet reading and --demo-run-id refusal
# --------------------------------------------------------------------------------------


def write_provenance_workbook(path: Path, fields: dict[str, str]) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Provenance"
    ws.append(["field", "value"])
    for field, value in fields.items():
        ws.append([field, value])
    wb.save(path)


def test_read_provenance_fields_reads_field_value_pairs(tmp_path: Path):
    path = tmp_path / "wb.xlsx"
    write_provenance_workbook(path, {
        "demo_run_id": "20260910-095226", "export_time": "2026-09-10T12:00:00+00:00",
    })
    fields = sc.read_provenance_fields(path)
    assert fields["demo_run_id"] == "20260910-095226"
    assert fields["export_time"] == "2026-09-10T12:00:00+00:00"


def test_read_provenance_fields_none_when_no_provenance_sheet(tmp_path: Path):
    path = tmp_path / "wb.xlsx"
    wb = openpyxl.Workbook()
    wb.active.title = "NotProvenance"
    wb.save(path)
    assert sc.read_provenance_fields(path) is None


def test_validate_demo_run_id_passes_when_expected_none():
    sc.validate_demo_run_id("20260909-071345", None)  # no raise


def test_validate_demo_run_id_passes_when_equal():
    sc.validate_demo_run_id("20260910-095226", "20260910-095226")  # no raise


def test_validate_demo_run_id_raises_on_mismatch():
    with pytest.raises(sc.ReviewSheetError, match="20260909-071345"):
        sc.validate_demo_run_id("20260909-071345", "20260910-095226")


# --------------------------------------------------------------------------------------
# --export-time refusal: --demo-run-id alone cannot tell two exports of the same demo run
# apart, since both share the same demo_run_id and workbook id (both depend only on the kept
# item order and the seed, not on when the export ran).
# --------------------------------------------------------------------------------------


def test_validate_export_time_passes_when_expected_none():
    sc.validate_export_time("2026-09-10T13:09:37+00:00", None)  # no raise


def test_validate_export_time_passes_when_equal():
    sc.validate_export_time(
        "2026-09-10T13:09:37+00:00", "2026-09-10T13:09:37+00:00",
    )  # no raise


def test_validate_export_time_raises_on_mismatch():
    with pytest.raises(sc.ReviewSheetError, match="2026-09-09T07:13:45"):
        sc.validate_export_time(
            "2026-09-09T07:13:45+00:00", "2026-09-10T13:09:37+00:00",
        )


# --------------------------------------------------------------------------------------
# Scoring -- Review sheet
# --------------------------------------------------------------------------------------


def make_key():
    return {
        "R-001": {"real_item_id": "hss-verbatim-01", "set": "constructed",
                  "construction_category": "verbatim", "expected_label": ["verified"]},
        "R-002": {"real_item_id": "real-test-01", "set": "real",
                  "construction_category": "", "expected_label": []},
        "R-003": {"real_item_id": "hss-altered-01", "set": "constructed",
                  "construction_category": "altered", "expected_label": ["unsupported"]},
    }


def make_rows():
    return [
        base_row(item_id="R-001", verifier_status="verified", human_verdict="verified",
                 human_agrees_with_verifier="yes"),
        base_row(item_id="R-002", verifier_status="needs_nuance", human_verdict="needs_nuance",
                 human_agrees_with_verifier="yes"),
        base_row(
            item_id="R-003", verifier_status="needs_nuance", human_verdict="unsupported",
            human_agrees_with_verifier="no", disagreement_axis="added detail",
            human_quote="q", human_justification="j",
        ),
    ]


def test_score_agreement_with_verifier_overall_and_per_set():
    result = sc.score(make_rows(), make_key())
    av = result["agreement_with_verifier"]
    assert av["overall"]["n"] == 3
    assert av["overall"]["n_correct"] == 2
    assert av["per_set"]["constructed"]["n"] == 2
    assert av["per_set"]["real"]["n"] == 1
    assert av["per_status"]["needs_nuance"]["n"] == 2
    # A judged status carries its own n_withheld (0 here) and no nested "withheld" key,
    # since nothing in it was withheld.
    assert av["per_status"]["needs_nuance"]["n_withheld"] == 0
    assert "withheld" not in av["per_status"]["needs_nuance"]
    assert av["n_items"] == 3
    assert av["n_judged"] == 3
    assert av["n_withheld"] == 0


# A no_full_text row's human_verdict is dictated by the Guide, not a judgement, and must not
# inflate the headline agreement/kappa by construction.
def test_score_agreement_with_verifier_withholds_no_full_text_from_overall_and_kappa():
    key = make_key()
    key["R-004"] = {
        "real_item_id": "hss-wrong_paper-01", "set": "constructed",
        "construction_category": "wrong_paper", "expected_label": ["no_full_text"],
    }
    rows = make_rows() + [
        base_row(
            item_id="R-004", verifier_status="no_full_text", human_verdict="no_full_text",
            human_agrees_with_verifier="yes",
        ),
    ]
    result = sc.score(rows, key)
    av = result["agreement_with_verifier"]
    assert av["n_items"] == 4
    assert av["n_judged"] == 3
    assert av["n_withheld"] == 1
    # unchanged from the three-row case: the withheld row is not in overall/kappa.
    assert av["overall"]["n"] == 3
    assert av["overall"]["n_correct"] == 2
    # The no_full_text bucket's own row is withheld by definition (every row in it is), so
    # its main block is empty and the row is reported nested under "withheld" instead, the
    # same convention score_screening's per_status gives a guard-demoted needs_review
    # bucket.
    assert av["per_status"]["no_full_text"]["n"] == 0
    assert av["per_status"]["no_full_text"]["n_withheld"] == 1
    assert av["per_status"]["no_full_text"]["withheld"]["n"] == 1
    assert av["per_status"]["no_full_text"]["withheld"]["n_correct"] == 1
    assert av["per_status"]["no_full_text"]["withheld"]["accuracy"] == 1.0


def test_score_agreement_with_verifier_kappa_unaffected_by_no_full_text_rows():
    key = make_key()
    key["R-004"] = {
        "real_item_id": "hss-wrong_paper-01", "set": "constructed",
        "construction_category": "wrong_paper", "expected_label": ["no_full_text"],
    }
    judged_only = sc.score(make_rows(), make_key())
    with_withheld = sc.score(
        make_rows() + [
            base_row(
                item_id="R-004", verifier_status="no_full_text",
                human_verdict="no_full_text", human_agrees_with_verifier="yes",
            ),
        ],
        key,
    )
    assert (
        with_withheld["agreement_with_verifier"]["kappa"]
        == judged_only["agreement_with_verifier"]["kappa"]
    )


# per_set must mirror the same withheld-row exclusion as the headline, and state its own
# denominators.
def test_score_agreement_with_verifier_per_set_withholds_no_full_text_and_states_denominators():
    key = make_key()
    key["R-004"] = {
        "real_item_id": "hss-wrong_paper-01", "set": "constructed",
        "construction_category": "wrong_paper", "expected_label": ["no_full_text"],
    }
    rows = make_rows() + [
        base_row(
            item_id="R-004", verifier_status="no_full_text", human_verdict="no_full_text",
            human_agrees_with_verifier="yes",
        ),
    ]
    result = sc.score(rows, key)
    per_set = result["agreement_with_verifier"]["per_set"]
    # constructed: R-001, R-003, R-004 -- R-004 is withheld, so n (judged) is 2.
    assert per_set["constructed"]["n"] == 2
    assert per_set["constructed"]["n_items"] == 3
    assert per_set["constructed"]["n_judged"] == 2
    assert per_set["constructed"]["n_withheld"] == 1
    assert per_set["real"]["n"] == 1
    assert per_set["real"]["n_items"] == 1
    assert per_set["real"]["n_withheld"] == 0


def test_score_disagreement_axis_counts():
    result = sc.score(make_rows(), make_key())
    assert result["disagreement_axis_counts"] == {"added detail": 1}


def test_score_overridden_rows_lists_disagreements_only():
    result = sc.score(make_rows(), make_key())
    assert len(result["overridden_rows"]) == 1
    assert result["overridden_rows"][0]["real_item_id"] == "hss-altered-01"


def test_score_agreement_with_model_labels_none_when_no_sheet():
    result = sc.score(make_rows(), make_key(), model_labels=None)
    assert result["agreement_with_model_labels"] is None


def test_score_agreement_with_model_labels_computed_when_present():
    model_labels = {
        "R-001": {"model_annotation_label": "verified"},
        "R-003": {"model_annotation_label": "unsupported"},
    }
    result = sc.score(make_rows(), make_key(), model_labels=model_labels)
    aml = result["agreement_with_model_labels"]
    assert aml["n"] == 2
    assert aml["n_correct"] == 2
    assert aml["n_items"] == 2
    assert aml["n_judged"] == 2
    assert aml["n_withheld"] == 0


# Mirrors agreement_with_verifier's withheld-row exclusion.
def test_score_agreement_with_model_labels_withholds_no_full_text_and_states_denominators():
    key = make_key()
    key["R-004"] = {
        "real_item_id": "hss-wrong_paper-01", "set": "constructed",
        "construction_category": "wrong_paper", "expected_label": ["no_full_text"],
    }
    rows = make_rows() + [
        base_row(
            item_id="R-004", verifier_status="no_full_text", human_verdict="no_full_text",
            human_agrees_with_verifier="yes",
        ),
    ]
    model_labels = {
        "R-001": {"model_annotation_label": "verified"},
        "R-003": {"model_annotation_label": "unsupported"},
        "R-004": {"model_annotation_label": "no_full_text"},
    }
    result = sc.score(rows, key, model_labels=model_labels)
    aml = result["agreement_with_model_labels"]
    assert aml["n_items"] == 3
    assert aml["n_judged"] == 2
    assert aml["n_withheld"] == 1
    assert aml["n"] == 2
    assert aml["n_correct"] == 2


def test_score_raises_on_blank_and_lists_row():
    rows = make_rows()
    rows[0]["human_verdict"] = ""
    with pytest.raises(sc.ReviewSheetError, match="row"):
        sc.score(rows, make_key())


def test_score_raises_on_key_mismatch():
    rows = make_rows()
    rows[0]["item_id"] = "R-999"
    with pytest.raises(sc.ReviewSheetError, match="not present in the key"):
        sc.score(rows, make_key())


def test_score_raises_on_workbook_binding_mismatch():
    with pytest.raises(sc.ReviewSheetError, match="does not match"):
        sc.score(make_rows(), make_key(), review_workbook_id="a", key_workbook_id="b")


# --------------------------------------------------------------------------------------
# Scoring -- Screening sheet
# --------------------------------------------------------------------------------------


def test_score_screening_overall_and_per_status_and_criterion():
    rows = [
        screening_row(record_id="S-001", status="include", human_status="include"),
        screening_row(record_id="S-002", status="exclude", human_status="exclude",
                       criterion="E1", criterion_is_right="yes"),
        screening_row(record_id="S-003", status="exclude", human_status="include",
                       human_agrees="no", human_note="wrong", criterion="E1",
                       criterion_is_right="no"),
    ]
    result = sc.score_screening(rows)
    assert result["overall"]["n"] == 3
    assert result["overall"]["n_correct"] == 2
    assert result["per_status"]["exclude"]["n"] == 2
    assert result["per_criterion"]["E1"]["n"] == 2
    assert len(result["overridden_records"]) == 1


def test_score_screening_raises_on_blank():
    with pytest.raises(sc.ReviewSheetError):
        sc.score_screening([screening_row(human_status="")])


# Unscreened census rows cannot agree with a human_status by construction and must not drag
# the headline accuracy/kappa down.
def test_score_screening_excludes_unscreened_from_overall_and_kappa():
    rows = [
        screening_row(record_id="S-001", status="include", human_status="include"),
        screening_row(record_id="S-002", status="include", human_status="include"),
        screening_row(record_id="S-003", status="unscreened", human_status="exclude"),
    ]
    result = sc.score_screening(rows)
    assert result["overall"]["n"] == 2
    assert result["overall"]["n_correct"] == 2
    assert result["overall"]["accuracy"] == 1.0
    assert result["n_items"] == 3
    assert result["n_scored"] == 2
    assert result["n_unscreened"] == 1
    # the unscreened row is still visible on its own line.
    assert result["per_status"]["unscreened"]["n"] == 1


def test_score_screening_kappa_unaffected_by_unscreened_rows():
    scored_only = sc.score_screening([
        screening_row(record_id="S-001", status="include", human_status="include"),
        screening_row(record_id="S-002", status="exclude", human_status="include",
                      human_agrees="no", human_note="wrong"),
    ])
    with_unscreened = sc.score_screening([
        screening_row(record_id="S-001", status="include", human_status="include"),
        screening_row(record_id="S-002", status="exclude", human_status="include",
                      human_agrees="no", human_note="wrong"),
        screening_row(record_id="S-003", status="unscreened", human_status="exclude"),
    ])
    assert with_unscreened["kappa"] == scored_only["kappa"]


# A needs_review row demoted onto the same criterion id as a real exclusion must not be
# pooled into per_criterion with it.
def test_score_screening_per_criterion_excludes_needs_review_rows():
    rows = [
        screening_row(record_id="S-001", status="exclude", human_status="exclude",
                      criterion="E1", criterion_is_right="yes"),
        screening_row(record_id="S-002", status="needs_review", human_status="needs_review",
                      criterion="E1", criterion_is_right="yes"),
    ]
    result = sc.score_screening(rows)
    assert result["per_criterion"]["E1"]["n"] == 1


def test_score_screening_screen_in_share_reads_stratum_column():
    rows = [
        screening_row(record_id="S-001", status="exclude", human_status="exclude",
                      criterion="E1", criterion_is_right="yes", stratum="excluded_numbered"),
        screening_row(record_id="S-002", status="exclude", human_status="include",
                      human_agrees="no", human_note="wrong", criterion="E1",
                      criterion_is_right="no", stratum="excluded_numbered"),
        screening_row(record_id="S-003", status="exclude", human_status="needs_review",
                      human_agrees="no", human_note="unsure", criterion="TOPIC",
                      criterion_is_right="no", stratum="excluded_off_topic"),
    ]
    result = sc.score_screening(rows)
    assert result["screen_in_share"]["excluded_numbered"] == {
        "n": 2, "n_screened_in": 1, "rate": 0.5,
    }
    assert result["screen_in_share"]["excluded_off_topic"] == {
        "n": 1, "n_screened_in": 1, "rate": 1.0,
    }


def test_score_screening_screen_in_share_none_when_stratum_not_sampled():
    result = sc.score_screening([screening_row(record_id="S-001", stratum="included")])
    assert result["screen_in_share"]["excluded_numbered"] == {
        "n": 0, "n_screened_in": 0, "rate": None,
    }


def test_score_screening_quote_and_criterion_rates():
    rows = [
        screening_row(
            record_id="S-001", quote="the quote", quote_is_verbatim="yes",
            criterion="E1", criterion_is_right="yes",
        ),
        screening_row(
            record_id="S-002", quote="another quote", quote_is_verbatim="no",
            criterion="E2", criterion_is_right="yes",
        ),
        # no quote/criterion on this row -- excluded from both denominators.
        screening_row(record_id="S-003"),
    ]
    result = sc.score_screening(rows)
    assert result["quote_is_verbatim_rate"]["n"] == 2
    assert result["quote_is_verbatim_rate"]["n_yes"] == 1
    assert result["quote_is_verbatim_rate"]["rate"] == 0.5
    assert result["quote_is_verbatim_rate"]["unanchored_exclude"] == {
        "n": 0, "n_yes": 0, "rate": None,
    }
    assert result["criterion_is_right_rate"] == {"n": 2, "n_yes": 2, "rate": 1.0}


def test_score_screening_rate_none_when_no_rows_answer():
    result = sc.score_screening([screening_row(record_id="S-001")])
    assert result["quote_is_verbatim_rate"]["n"] == 0
    assert result["quote_is_verbatim_rate"]["n_yes"] == 0
    assert result["quote_is_verbatim_rate"]["rate"] is None
    assert result["criterion_is_right_rate"] == {"n": 0, "n_yes": 0, "rate": None}


# A guard-demoted unanchored_exclude row must not pool with the genuine quote-fidelity rows,
# and must be reported on its own.
def test_score_screening_quote_is_verbatim_rate_excludes_unanchored_exclude_rows():
    rows = [
        screening_row(
            record_id="S-001", quote="the quote", quote_is_verbatim="yes",
        ),
        screening_row(
            record_id="S-002", quote="another quote", quote_is_verbatim="yes",
        ),
        screening_row(
            record_id="S-003", status="needs_review", human_status="needs_review",
            needs_review_reason="unanchored_exclude", quote="unanchored quote",
            quote_is_verbatim="no",
        ),
    ]
    result = sc.score_screening(rows)
    # the two genuine rows both say "yes" -- a rate of 1.0, not dragged down by the guard row.
    rate = result["quote_is_verbatim_rate"]
    assert rate["n"] == 2
    assert rate["n_yes"] == 2
    assert rate["rate"] == 1.0
    assert rate["unanchored_exclude"] == {"n": 1, "n_yes": 0, "rate": 0.0}


# A row the Guide's UNANCHORED_EXCLUDE_NOTE tells the reviewer to answer with a dictated
# human_status/human_agrees (exclude/no) even though the tool's own status is needs_review
# must not count as a disagreement in overall or in needs_review's own per_status block; it
# is reported separately instead, in the same accuracy-block shape.
def test_score_screening_excludes_unanchored_exclude_from_overall_and_kappa():
    rows = [
        screening_row(record_id="S-001", status="include", human_status="include"),
        screening_row(record_id="S-002", status="include", human_status="include"),
        screening_row(
            record_id="S-003", status="needs_review", human_status="exclude",
            human_agrees="no", human_note="guard demoted this",
            needs_review_reason="unanchored_exclude",
        ),
    ]
    result = sc.score_screening(rows)
    assert result["overall"]["n"] == 2
    assert result["overall"]["n_correct"] == 2
    assert result["overall"]["accuracy"] == 1.0
    assert result["n_items"] == 3
    assert result["n_scored"] == 2
    assert result["n_unanchored_exclude"] == 1
    assert result["overall"]["unanchored_exclude"]["n"] == 1
    assert result["overall"]["unanchored_exclude"]["n_correct"] == 0


def test_score_screening_per_status_needs_review_excludes_unanchored_exclude_rows():
    rows = [
        screening_row(record_id="S-001", status="needs_review", human_status="needs_review"),
        screening_row(
            record_id="S-002", status="needs_review", human_status="exclude",
            human_agrees="no", human_note="guard demoted this",
            needs_review_reason="unanchored_exclude",
        ),
    ]
    result = sc.score_screening(rows)
    needs_review_block = result["per_status"]["needs_review"]
    assert needs_review_block["n"] == 1
    assert needs_review_block["n_correct"] == 1
    assert needs_review_block["unanchored_exclude"]["n"] == 1
    assert needs_review_block["unanchored_exclude"]["n_correct"] == 0


def test_score_screening_kappa_unaffected_by_unanchored_exclude_rows():
    scored_only = sc.score_screening([
        screening_row(record_id="S-001", status="include", human_status="include"),
        screening_row(record_id="S-002", status="exclude", human_status="include",
                      human_agrees="no", human_note="wrong"),
    ])
    with_unanchored = sc.score_screening([
        screening_row(record_id="S-001", status="include", human_status="include"),
        screening_row(record_id="S-002", status="exclude", human_status="include",
                      human_agrees="no", human_note="wrong"),
        screening_row(
            record_id="S-003", status="needs_review", human_status="exclude",
            human_agrees="no", human_note="guard demoted this",
            needs_review_reason="unanchored_exclude",
        ),
    ])
    assert with_unanchored["kappa"] == scored_only["kappa"]


def test_score_screening_per_status_has_no_unanchored_exclude_key_when_none_sampled():
    result = sc.score_screening([
        screening_row(record_id="S-001", status="needs_review", human_status="needs_review"),
    ])
    assert "unanchored_exclude" not in result["per_status"]["needs_review"]


# --------------------------------------------------------------------------------------
# Delivered sheet: the final product a user receives, judged sentence by sentence
# --------------------------------------------------------------------------------------

DELIVERED_FIXTURE_PATH = (
    Path(__file__).resolve().parent / "fixtures" / "delivered_evidence_sample.json"
)


def delivered_row(**overrides):
    row = {
        "_row_number": 2, "row": 1, "sentence": "the delivered sentence",
        "citation_text": "(Author, 2020)", "paper_title": "Paper", "paper_doi": "10.1/p",
        "evidence_quotes": "an evidence quote", "source_passage": "the source passage",
        "human_sentence_correct": "yes", "human_quote_supports": "yes", "human_note": "",
    }
    row.update(overrides)
    return row


def test_find_delivered_blanks_clean_row_has_no_problems():
    assert sc.find_delivered_blanks([delivered_row()]) == []


def test_find_delivered_blanks_empty_sentence_correct():
    problems = sc.find_delivered_blanks([delivered_row(human_sentence_correct="")])
    assert any("empty human_sentence_correct" in p for p in problems)


def test_find_delivered_blanks_bad_sentence_correct_value():
    problems = sc.find_delivered_blanks([delivered_row(human_sentence_correct="maybe")])
    assert any("human_sentence_correct must be yes or no" in p for p in problems)


def test_find_delivered_blanks_empty_quote_supports():
    problems = sc.find_delivered_blanks([delivered_row(human_quote_supports="")])
    assert any("empty human_quote_supports" in p for p in problems)


def test_find_delivered_blanks_bad_quote_supports_value():
    problems = sc.find_delivered_blanks([delivered_row(human_quote_supports="maybe")])
    assert any("human_quote_supports must be yes or no" in p for p in problems)


def test_find_delivered_blanks_no_on_sentence_correct_requires_note():
    problems = sc.find_delivered_blanks(
        [delivered_row(human_sentence_correct="no", human_note="")]
    )
    assert any("empty human_note" in p for p in problems)


def test_find_delivered_blanks_no_on_quote_supports_requires_note():
    problems = sc.find_delivered_blanks(
        [delivered_row(human_quote_supports="no", human_note="")]
    )
    assert any("empty human_note" in p for p in problems)


def test_find_delivered_blanks_no_with_note_is_clean():
    problems = sc.find_delivered_blanks(
        [delivered_row(human_sentence_correct="no", human_note="the sentence overstates it")]
    )
    assert problems == []


def test_find_delivered_blanks_yes_does_not_require_note():
    assert sc.find_delivered_blanks([delivered_row()]) == []


def test_validate_delivered_rows_raises_on_any_problem():
    with pytest.raises(sc.ReviewSheetError):
        sc.validate_delivered_rows([delivered_row(human_sentence_correct="")])


def test_validate_delivered_rows_passes_clean_rows():
    sc.validate_delivered_rows([delivered_row()])  # no raise


def test_read_delivered_sheet_none_when_no_sheet(tmp_path: Path):
    wb = openpyxl.Workbook()
    path = tmp_path / "workbook.xlsx"
    wb.save(path)
    assert sc.read_delivered_sheet(path) is None


def test_read_delivered_sheet_reads_rows(tmp_path: Path):
    wb = openpyxl.Workbook()
    ex.build_delivered_sheet(wb, DELIVERED_FIXTURE_PATH, max_cell_chars=32000)
    path = tmp_path / "workbook.xlsx"
    wb.save(path)
    rows = sc.read_delivered_sheet(path)
    assert len(rows) == 2
    assert rows[0]["row"] == 1
    assert rows[0]["human_sentence_correct"] is None


def make_delivered_rows():
    return [
        delivered_row(row=1, human_sentence_correct="yes", human_quote_supports="yes"),
        delivered_row(row=2, human_sentence_correct="yes", human_quote_supports="no",
                      human_note="the quote is about a different claim"),
        delivered_row(row=3, human_sentence_correct="no", human_quote_supports="yes",
                      human_note="the sentence overgeneralises"),
        delivered_row(row=4, human_sentence_correct="yes", human_quote_supports="yes"),
    ]


def test_score_delivered_counts_and_rates():
    result = sc.score_delivered(make_delivered_rows())
    assert result["n_rows"] == 4
    assert result["n_judged"] == 4
    assert result["sentence_correct_count"] == 3
    assert result["sentence_correct_rate"] == pytest.approx(0.75)
    assert result["quote_supports_count"] == 3
    assert result["quote_supports_rate"] == pytest.approx(0.75)


def test_score_delivered_wilson_interval_on_each_rate():
    result = sc.score_delivered(make_delivered_rows())
    lo, hi = sc.wilson_interval(3, 4)
    assert result["sentence_correct_wilson_95"]["lo"] == pytest.approx(lo)
    assert result["sentence_correct_wilson_95"]["hi"] == pytest.approx(hi)
    assert result["quote_supports_wilson_95"]["lo"] == pytest.approx(lo)
    assert result["quote_supports_wilson_95"]["hi"] == pytest.approx(hi)


def test_score_delivered_lists_rows_judged_no():
    result = sc.score_delivered(make_delivered_rows())
    rows_judged_no = result["rows_judged_no"]
    assert {r["row"] for r in rows_judged_no} == {2, 3}


# --------------------------------------------------------------------------------------
# Per-section breakdown: a single run can deliver rows from more than one generated section
# into one Delivered sheet; the headline stays over every row, and a per-section table sits
# beside it.
# --------------------------------------------------------------------------------------


def test_delivered_section_order_returns_distinct_labels_in_first_seen_order():
    rows = [
        delivered_row(row=1, section_title="Introduction"),
        delivered_row(row=2, section_title="Literature Review"),
        delivered_row(row=3, section_title="Introduction"),
    ]
    assert sc.delivered_section_order(rows) == ["Introduction", "Literature Review"]


def test_delivered_section_order_labels_a_row_with_no_section_title():
    rows = [delivered_row(row=1)]
    assert sc.delivered_section_order(rows) == [sc.UNRECORDED_SECTION_LABEL]


def test_score_delivered_reports_n_sections_and_a_by_section_table():
    rows = [
        delivered_row(row=1, section_title="Introduction",
                      human_sentence_correct="yes", human_quote_supports="yes"),
        delivered_row(row=2, section_title="Introduction",
                      human_sentence_correct="no", human_quote_supports="yes",
                      human_note="overstates the finding"),
        delivered_row(row=3, section_title="Literature Review",
                      human_sentence_correct="yes", human_quote_supports="yes"),
    ]
    result = sc.score_delivered(rows)
    assert result["n_sections"] == 2
    by_section = result["by_section"]
    assert set(by_section) == {"Introduction", "Literature Review"}
    assert by_section["Introduction"]["n_rows"] == 2
    assert by_section["Introduction"]["sentence_correct_count"] == 1
    assert by_section["Introduction"]["sentence_correct_rate"] == pytest.approx(0.5)
    assert by_section["Literature Review"]["n_rows"] == 1
    assert by_section["Literature Review"]["sentence_correct_rate"] == pytest.approx(1.0)


def test_score_delivered_by_section_excludes_its_own_fidelity_failed_rows():
    rows = [
        delivered_row(row=1, section_title="Introduction", source_located="no"),
        delivered_row(row=2, section_title="Introduction"),
    ]
    result = sc.score_delivered(rows)
    assert result["by_section"]["Introduction"]["n_judged"] == 1
    assert result["by_section"]["Introduction"]["n_fidelity_excluded"] == 1


def test_score_delivered_defaults_to_one_unrecorded_section_when_no_section_title():
    result = sc.score_delivered(make_delivered_rows())
    assert result["n_sections"] == 1
    assert sc.UNRECORDED_SECTION_LABEL in result["by_section"]
    assert result["by_section"][sc.UNRECORDED_SECTION_LABEL]["n_rows"] == 4


def test_score_delivered_raises_on_blank_human_cell():
    with pytest.raises(sc.ReviewSheetError):
        sc.score_delivered([delivered_row(human_sentence_correct="")])


# --------------------------------------------------------------------------------------
# source_located/sentence_in_draft are excluded from the headline and reported separately,
# rather than silently forced into the rate either way.
# --------------------------------------------------------------------------------------


def test_score_delivered_rows_with_blank_fidelity_fields_are_not_excluded():
    """Neither field is present on `delivered_row`'s own defaults (the pre-existing schema
    shape, or a demo run that never recorded the flag) -- a blank cell is not treated as a
    fidelity failure, only an explicit no is."""
    result = sc.score_delivered(make_delivered_rows())
    assert result["n_fidelity_excluded"] == 0
    assert result["n_judged"] == 4
    assert result["fidelity_excluded_rows"] == []


def test_score_delivered_excludes_a_row_with_source_not_located_from_the_headline():
    rows = make_delivered_rows()
    rows[0] = delivered_row(
        row=1, human_sentence_correct="yes", human_quote_supports="yes",
        source_located="no",
    )
    result = sc.score_delivered(rows)
    assert result["n_rows"] == 4
    assert result["n_judged"] == 3
    assert result["n_fidelity_excluded"] == 1
    # Row 1 was itself a yes/yes row; excluding it must not move the yes counts up.
    assert result["sentence_correct_count"] == 2
    assert result["quote_supports_count"] == 2


def test_score_delivered_excludes_a_row_with_sentence_not_in_draft_from_the_headline():
    rows = make_delivered_rows()
    rows[3] = delivered_row(
        row=4, human_sentence_correct="yes", human_quote_supports="yes",
        sentence_in_draft="no",
    )
    result = sc.score_delivered(rows)
    assert result["n_judged"] == 3
    assert result["n_fidelity_excluded"] == 1


def test_score_delivered_reports_fidelity_excluded_rows_separately():
    rows = make_delivered_rows()
    rows[0] = delivered_row(
        row=1, human_sentence_correct="yes", human_quote_supports="yes",
        source_located="no",
    )
    result = sc.score_delivered(rows)
    excluded = result["fidelity_excluded_rows"]
    assert {r["row"] for r in excluded} == {1}
    assert excluded[0]["source_located"] == "no"


def test_score_delivered_excluded_row_is_not_double_counted_in_rows_judged_no():
    rows = make_delivered_rows()
    rows[1] = delivered_row(
        row=2, human_sentence_correct="yes", human_quote_supports="no",
        human_note="the quote is about a different claim", source_located="no",
    )
    result = sc.score_delivered(rows)
    assert {r["row"] for r in result["rows_judged_no"]} == {3}
    assert {r["row"] for r in result["fidelity_excluded_rows"]} == {2}


def test_score_delivered_rows_judged_no_carries_its_own_section_label():
    """``row`` restarts at 1 in every section (`_delivered_section_label`), so a row must
    carry its own section alongside it, or a multi-section sheet's rows judged no cannot be
    told apart by row number alone."""
    rows = [
        delivered_row(
            row=2, section_title="Section A", human_quote_supports="no",
            human_note="the quote is about a different claim",
        ),
        delivered_row(
            row=2, section_title="Section B", human_sentence_correct="no",
            human_note="the sentence overgeneralises",
        ),
    ]
    result = sc.score_delivered(rows)
    sections_by_row = {(r["row"], r["section"]) for r in result["rows_judged_no"]}
    assert sections_by_row == {(2, "Section A"), (2, "Section B")}


def test_score_delivered_fidelity_excluded_rows_carry_their_own_section_label():
    """Section labels also identify rows in the fidelity-excluded table."""
    rows = [
        delivered_row(row=2, section_title="Section A", source_located="no"),
        delivered_row(row=2, section_title="Section B", passage_located="no"),
    ]
    result = sc.score_delivered(rows)
    sections_by_row = {(r["row"], r["section"]) for r in result["fidelity_excluded_rows"]}
    assert sections_by_row == {(2, "Section A"), (2, "Section B")}


# --------------------------------------------------------------------------------------
# passage_located is a fidelity failure in its own right: source_located alone only records
# that the chunk was fetched, not that a passage was found in it.
# --------------------------------------------------------------------------------------


def test_score_delivered_excludes_a_row_with_passage_not_located_from_the_headline():
    """A row with source_located true (the chunk fetch succeeded) but passage_located
    false (no evidence quote could be found in it, e.g. an en dash on the chunk side
    against a hyphen on the quote side) must still be excluded from the headline -- this
    is the exact case round 2's own finding text named and round 3 found still not
    excluded."""
    rows = make_delivered_rows()
    rows[0] = delivered_row(
        row=1, human_sentence_correct="yes", human_quote_supports="yes",
        source_located="yes", passage_located="no",
    )
    result = sc.score_delivered(rows)
    assert result["n_rows"] == 4
    assert result["n_judged"] == 3
    assert result["n_fidelity_excluded"] == 1
    assert result["sentence_correct_count"] == 2
    assert result["quote_supports_count"] == 2


def test_score_delivered_rows_with_blank_passage_located_are_not_excluded():
    """A row that never recorded passage_located (an older delivered_evidence.json) is
    not treated as a failure -- only an explicit no is."""
    result = sc.score_delivered(make_delivered_rows())
    assert result["n_fidelity_excluded"] == 0
    assert result["n_judged"] == 4


def test_score_delivered_fidelity_excluded_rows_report_passage_located_and_unlocated_quotes():
    rows = make_delivered_rows()
    rows[0] = delivered_row(
        row=1, human_sentence_correct="yes", human_quote_supports="yes",
        source_located="yes", passage_located="no",
        unlocated_quotes="a quote that could not be found",
    )
    result = sc.score_delivered(rows)
    excluded = result["fidelity_excluded_rows"]
    assert {r["row"] for r in excluded} == {1}
    assert excluded[0]["passage_located"] == "no"
    assert excluded[0]["unlocated_quotes"] == "a quote that could not be found"


# --------------------------------------------------------------------------------------
# The headline is selected by testing each row object itself, not by round-tripping through
# the sheet's own (unprotected) row cell.
# --------------------------------------------------------------------------------------


def test_score_delivered_headline_not_dropped_by_a_duplicated_row_cell():
    """Two rows sharing the same row cell value, one of them flagged, must not drop the
    other, unflagged one from the headline (reproduced by round 3's review: the old
    row-cell round trip gave n_judged 1 of 3 here instead of 2 of 3)."""
    rows = [
        delivered_row(row=1, human_sentence_correct="yes", human_quote_supports="yes"),
        delivered_row(
            row=1, human_sentence_correct="yes", human_quote_supports="yes",
            source_located="no",
        ),
        delivered_row(row=1, human_sentence_correct="yes", human_quote_supports="yes"),
    ]
    result = sc.score_delivered(rows)
    assert result["n_rows"] == 3
    assert result["n_fidelity_excluded"] == 1
    assert result["n_judged"] == 2


def test_score_delivered_headline_not_dropped_by_a_blank_row_cell():
    rows = [
        delivered_row(row=None, human_sentence_correct="yes", human_quote_supports="yes"),
        delivered_row(
            row=None, human_sentence_correct="yes", human_quote_supports="yes",
            source_located="no",
        ),
        delivered_row(row=None, human_sentence_correct="yes", human_quote_supports="yes"),
    ]
    result = sc.score_delivered(rows)
    assert result["n_rows"] == 3
    assert result["n_fidelity_excluded"] == 1
    assert result["n_judged"] == 2


# --------------------------------------------------------------------------------------
# location_section_mismatch is listed beside the fidelity-excluded rows, not folded into
# the same exclusion and not silently dropped.
# --------------------------------------------------------------------------------------


def test_delivered_location_mismatch_rows_lists_flagged_rows():
    rows = make_delivered_rows()
    rows[0] = delivered_row(
        row=1, human_sentence_correct="yes", human_quote_supports="yes",
        location_section_mismatch="yes",
    )
    assert {r["row"] for r in sc.delivered_location_mismatch_rows(rows)} == {1}


def test_delivered_location_mismatch_rows_empty_when_none_flagged():
    assert sc.delivered_location_mismatch_rows(make_delivered_rows()) == []


# delivered_fidelity_excluded_rows and delivered_rows_judged_no both carry a "section" key
# so a row is identifiable on a multi-section sheet where "row" restarts at 1 in every
# section; location_mismatch_rows must do the same.
def test_delivered_location_mismatch_rows_carries_section_label():
    rows = make_delivered_rows()
    rows[0] = delivered_row(
        row=1, section_title="Introduction", human_sentence_correct="yes",
        human_quote_supports="yes", location_section_mismatch="yes",
    )
    result = sc.delivered_location_mismatch_rows(rows)
    assert result[0]["section"] == "Introduction"


def test_score_delivered_location_mismatch_rows_present_and_not_excluded():
    rows = make_delivered_rows()
    rows[0] = delivered_row(
        row=1, human_sentence_correct="yes", human_quote_supports="yes",
        location_section_mismatch="yes",
    )
    result = sc.score_delivered(rows)
    assert {r["row"] for r in result["location_mismatch_rows"]} == {1}
    # Flagged on its own, with no other fidelity flag set to no, so still in the headline.
    assert result["n_judged"] == 4
    assert result["n_fidelity_excluded"] == 0


# --------------------------------------------------------------------------------------
# Markdown summary smoke test
# --------------------------------------------------------------------------------------


def test_render_markdown_summary_contains_key_sections():
    result = sc.score(make_rows(), make_key())
    md = sc.render_markdown_summary(result)
    assert "Agreement with the verifier" in md
    assert "Disagreement axis counts" in md


def test_render_markdown_summary_includes_screening_section_when_present():
    result = {"screening": sc.score_screening([screening_row()])}
    md = sc.render_markdown_summary(result)
    assert "Screening sheet" in md


# The Delivered sheet is the final product a user receives, so its own numbers are the
# headline, printed above the Review/Screening component numbers, not folded in with them.
def test_render_markdown_summary_delivered_block_is_the_headline_above_the_rest():
    result = {
        "delivered": sc.score_delivered(make_delivered_rows()),
        **sc.score(make_rows(), make_key()),
    }
    md = sc.render_markdown_summary(result)
    assert "Delivered" in md
    delivered_at = md.index("Delivered")
    agreement_at = md.index("Agreement with the verifier")
    assert delivered_at < agreement_at


def test_render_markdown_summary_delivered_block_states_counts_and_rows_judged_no():
    result = {"delivered": sc.score_delivered(make_delivered_rows())}
    md = sc.render_markdown_summary(result)
    assert "4" in md  # n_rows
    assert "0.75" in md or "75" in md  # sentence_correct_rate / quote_supports_rate
    assert "2" in md  # rows judged no on at least one question


def test_render_markdown_summary_omits_delivered_block_when_absent():
    result = sc.score(make_rows(), make_key())
    md = sc.render_markdown_summary(result)
    assert "Delivered" not in md


def test_render_markdown_summary_delivered_block_states_fidelity_excluded_rows():
    rows = make_delivered_rows()
    rows[0] = delivered_row(
        row=1, human_sentence_correct="yes", human_quote_supports="yes",
        source_located="no",
    )
    result = {"delivered": sc.score_delivered(rows)}
    md = sc.render_markdown_summary(result)
    assert "excluded" in md
    assert "1 row(s)" in md or "1 excluded" in md


def test_render_markdown_summary_delivered_block_states_location_mismatch_rows():
    rows = make_delivered_rows()
    rows[0] = delivered_row(
        row=1, human_sentence_correct="yes", human_quote_supports="yes",
        location_section_mismatch="yes",
    )
    result = {"delivered": sc.score_delivered(rows)}
    md = sc.render_markdown_summary(result)
    assert "location_section_mismatch" in md
    assert "1 row(s)" in md


def test_render_markdown_summary_delivered_block_no_location_mismatch_rows_by_default():
    result = {"delivered": sc.score_delivered(make_delivered_rows())}
    md = sc.render_markdown_summary(result)
    assert "no rows flagged location_section_mismatch" in md


def test_render_markdown_summary_delivered_block_location_mismatch_table_has_section_column():
    rows = make_delivered_rows()
    rows[0] = delivered_row(
        row=1, section_title="Introduction", human_sentence_correct="yes",
        human_quote_supports="yes", location_section_mismatch="yes",
    )
    result = {"delivered": sc.score_delivered(rows)}
    md = sc.render_markdown_summary(result)
    assert "Introduction" in md


def test_render_markdown_summary_delivered_block_includes_per_section_table():
    rows = [
        delivered_row(row=1, section_title="Introduction",
                      human_sentence_correct="yes", human_quote_supports="yes"),
        delivered_row(row=2, section_title="Literature Review",
                      human_sentence_correct="yes", human_quote_supports="yes"),
    ]
    result = {"delivered": sc.score_delivered(rows)}
    md = sc.render_markdown_summary(result)
    assert "Introduction" in md
    assert "Literature Review" in md
    assert "2" in md  # n_sections


def test_render_markdown_summary_delivered_block_per_section_table_absent_sections_key():
    # score_delivered always fills in n_sections/by_section, so the markdown renderer must
    # not assume the older shape (no "by_section" key at all) can still occur; this
    # documents that the renderer reads it via .get, tolerating a hand-built dict that omits
    # it (e.g. an older cached json re-rendered).
    result = {"delivered": {**sc.score_delivered(make_delivered_rows())}}
    del result["delivered"]["by_section"]
    del result["delivered"]["n_sections"]
    sc.render_markdown_summary(result)  # must not raise


# agreement_with_verifier.per_status must give its own no_full_text bucket the same nested
# treatment score_screening gives a guard-demoted needs_review bucket, and the markdown must
# state it next to the per-status table.
def test_render_markdown_summary_states_no_full_text_withheld_split_in_per_status():
    key = make_key()
    key["R-004"] = {
        "real_item_id": "hss-wrong_paper-01", "set": "constructed",
        "construction_category": "wrong_paper", "expected_label": ["no_full_text"],
    }
    rows = make_rows() + [
        base_row(
            item_id="R-004", verifier_status="no_full_text", human_verdict="no_full_text",
            human_agrees_with_verifier="yes",
        ),
    ]
    result = sc.score(rows, key)
    md = sc.render_markdown_summary(result)
    lines = md.splitlines()
    header = next(line for line in lines if line.startswith("### Per verifier_status"))
    assert "no_full_text excludes its own withheld row(s)" in header
    assert any(
        "no_full_text's own" in line and "withheld row(s)" in line for line in lines
    )


def test_render_markdown_summary_per_verifier_status_header_plain_when_nothing_withheld():
    result = sc.score(make_rows(), make_key())
    md = sc.render_markdown_summary(result)
    lines = md.splitlines()
    header = next(line for line in lines if line.startswith("### Per verifier_status"))
    assert header == "### Per verifier_status"


# The screening "Per status" header and the quote_is_verbatim rate clause must not claim an
# unanchored_exclude exclusion happened on a run that sampled none of that reason; both must
# gate on n_unanchored_exclude the same way the per-status nested "needs_review's own ..."
# prose line already does.
def test_render_markdown_summary_screening_per_status_header_plain_when_none_sampled():
    result = {"screening": sc.score_screening([screening_row(record_id="S-001")])}
    md = sc.render_markdown_summary(result)
    lines = md.splitlines()
    header = next(line for line in lines if line.startswith("### Per status"))
    assert header == "### Per status"
    quote_line = next(line for line in lines if line.startswith("quote_is_verbatim rate"))
    assert "unanchored_exclude" not in quote_line


def test_render_markdown_summary_screening_per_status_header_states_split_when_sampled():
    rows = [
        screening_row(record_id="S-001", status="needs_review", human_status="needs_review"),
        screening_row(
            record_id="S-002", status="needs_review", human_status="exclude",
            human_agrees="no", human_note="guard demoted this",
            needs_review_reason="unanchored_exclude",
        ),
    ]
    result = {"screening": sc.score_screening(rows)}
    md = sc.render_markdown_summary(result)
    lines = md.splitlines()
    header = next(line for line in lines if line.startswith("### Per status"))
    assert "guard-demoted unanchored_exclude row(s)" in header
    quote_line = next(line for line in lines if line.startswith("quote_is_verbatim rate"))
    assert "unanchored_exclude" in quote_line


# The markdown summary must state the unanchored_exclude split next to both the headline
# overall accuracy and the per-status table.
def test_render_markdown_summary_states_unanchored_exclude_split_next_to_both_numbers():
    rows = [
        screening_row(record_id="S-001", status="needs_review", human_status="needs_review"),
        screening_row(
            record_id="S-002", status="needs_review", human_status="exclude",
            human_agrees="no", human_note="guard demoted this",
            needs_review_reason="unanchored_exclude",
        ),
    ]
    result = {"screening": sc.score_screening(rows)}
    md = sc.render_markdown_summary(result)
    assert "guard-demoted unanchored_exclude row(s)" in md
    lines = md.splitlines()
    overall_line = next(line for line in lines if line.startswith("overall accuracy"))
    assert "unanchored_exclude" in overall_line
    assert any(
        "needs_review's own" in line and "unanchored_exclude" in line for line in lines
    )


# --------------------------------------------------------------------------------------
# End-to-end: export then score a tiny in-memory workbook
# --------------------------------------------------------------------------------------


CONSTRUCTED_ITEM = {
    "item_id": "hss-verbatim-01", "rule": "verbatim", "expected": ["verified"],
    "claim": "Scores were significantly higher for the treatment group after training.",
    "source_doi": "10.1000/paper-a", "chunk_doi": "10.1000/paper-a", "chunk_index": 1,
    "title": "Paper A", "authors": ["A. One"],
}
REAL_ITEM = {
    "item_id": "real-test-01", "rule": "real", "expected": None,
    "claim": "Feedback improved learner motivation across the term for most participants.",
    "cited_title": "Paper B", "cited_doi": "10.1000/paper-b",
    "source_doi": "10.1000/paper-b", "chunk_doi": "10.1000/paper-b", "chunk_index": 0,
    "title": "Paper B", "authors": ["B. Two"],
}


def write_jsonl(path: Path, rows) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as fh:
        for row in rows:
            fh.write(json.dumps(row) + "\n")


def result_row(item_id, status):
    return {
        "item_id": item_id, "predicted_status": status, "model_status": status,
        "evidence_quotes": ["an evidence quote"], "evidence_quote": "an evidence quote",
        "machine_reasons": [], "diagnostics": [], "unstated_details": [], "assertions": [],
    }


def test_end_to_end_export_then_score(tmp_path: Path):
    items_path = tmp_path / "items.jsonl"
    write_jsonl(items_path, [CONSTRUCTED_ITEM, REAL_ITEM])
    cache_dir = tmp_path / "cache"
    cache_dir.mkdir()
    results_dir = tmp_path / "results"
    write_jsonl(results_dir / "hss_runA.jsonl", [result_row("hss-verbatim-01", "verified")])
    write_jsonl(results_dir / "real_runA.jsonl", [result_row("real-test-01", "unsupported")])

    item_specs = ex.resolve_item_specs([items_path], [cache_dir])
    wb, kept_items, opaque_ids, missing, workbook_id = ex.build_workbook(
        item_specs, results_dir, ["A"], max_cell_chars=32000, sources_max_chars=30000,
        allow_missing_chunks=True,
    )
    out_path = tmp_path / "workbook.xlsx"
    wb.save(out_path)
    key_path = tmp_path / "key.csv"
    ex.write_key_csv(key_path, kept_items, opaque_ids, workbook_id=workbook_id)

    # simulate a human filling in the sheet: agree on both rows
    filled = openpyxl.load_workbook(out_path)
    ws = filled["Review"]
    header = [c.value for c in ws[1]]
    verdict_col = header.index("human_verdict") + 1
    agrees_col = header.index("human_agrees_with_verifier") + 1
    status_col = header.index("verifier_status") + 1
    for r in range(2, ws.max_row + 1):
        ws.cell(row=r, column=verdict_col, value=ws.cell(row=r, column=status_col).value)
        ws.cell(row=r, column=agrees_col, value="yes")
    filled.save(out_path)

    rows = sc.read_review_sheet(out_path)
    key = sc.load_key(key_path)
    review_workbook_id = sc.read_workbook_id(out_path)
    key_workbook_id = sc.load_key_workbook_id(key_path)
    result = sc.score(
        rows, key, review_workbook_id=review_workbook_id, key_workbook_id=key_workbook_id,
    )
    assert result["agreement_with_verifier"]["overall"]["accuracy"] == 1.0
    assert result["n_items"] == 2


# The CLI must report the workbook's own demo_run_id/export_time and refuse a --demo-run-id
# that does not match the workbook's own Provenance.
def test_main_reports_demo_run_id_and_export_time_and_refuses_on_mismatch(tmp_path: Path):
    items_path = tmp_path / "items.jsonl"
    write_jsonl(items_path, [CONSTRUCTED_ITEM, REAL_ITEM])
    cache_dir = tmp_path / "cache"
    cache_dir.mkdir()
    results_dir = tmp_path / "results"
    write_jsonl(results_dir / "hss_runA.jsonl", [result_row("hss-verbatim-01", "verified")])
    write_jsonl(results_dir / "real_runA.jsonl", [result_row("real-test-01", "unsupported")])

    item_specs = ex.resolve_item_specs([items_path], [cache_dir])
    wb, kept_items, opaque_ids, missing, workbook_id = ex.build_workbook(
        item_specs, results_dir, ["A"], max_cell_chars=32000, sources_max_chars=30000,
        allow_missing_chunks=True, demo_run_id="20260910-095226",
        export_time="2026-09-10T12:00:00+00:00",
    )
    out_path = tmp_path / "workbook.xlsx"
    wb.save(out_path)
    key_path = tmp_path / "key.csv"
    ex.write_key_csv(key_path, kept_items, opaque_ids, workbook_id=workbook_id)

    filled = openpyxl.load_workbook(out_path)
    ws = filled["Review"]
    header = [c.value for c in ws[1]]
    verdict_col = header.index("human_verdict") + 1
    agrees_col = header.index("human_agrees_with_verifier") + 1
    status_col = header.index("verifier_status") + 1
    for r in range(2, ws.max_row + 1):
        ws.cell(row=r, column=verdict_col, value=ws.cell(row=r, column=status_col).value)
        ws.cell(row=r, column=agrees_col, value="yes")
    filled.save(out_path)

    mismatch_out = tmp_path / "mismatch.json"
    exit_code = sc.main([
        "--workbook", str(out_path), "--key", str(key_path), "--out", str(mismatch_out),
        "--no-markdown", "--demo-run-id", "some-other-run",
    ])
    assert exit_code == 2
    assert not mismatch_out.exists()

    match_out = tmp_path / "match.json"
    exit_code = sc.main([
        "--workbook", str(out_path), "--key", str(key_path), "--out", str(match_out),
        "--no-markdown", "--demo-run-id", "20260910-095226",
    ])
    assert exit_code == 0
    result = json.loads(match_out.read_text(encoding="utf-8"))
    assert result["demo_run_id"] == "20260910-095226"
    assert result["export_time"] == "2026-09-10T12:00:00+00:00"

    # A matching --demo-run-id alone must not be enough: two exports of the same demo run
    # share it (and the workbook id), so a wrong --export-time must still refuse.
    export_time_mismatch_out = tmp_path / "export-time-mismatch.json"
    exit_code = sc.main([
        "--workbook", str(out_path), "--key", str(key_path),
        "--out", str(export_time_mismatch_out), "--no-markdown",
        "--demo-run-id", "20260910-095226", "--export-time", "2026-09-09T07:13:45+00:00",
    ])
    assert exit_code == 2
    assert not export_time_mismatch_out.exists()

    export_time_match_out = tmp_path / "export-time-match.json"
    exit_code = sc.main([
        "--workbook", str(out_path), "--key", str(key_path),
        "--out", str(export_time_match_out), "--no-markdown",
        "--demo-run-id", "20260910-095226", "--export-time", "2026-09-10T12:00:00+00:00",
    ])
    assert exit_code == 0
    assert json.loads(export_time_match_out.read_text(encoding="utf-8"))["export_time"] == (
        "2026-09-10T12:00:00+00:00"
    )


# End to end: a workbook with a Delivered sheet, filled in and scored through main().
def test_main_scores_delivered_sheet_end_to_end(tmp_path: Path):
    items_path = tmp_path / "items.jsonl"
    write_jsonl(items_path, [CONSTRUCTED_ITEM, REAL_ITEM])
    cache_dir = tmp_path / "cache"
    cache_dir.mkdir()
    results_dir = tmp_path / "results"
    write_jsonl(results_dir / "hss_runA.jsonl", [result_row("hss-verbatim-01", "verified")])
    write_jsonl(results_dir / "real_runA.jsonl", [result_row("real-test-01", "unsupported")])

    item_specs = ex.resolve_item_specs([items_path], [cache_dir])
    wb, kept_items, opaque_ids, missing, workbook_id = ex.build_workbook(
        item_specs, results_dir, ["A"], max_cell_chars=32000, sources_max_chars=30000,
        allow_missing_chunks=True,
    )
    ex.build_delivered_sheet(wb, DELIVERED_FIXTURE_PATH, max_cell_chars=32000)
    out_path = tmp_path / "workbook.xlsx"
    wb.save(out_path)
    key_path = tmp_path / "key.csv"
    ex.write_key_csv(key_path, kept_items, opaque_ids, workbook_id=workbook_id)

    filled = openpyxl.load_workbook(out_path)
    review_ws = filled["Review"]
    header = [c.value for c in review_ws[1]]
    verdict_col = header.index("human_verdict") + 1
    agrees_col = header.index("human_agrees_with_verifier") + 1
    status_col = header.index("verifier_status") + 1
    for r in range(2, review_ws.max_row + 1):
        review_ws.cell(row=r, column=verdict_col,
                       value=review_ws.cell(row=r, column=status_col).value)
        review_ws.cell(row=r, column=agrees_col, value="yes")
    delivered_ws = filled["Delivered"]
    d_header = [c.value for c in delivered_ws[1]]
    correct_col = d_header.index("human_sentence_correct") + 1
    supports_col = d_header.index("human_quote_supports") + 1
    for r in range(2, delivered_ws.max_row + 1):
        delivered_ws.cell(row=r, column=correct_col, value="yes")
        delivered_ws.cell(row=r, column=supports_col, value="yes")
    filled.save(out_path)

    out = tmp_path / "score.json"
    exit_code = sc.main([
        "--workbook", str(out_path), "--key", str(key_path), "--out", str(out),
    ])
    assert exit_code == 0
    result = json.loads(out.read_text(encoding="utf-8"))
    assert result["delivered"]["n_rows"] == 2
    assert result["delivered"]["sentence_correct_rate"] == 1.0
    assert result["delivered"]["quote_supports_rate"] == 1.0
    md = out.with_suffix(".md").read_text(encoding="utf-8")
    assert "Delivered" in md


# --------------------------------------------------------------------------------------
# v4: claim_checked column pass-through and the per-sentence secondary block
# --------------------------------------------------------------------------------------


def test_read_delivered_sheet_reads_claim_checked_and_paper_authors(tmp_path: Path):
    wb = openpyxl.Workbook()
    ex.build_delivered_sheet(
        wb, DELIVERED_FIXTURE_PATH, max_cell_chars=32000,
        paper_authors={"10.1234/fixture.2026": "Ada Fixture; Sam Sample (2026)"},
    )
    path = tmp_path / "workbook.xlsx"
    wb.save(path)
    rows = sc.read_delivered_sheet(path)
    assert rows[0]["claim_checked"]
    assert rows[0]["paper_authors"] == "Ada Fixture; Sam Sample (2026)"


def _sentence_group_rows():
    return [
        delivered_row(row=1, sentence="sentence A", human_sentence_correct="yes",
                      human_quote_supports="yes"),
        delivered_row(row=2, sentence="sentence A", human_sentence_correct="no",
                      human_quote_supports="yes", human_note="a reason"),
        delivered_row(row=3, sentence="sentence B", human_sentence_correct="yes",
                      human_quote_supports="yes"),
    ]


def test_delivered_sentence_groups_groups_by_draft_and_sentence_in_first_seen_order():
    rows = _sentence_group_rows()
    for r in rows:
        r["draft_id"] = "d1"
    groups = sc.delivered_sentence_groups(rows)
    assert [key[1] for key, _ in groups] == ["sentence A", "sentence B"]
    assert len(groups[0][1]) == 2
    assert len(groups[1][1]) == 1


def test_delivered_sentence_groups_drops_fidelity_failed_rows():
    rows = [
        delivered_row(row=1, sentence="s1", source_located="no"),
        delivered_row(row=2, sentence="s1"),
    ]
    groups = sc.delivered_sentence_groups(rows)
    assert len(groups) == 1
    assert len(groups[0][1]) == 1


def test_delivered_sentence_groups_omits_a_sentence_whose_every_row_is_excluded():
    rows = [delivered_row(row=1, sentence="s1", source_located="no")]
    groups = sc.delivered_sentence_groups(rows)
    assert groups == []


def test_score_delivered_by_sentence_counts_a_sentence_yes_only_when_every_row_is_yes():
    rows = _sentence_group_rows()
    result = sc.score_delivered_by_sentence(rows)
    assert result["n_sentences"] == 2
    assert result["sentence_correct"]["n_yes"] == 1
    assert result["sentence_correct"]["n"] == 2


def test_score_delivered_by_sentence_one_no_row_makes_the_sentence_no():
    rows = _sentence_group_rows()
    result = sc.score_delivered_by_sentence(rows)
    assert result["quote_supports"]["n_yes"] == 2  # both sentences all-yes on this question


def test_score_delivered_by_sentence_reports_a_wilson_interval_for_both_columns():
    rows = _sentence_group_rows()
    result = sc.score_delivered_by_sentence(rows)
    assert "lo" in result["sentence_correct"]["wilson_95"]
    assert "hi" in result["quote_supports"]["wilson_95"]


def test_score_delivered_by_sentence_counts_the_same_text_in_two_drafts_twice():
    rows = [
        delivered_row(row=1, sentence="same", human_sentence_correct="yes",
                      human_quote_supports="yes"),
        delivered_row(row=2, sentence="same", human_sentence_correct="yes",
                      human_quote_supports="yes"),
    ]
    rows[0]["draft_id"] = "d1"
    rows[1]["draft_id"] = "d2"
    result = sc.score_delivered_by_sentence(rows)
    assert result["n_sentences"] == 2


def test_score_delivered_includes_by_sentence_without_changing_the_headline_keys():
    result = sc.score_delivered(make_delivered_rows())
    assert "by_sentence" in result
    # make_delivered_rows's own rows all share one sentence text/draft_id, so they form one
    # sentence group, made no by row 3's own no.
    assert result["by_sentence"]["n_sentences"] == 1
    assert result["by_sentence"]["sentence_correct"]["n_yes"] == 0
    assert result["n_rows"] == 4
    assert result["sentence_correct_rate"] == pytest.approx(0.75)


def test_score_delivered_rows_judged_no_carries_claim_checked():
    rows = make_delivered_rows()
    rows[1]["claim_checked"] = "the quote-only claim"
    result = sc.score_delivered(rows)
    judged_no = {r["row"]: r for r in result["rows_judged_no"]}
    assert judged_no[2]["claim_checked"] == "the quote-only claim"


def test_score_delivered_fidelity_excluded_rows_carry_claim_checked():
    rows = [
        delivered_row(row=1, sentence="s1", source_located="no",
                      claim_checked="the excluded claim"),
    ]
    result = sc.score_delivered(rows)
    assert result["fidelity_excluded_rows"][0]["claim_checked"] == "the excluded claim"


def test_score_delivered_scores_a_workbook_with_no_claim_checked_column():
    rows = make_delivered_rows()  # delivered_row() never sets claim_checked
    result = sc.score_delivered(rows)
    assert result["rows_judged_no"][0]["claim_checked"] == ""


def test_render_markdown_summary_prints_the_per_sentence_block_under_the_headline():
    result = {"delivered": sc.score_delivered(make_delivered_rows())}
    md = sc.render_markdown_summary(result)
    assert "### Per delivered sentence (secondary)" in md
    headline_idx = md.index("sentence_correct: ")
    secondary_idx = md.index("### Per delivered sentence (secondary)")
    assert headline_idx < secondary_idx


def test_render_markdown_summary_prints_the_claim_granularity_disclosure():
    result = {"delivered": sc.score_delivered(make_delivered_rows())}
    md = sc.render_markdown_summary(result)
    assert "Claim granularity is the tool's own" in md


def test_render_markdown_summary_without_by_sentence_still_renders():
    d = sc.score_delivered(make_delivered_rows())
    del d["by_sentence"]
    result = {"delivered": d}
    md = sc.render_markdown_summary(result)
    assert "### Per delivered sentence (secondary)" not in md
    assert "Delivered" in md


def test_main_still_refuses_a_provenance_demo_run_id_mismatch_on_a_delivered_only_workbook(
    tmp_path: Path,
):
    authors = ex.load_paper_authors(
        Path(__file__).resolve().parent / "fixtures" / "paper_authors_sample.json"
    )
    wb, _record, _rows, workbook_id = ex.build_delivered_only_workbook(
        DELIVERED_FIXTURE_PATH, max_cell_chars=32000, demo_run_id="run-actual",
        export_time="2026-09-16T00:00:00Z", paper_authors=authors,
    )
    out_path = tmp_path / "wb.xlsx"
    wb.save(out_path)
    out = tmp_path / "score.json"
    rc = sc.main([
        "--workbook", str(out_path), "--out", str(out),
        "--demo-run-id", "run-wrong", "--export-time", "2026-09-16T00:00:00Z",
    ])
    assert rc == 2
    assert not out.exists()
